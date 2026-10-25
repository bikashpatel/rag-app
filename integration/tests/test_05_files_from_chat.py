"""Flow 5: files from chat. The fake calls create_document (arguments built from the schema the backend sent), the
backend renders the file in its child process, stores it as the user's private generated document, replays the call
with the tool result, the worker indexes it, and the download is a valid file of that format."""

import io
import json

import pytest
from docx import Document as WordDocument
from harness import assert_sse_contract, chat, check, system_prompt, unique, upload, wait_document
from openpyxl import load_workbook
from pypdf import PdfReader
from samples import txt

WORDS = {"pdf": "PDF", "docx": "Word", "xlsx": "Excel", "md": "Markdown"}
CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "md": "application/octet-stream",
}


def _file_text(fmt: str, data: bytes) -> str:
    """The text of a downloaded file, parsed with the library for its format."""
    if fmt == "pdf":
        reader = PdfReader(io.BytesIO(data))
        assert len(reader.pages) >= 1
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if fmt == "docx":
        document = WordDocument(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if fmt == "xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True)
        sheet = workbook["Data"]
        rows = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
        assert rows[0] == ("Item", "Value") or rows[0] == ["Item", "Value"]
        assert rows[1][0] == "alpha" and rows[1][1] in (1, "1")
        return "\n".join(" ".join(str(cell) for cell in row if cell is not None) for row in rows)
    return data.decode("utf-8")


@pytest.mark.deadline(240)
@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx", "md"])
def test_file_from_chat_is_stored_indexed_and_downloadable(make_user, fake, fmt: str) -> None:
    user = make_user(f"files-{fmt}")
    marker = unique("lantern")
    source = upload(user.client, f"{marker}.txt", txt(f"The {marker} lantern inventory counts forty brass lanterns."))
    wait_document(user.client, source["id"])
    title = f"Lantern report {marker}"
    message = f'Please create a {WORDS[fmt]} file titled "{title}" summarizing the {marker} lantern inventory.'

    mark = fake.mark()
    result = chat(user.client, message)
    assert_sse_contract(result)
    assert result.names[-1] == "done", result.events[-1]
    events = result.all("file")
    statuses = [event["status"] for event in events]
    assert statuses[0] == "writing" and statuses[-1] == "created" and "failed" not in statuses, statuses
    assert all(event["call_id"] == "f1" and event["kind"] == "file" for event in events)
    created = events[-1]
    assert created["file"]["format"] == fmt and created["file"]["filename"] == f"{title}.{fmt}"
    assert created["already_existed"] is False
    doc = created["document"]
    assert doc["generated"] is True and doc["scope"] == "private" and doc["owner_id"] == user.id

    first, second = fake.log("chat", after=mark)
    call = first["tool_call"]
    assert call["name"] == "create_document" and call["missing_required"] == []
    assert call["arguments"]["format"] == fmt and call["arguments"]["source_numbers"] == [1]
    assert first["parallel_tool_calls"] is False and first["tools"] == second["tools"]
    assistant, tool = second["messages"][-2], second["messages"][-1]
    replayed = assistant["tool_calls"][0]
    assert assistant["role"] == "assistant" and replayed["function"]["name"] == "create_document"
    assert json.loads(replayed["function"]["arguments"]) == call["arguments"]
    assert tool["role"] == "tool" and tool["tool_call_id"] == replayed["id"]
    assert "error" not in json.loads(tool["content"])
    assert second["answer"] in result.text
    usage = result.done["usage"]
    assert usage["prompt_tokens"] == first["usage"]["prompt_tokens"] + second["usage"]["prompt_tokens"]
    assert usage["completion_tokens"] == first["usage"]["completion_tokens"] + second["usage"]["completion_tokens"]

    ready = wait_document(user.client, doc["id"])
    assert ready["num_chunks"] >= 1 and ready["generated"] is True
    response = check(user.client.get(f"/api/documents/{doc['id']}/download"))
    assert response.headers["content-type"] == CONTENT_TYPES[fmt]
    assert response.headers["content-disposition"].startswith("attachment;")
    # nginx adds its own nosniff on top of the backend's (a known, harmless repeat; docs/HANDOFF.md)
    assert set(response.headers["x-content-type-options"].replace(" ", "").split(",")) == {"nosniff"}
    assert len(response.content) == ready["size_bytes"]
    text = _file_text(fmt, response.content)
    assert title in text
    if fmt != "xlsx":
        assert marker in text  # the passage the model used, with its [1] and a Sources section

    detail = check(user.client.get(f"/api/conversations/{result.meta['conversation_id']}")).json()
    (saved,) = detail["messages"][-1]["files"]
    assert (saved["document_id"], saved["format"], saved["status"]) == (doc["id"], fmt, "ready")

    other = make_user("files-other")
    assert other.client.get(f"/api/documents/{doc['id']}/download").status_code == 404

    mark = fake.mark()
    later = chat(user.client, f"What does the lantern report {marker} say?", document_ids=[doc["id"]])
    assert later.sources and later.sources[0]["document_id"] == doc["id"]
    assert "a file the assistant created on" in system_prompt(fake.log("chat", after=mark)[0])


def test_no_file_without_a_request(make_user, fake) -> None:
    user = make_user("files-none")
    mark = fake.mark()
    result = chat(user.client, "Summarize what a lantern is in one sentence.")
    assert result.names[-1] == "done" and result.all("file") == []
    (call,) = fake.log("chat", after=mark)
    assert "create_document" in call["tools"] and "tool_call" not in call
