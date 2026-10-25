"""Test documents, built at test time: one per supported upload format, each carrying the given text."""

import csv
import io
import json

from docx import Document as WordDocument
from fpdf import FPDF
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


def pdf(text: str) -> bytes:
    document = FPDF()
    document.add_page()
    document.set_font("helvetica", size=12)
    document.multi_cell(0, 8, text)
    return bytes(document.output())


def scanned_pdf(text: str) -> bytes:
    """A one-page PDF that is only an image of the text (no text layer), like a scan."""
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    draw.multiline_text((80, 120), text, fill="black", font=ImageFont.load_default(size=40), spacing=18)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    document = FPDF(unit="pt", format="A4")
    document.add_page()
    document.image(buffer, x=0, y=0, w=document.w, h=document.h)
    return bytes(document.output())


def docx(text: str) -> bytes:
    document = WordDocument()
    document.add_heading("Integration test document", level=1)
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx(text: str) -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Integration test slide"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    box.text_frame.text = text
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def xlsx(text: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Notes"
    sheet.append(["Topic", "Note"])
    sheet.append(["integration", text])
    sheet.append(["count", 3])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def html(text: str) -> bytes:
    return (
        "<!doctype html><html><head><title>Integration test page</title><style>p{color:red}</style></head>"
        f"<body><h1>Integration test page</h1><p>{text}</p><script>var ignored = 1;</script></body></html>"
    ).encode()


def md(text: str) -> bytes:
    return f"# Integration test notes\n\n{text}\n\n- one\n- two\n".encode()


def txt(text: str) -> bytes:
    return f"Integration test notes\n\n{text}\n".encode()


def csv_bytes(text: str) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["topic", "note"])
    writer.writerow(["integration", text])
    return buffer.getvalue().encode()


def json_bytes(text: str) -> bytes:
    return json.dumps({"title": "Integration test notes", "items": [{"note": text}, {"note": "second"}]}).encode()


BUILDERS = {
    ".pdf": pdf,
    ".docx": docx,
    ".pptx": pptx,
    ".xlsx": xlsx,
    ".html": html,
    ".htm": html,
    ".md": md,
    ".markdown": md,
    ".txt": txt,
    ".csv": csv_bytes,
    ".json": json_bytes,
}


def sample(extension: str, text: str) -> bytes:
    return BUILDERS[extension](text)
