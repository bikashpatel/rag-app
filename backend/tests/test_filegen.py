"""The files the chat assistant creates (app.services.filegen): every format parsed back, the limits, the fonts and
scripts of PDFs, no active content, and rendering in the time- and memory-limited child process."""

import asyncio
import io
import json
import logging
import re
import time
import zipfile
from datetime import datetime, timedelta, timezone
from typing import Any

import docx
import openpyxl
import pypdfium2 as pdfium
import pytest
from conftest import FileChild
from docx.oxml.ns import qn
from pypdf import PdfReader

from app.config import settings
from app.services import document_tool, filegen
from app.services.filegen import RenderError, RenderRequest, SheetData, SourceRef, file_name
from app.services.filegen import fonts as fontmod
from app.services.filegen.markdown import parse
from app.services.filegen.mdwriter import write

SAMPLE = """\
# Quarterly report

Revenue grew **12%** in the *northern* region [1]. ~~Old figure~~ and `inline code`.
See [the dashboard](https://www.example.com/q3?a=1), <https://example.org/>, [mail us](mailto:team@example.com) and
[a script](javascript:alert(1)).

## Lists

1. First item
2. Second item
   - Nested bullet
     - Deeper bullet

- Plain bullet

> A quoted remark.

```
def total(values):
    return sum(values)
```

---

| Region | Revenue |
|:--|--:|
| North | 1200 |
| South | 980 |

![chart](https://evil.example/x.png) <img src=x onerror=alert(1)>
"""
ACTIVE_CONTENT = (
    b"/JavaScript",
    b"/JS",
    b"/OpenAction",
    b"/AA",
    b"/Launch",
    b"/EmbeddedFile",
    b"/RichMedia",
    b"/AcroForm",
)


def request(fmt: str = "pdf", content: str = SAMPLE, **fields: Any) -> RenderRequest:
    return RenderRequest(
        title=fields.get("title", "Quarterly report"),
        format=fmt,  # type: ignore[arg-type]
        content=content,
        sheets=fields.get("sheets", []),
        sources=fields.get("sources", []),
    )


def pdf_text(data: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)


def pdfium_text(data: bytes) -> str:
    document = pdfium.PdfDocument(data)
    try:
        return "\n".join(document[number].get_textpage().get_text_range() for number in range(len(document)))
    finally:
        document.close()


def link_targets(data: bytes) -> list[str]:
    targets = []
    for page in PdfReader(io.BytesIO(data)).pages:
        for annotation in page.get("/Annots") or []:
            action = annotation.get_object().get("/A") or {}
            targets.append(str(action.get("/URI")))
    return targets


# ---------------- PDF ----------------


def test_pdf_has_every_construct_and_no_active_content() -> None:
    rendered = filegen.render(request(sources=[SourceRef([1], "notes.pdf", 2)]))

    assert rendered.data.startswith(b"%PDF")
    assert rendered.pages == len(PdfReader(io.BytesIO(rendered.data)).pages) == 1
    text = pdf_text(rendered.data)
    for expected in (
        "Quarterly report",
        "Revenue grew 12% in the northern region [1].",
        "First item",
        "Nested bullet",
        "Deeper bullet",
        "A quoted remark.",
        "return sum(values)",
        "North",
        "1200",
        "[Image: chart]",
        "<img src=x onerror=alert(1)>",
        "Sources",
        "[1] notes.pdf, page 2",
        "1 / 1",
    ):
        assert expected in text, expected
    assert text.count("Quarterly report") == 1  # the content's own title heading isn't doubled
    assert not any(marker in rendered.data for marker in ACTIVE_CONTENT)
    assert link_targets(rendered.data) == [
        "https://www.example.com/q3?a=1",
        "https://example.org/",
        "mailto:team@example.com",
    ]
    flat = " ".join(text.split())
    assert "the dashboard (example.com)" in flat and "mail us (team@example.com)" in flat
    assert "https://example.org/ (" not in flat  # the text is the URL: no host after it
    assert rendered.source_markdown == filegen.document_markdown(request(sources=[SourceRef([1], "notes.pdf", 2)]))
    reader = PdfReader(io.BytesIO(rendered.data))
    assert reader.metadata is not None
    assert (reader.metadata.title, reader.metadata.creator, reader.metadata.author) == (
        "Quarterly report",
        "AskNuvexusDoc",
        None,
    )


def test_pdf_title_heading_is_added_when_the_content_has_none() -> None:
    rendered = filegen.render(request(content="Just a paragraph.", title="Board letter"))

    text = pdf_text(rendered.data)
    assert text.index("Board letter") < text.index("Just a paragraph.")


def test_long_pdf_gets_more_pages_with_footers() -> None:
    content = "# Long\n\n" + "\n\n".join(f"Paragraph {n} " + "lorem ipsum dolor " * 40 for n in range(60))

    rendered = filegen.render(request(content=content))

    assert rendered.pages is not None and rendered.pages > 1
    text = pdf_text(rendered.data)
    assert f"1 / {rendered.pages}" in text and f"{rendered.pages} / {rendered.pages}" in text


def test_long_code_block_keeps_its_lines_across_pages() -> None:
    code = "\n".join(f"line_{n} = {n}" for n in range(200))

    rendered = filegen.render(request(content=f"# Code\n\n```\n{code}\n```"))

    text = pdf_text(rendered.data)
    assert rendered.pages is not None and rendered.pages > 1
    assert "line_0 = 0" in text and "line_199 = 199" in text


@pytest.mark.parametrize(
    "table",
    [
        "| a |\n|---|\n" + "| x |\n" * 501,
        "|" + "|".join(f"h{n}" for n in range(21)) + "|\n|" + "|".join("-" * 21) + "|\n",
        "|"
        + "|".join(f"h{n}" for n in range(20))
        + "|\n|"
        + "|".join("-" * 20)
        + "|\n"
        + ("|" + "|".join("x" * 20) + "|\n") * 500
        + "\n\n| a |\n|---|\n| y |\n",
        "| a |\n|---|\n| " + "w" * 2001 + " |\n",
    ],
    ids=["501 rows", "21 columns", "10001 cells", "2001 characters"],
)
@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_tables_over_the_caps_are_refused(table: str, fmt: str) -> None:
    with pytest.raises(RenderError, match=r"at most 20 columns, 500 rows .* use xlsx for large tables"):
        filegen.render(request(fmt, content="# T\n\n" + table))


def test_tables_at_the_caps_render_but_md_has_no_caps() -> None:
    table = "|" + "|".join(f"h{n}" for n in range(20)) + "|\n|" + "|".join("-" * 20) + "|\n"
    table += ("|" + "|".join("x" * 20) + "|\n") * 500
    started = time.perf_counter()
    rendered = filegen.render(request(content="# T\n\n" + table))
    assert time.perf_counter() - started < 30  # speed guard for the capped table
    assert rendered.pages is not None and rendered.pages > 1
    huge = "| a |\n|---|\n" + "| x |\n" * 600
    assert filegen.render(request("md", content=huge)).data.count(b"| x |") == 600


def test_a_table_row_taller_than_a_page_is_a_render_error() -> None:
    columns = 20
    header = "|" + "|".join(f"h{n}" for n in range(columns)) + "|\n|" + "|".join("-" * columns) + "|\n"
    row = "|" + "|".join(" ".join(["word"] * 390) for _ in range(columns)) + "|\n"

    with pytest.raises(RenderError, match="A table row is too tall to fit on one page"):
        filegen.render(request(content="# T\n\n" + header + row))


def test_runs_only_put_characters_into_fonts_that_have_them() -> None:
    sample = (
        "Hello नमस्ते (क्षत्रिय) বাংলা ਪੰਜਾਬੀ ગુજરાતી ଓଡ଼ିଆ தமிழ் తెలుగు ಕನ್ನಡ മലയാളം සිංහල "
        "Ελληνικά Кириллица ≤ ≥ ≠ √ ∞ ∑ ⇒ ✓ ✔ ✗ ★ ● 12,5% — “quotes” 東京 😀👍🏽 ✔️ مرحبا"
    )
    runs = fontmod.Runs()
    for base in (fontmod.SANS, fontmod.MONO):
        for family, text in runs(sample, base):
            missing = [ch for ch in text if ch != "\n" and ord(ch) not in fontmod.cmap(family)]
            assert missing == [], (family, missing)
    assert runs.placeholders["Han"] == 4 and runs.placeholders["Arabic"] == 10


def test_lines_after_a_fallback_script_keep_their_font() -> None:
    # fpdf2 2.8.9 lost the font of lines after one starting with a fallback font: "Hello" and "Plain" vanished.
    content = "নমস্কার বাংলা\n\nHello নমস্কার\n\nPlain"

    rendered = filegen.render(request(content=content, title="নমস্কার"))

    text = pdfium_text(rendered.data)
    assert "Hello" in text and "Plain" in text


@pytest.mark.parametrize(
    ("sample", "font"),
    [
        ("नमस्ते दुनिया, यह एक रिपोर्ट है।", b"NotoSansDevanagari"),
        ("নমস্কার, এটি একটি প্রতিবেদন।", b"NotoSansBengali"),
        ("வணக்கம், இது ஒரு அறிக்கை.", b"NotoSansTamil"),
    ],
)
def test_indian_scripts_embed_their_noto_font(sample: str, font: bytes) -> None:
    rendered = filegen.render(request(content=f"# Report\n\n{sample}\n\n| A |\n|---|\n| {sample} |"))

    assert font in rendered.data
    assert rendered.notes == []


def test_latin_documents_load_only_the_fonts_they_need() -> None:
    rendered = filegen.render(request(content="# Plain\n\nOnly Latin text, without code."))

    assert b"NotoSans-Regular" in rendered.data or b"NotoSans" in rendered.data
    for family in (b"Devanagari", b"Bengali", b"Tamil", b"NotoSansMono", b"NotoSansMath", b"Symbols"):
        assert family not in rendered.data


def test_a_few_characters_of_an_unsupported_script_become_placeholders() -> None:
    content = "# Trip\n\n" + "We visited the capital, as planned, and met the regional team there. " * 3 + "東京"

    rendered = filegen.render(request(content=content))

    assert "□□" in pdfium_text(rendered.data)
    assert rendered.notes == [
        "2 characters in Han couldn't be shown in the PDF and appear as □. A docx or md file keeps them."
    ]


@pytest.mark.parametrize("block", ["```\nprint('東')\nx = 1\n```", "| a |\n|--|\n| **x** 東 |"])
def test_placeholders_are_counted_as_drawn(block: str) -> None:
    """A code block is drawn line by line and a table cell as one text: each placeholder counts once."""
    content = "# Notes\n\n" + "A paragraph of Latin text that keeps the share of Han letters low. " * 3 + "\n\n" + block

    rendered = filegen.render(request(content=content))

    assert rendered.notes == [
        "1 characters in Han couldn't be shown in the PDF and appear as □. A docx or md file keeps them."
    ]


@pytest.mark.parametrize(
    ("sample", "script"), [("مرحبا بالعالم " * 20, "Arabic"), ("東京は日本の首都です。" * 10, "Han")]
)
def test_mostly_unsupported_scripts_are_refused_with_a_hint(sample: str, script: str) -> None:
    with pytest.raises(RenderError, match=f"this content has {script} text. Create a docx or md file instead."):
        filegen.render(request(content=f"# T\n\n{sample}"))


def test_emoji_are_left_out_of_pdfs() -> None:
    rendered = filegen.render(request(content="# Fun 🎉\n\nGreat work 👍🏽✨ team ✔️ and ✓"))

    text = pdfium_text(rendered.data)
    assert "Great work" in text and "✓" in text and "👍" not in text
    assert rendered.notes == []


# ---------------- DOCX ----------------


def test_docx_has_styles_runs_lists_tables_links_and_properties() -> None:
    content = (
        SAMPLE
        + "\n\n3. third\n4. fourth\n\nText between.\n\n1. again one\n2. again two\n\nمرحبا بالعالم هذا نص\n\n"
        + "Bold हिन्दी: **मोटा** and *तिरछा*"
    )

    data = filegen.render(request("docx", content=content)).data

    document = docx.Document(io.BytesIO(data))
    styles = [(paragraph.style.name, paragraph.text) for paragraph in document.paragraphs]
    assert ("Heading 1", "Quarterly report") in styles and ("Heading 2", "Lists") in styles
    assert ("List Number", "First item") in styles and ("List Bullet 2", "Nested bullet") in styles
    assert ("List Bullet 3", "Deeper bullet") in styles and ("Quote", "A quoted remark.") in styles
    assert any(text.startswith("Revenue grew 12%") and "the dashboard (example.com)" in text for _, text in styles)
    assert any("[Image: chart] <img src=x onerror=alert(1)>" in text for _, text in styles)
    runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
    bold = next(run for run in runs if run.text == "12%")
    assert bold.bold and bold.font.cs_bold
    italic = next(run for run in runs if run.text == "northern")
    assert italic.italic and italic.font.cs_italic
    assert next(run for run in runs if run.text == "Old figure").font.strike
    code = next(run for run in runs if run.text == "inline code")
    assert code.font.name == "Consolas"
    assert code._r.rPr.find(qn("w:szCs")).get(qn("w:val")) == "19"
    code_block = next(paragraph for paragraph in document.paragraphs if "def total" in paragraph.text)
    assert {run.font.name for run in code_block.runs if run.text} == {"Consolas"}
    assert code_block._p.pPr.find(qn("w:shd")).get(qn("w:fill")) == "F3F4F6"
    hindi = next(run for run in runs if run.text == "मोटा")
    assert hindi.bold and hindi.font.cs_bold
    numbered = [
        paragraph._p.pPr.numPr.numId.val
        for paragraph in document.paragraphs
        if paragraph.style.name == "List Number" and paragraph._p.pPr.numPr is not None
    ]
    assert len(set(numbered)) == 3  # every ordered list restarts its numbering
    arabic = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("مرحبا"))
    assert arabic._p.pPr.find(qn("w:bidi")) is not None
    assert all(run.font.rtl for run in arabic.runs)
    [table] = document.tables
    assert table.style.name == "Table Grid"
    assert [cell.text for cell in table.rows[0].cells] == ["Region", "Revenue"]
    assert all(run.bold for cell in table.rows[0].cells for run in cell.paragraphs[0].runs)
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert table.rows[1].cells[1].paragraphs[0].alignment == 2  # right
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/_rels/document.xml.rels").decode()
    assert 'Target="https://www.example.com/q3?a=1" TargetMode="External"' in xml
    assert "javascript" not in xml
    properties = document.core_properties
    assert (properties.title, properties.author, properties.last_modified_by) == (
        "Quarterly report",
        "AskNuvexusDoc",
        "AskNuvexusDoc",
    )
    assert properties.created is not None
    assert datetime.now(timezone.utc) - properties.created.replace(tzinfo=timezone.utc) < timedelta(minutes=5)
    assert properties.comments == ""
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert not any("vbaProject" in name for name in names)
    assert b"macroEnabled" not in zipfile.ZipFile(io.BytesIO(data)).read("[Content_Types].xml")


# ---------------- XLSX ----------------


def test_xlsx_types_cells_and_never_writes_formulas() -> None:
    sheets = [
        SheetData(
            "Q3: sales/North?",
            ["Region", "Revenue", "Date", "Note"],
            [
                ["North", "12.5", "2026-10-15", True],
                ["=1+1", "-3", "2026-02-30", None],
                ["+cmd", "007", "", False],
                ["@SUM(A1)", "1,200", 12.5, "-x"],
                ["\tx", "a\x01b", 4, ""],
            ],
        ),
        SheetData("History", [], [["no header"]]),
        SheetData("x" * 40, ["a"], []),
        SheetData("x" * 40, ["a"], []),
        SheetData("שלום", ["עמודה"], [["שלום"], ["עולם"], ["x"]]),
    ]
    data = filegen.render(
        request("xlsx", content="# Notes\n\n=not a formula", sheets=sheets, sources=[SourceRef([2, 5], "a.pdf", 3)])
    ).data

    workbook = openpyxl.load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["Q3 salesNorth", "Sheet2", "x" * 31, "x" * 27 + " (2)", "שלום", "Notes", "Sources"]
    assert (workbook.properties.title, workbook.properties.creator) == ("Quarterly report", "AskNuvexusDoc")
    sheet = workbook["Q3 salesNorth"]
    header = sheet[1]
    assert [cell.value for cell in header] == ["Region", "Revenue", "Date", "Note"]
    assert all(cell.font.bold and cell.fill.start_color.rgb.endswith("F3F4F6") for cell in header)
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:D6"
    values = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
    assert values[0][:2] == ["North", 12.5] and values[0][2] == datetime(2026, 10, 15) and values[0][3] is True
    assert sheet["C2"].number_format == "yyyy-mm-dd"
    assert values[1] == ["=1+1", -3, "2026-02-30", None]
    assert values[2] == ["+cmd", "007", None, False]
    assert values[3] == ["@SUM(A1)", "1,200", 12.5, "-x"]
    assert values[4] == ["\tx", "ab", 4, None]
    for reference in ("A3", "A4", "A5", "A6", "D5"):
        assert sheet[reference].data_type == "s" and sheet[reference].quotePrefix, reference
    assert not sheet["A2"].quotePrefix
    assert all(8 <= sheet.column_dimensions[column].width <= 60 for column in "ABCD")
    assert workbook["שלום"].sheet_view.rightToLeft is True and not sheet.sheet_view.rightToLeft
    assert [[cell.value for cell in row] for row in workbook["Notes"].iter_rows()] == [
        ["# Notes"],
        [None],
        ["\\=not a formula"],
    ]
    assert workbook["Notes"]["A3"].data_type == "s"
    assert [[cell.value for cell in row] for row in workbook["Sources"].iter_rows()] == [
        ["Refs", "Document", "Page"],
        ["[2], [5]", "a.pdf", 3],
    ]
    archive = zipfile.ZipFile(io.BytesIO(data))
    for name in archive.namelist():
        if name.endswith(".xml"):
            xml = archive.read(name).decode()
            assert "<f>" not in xml and "<f " not in xml, name
    # The only defined names are the hidden ranges of the auto-filters (openpyxl writes them, as Excel does).
    names = re.findall(r'<definedName name="([^"]+)"', archive.read("xl/workbook.xml").decode())
    assert names and set(names) == {"_xlnm._FilterDatabase"}
    assert not any("externalLink" in name for name in archive.namelist())


# ---------------- Characters XML can't hold ----------------

# Vertical tab, form feed, ESC, a C0 control and the noncharacters U+FFFE and U+FFFF: lxml refuses all of them, so
# python-docx and openpyxl used to crash (the tool said "server error"). PDF and MD rendered them already.
XML_INVALID = "\x0b\x0c\x1b\x01￾￿"


def _docx_texts(data: bytes) -> list[str]:
    document = docx.Document(io.BytesIO(data))
    return [paragraph.text for paragraph in document.paragraphs]


def test_tool_arguments_lose_the_characters_xml_cant_hold() -> None:
    raw = json.dumps(
        {
            "title": "Memo￿\x0b2026￾",
            "format": "docx",
            "content": "# Memo\n\nLine one\x0bline two\x0cthree \x1b[31mred\x1b[0m\x01 end￿",
        }
    )

    args = document_tool.parse_arguments(raw)

    assert args.title == "Memo 2026"
    assert args.content == "# Memo\n\nLine one\nline two\nthree [31mred[0m end"  # \v and \f are line feeds
    data = filegen.render(RenderRequest(args.title, "docx", args.content, [], [SourceRef([1], "notes￿.pdf", 2)]))
    texts = _docx_texts(data.data)
    assert "Line one line two three [31mred[0m end" in texts  # one paragraph: Markdown soft line breaks
    assert texts[-1] == "[1] notes.pdf, page 2"
    assert docx.Document(io.BytesIO(data.data)).core_properties.title == "Memo 2026"
    assert file_name(args.title, ".docx") == "Memo 2026.docx"


def test_xlsx_arguments_lose_the_characters_xml_cant_hold() -> None:
    raw = json.dumps(
        {
            "title": "Book￿",
            "format": "xlsx",
            "content": "Notes￿\x0bmore",
            "sheets": [
                {
                    "name": "Q1\x0bdata￿",
                    "columns": ["Note￾", "Value"],
                    "rows": [["x\x0by￿", "12￾"], ["\x1b=1+1", "\x0c"]],
                }
            ],
        }
    )

    args = document_tool.parse_arguments(raw)
    data = filegen.render(
        RenderRequest(args.title, "xlsx", args.content, args.sheets, [SourceRef([1], "notes￿.pdf", None)])
    ).data

    workbook = openpyxl.load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["Q1 data", "Notes", "Sources"]
    assert workbook.properties.title == "Book"
    sheet = workbook["Q1 data"]
    assert [[cell.value for cell in row] for row in sheet.iter_rows()] == [
        ["Note", "Value"],
        ["x\ny", 12],  # a vertical tab is a line break; the number stays a number
        ["=1+1", "\n"],
    ]
    assert sheet["A3"].data_type == "s" and sheet["A3"].quotePrefix  # still never a formula
    assert [[cell.value for cell in row] for row in workbook["Notes"].iter_rows()] == [["Notes more"]]  # Markdown
    assert [[cell.value for cell in row] for row in workbook["Sources"].iter_rows()] == [
        ["Refs", "Document", "Page"],
        ["[1]", "notes.pdf", None],
    ]


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pdf", "md"])
def test_renderers_drop_the_characters_xml_cant_hold_themselves(fmt: str) -> None:
    """The renderers' own guard, for text that reaches them without the tool's argument cleaning."""
    bad = XML_INVALID
    sheets = [SheetData(f"Sheet{bad}", [f"a{bad}"], [[f"b{bad}", 1, None]])]
    rendered = filegen.render(
        request(
            fmt,
            content=f"# Title{bad}\n\nText{bad} and `code{bad}`\n\n```\nblock{bad}\n```",
            title=f"Report{bad}",
            sheets=sheets if fmt == "xlsx" else [],
            sources=[SourceRef([1], f"shared{bad}.pdf", None)],
        )
    )

    if fmt == "docx":
        text = "\n".join(_docx_texts(rendered.data))
        assert "Text" in text and "[1] shared" in text
    elif fmt == "xlsx":
        workbook = openpyxl.load_workbook(io.BytesIO(rendered.data))
        values = [cell.value for sheet in workbook for row in sheet.iter_rows() for cell in row]
        text = "\n".join(str(value) for value in values)
        assert workbook.sheetnames[0] == "Sheet" and workbook.properties.title == "Report\n\n"
        assert "b\n\n" in values
    elif fmt == "pdf":
        text = rendered.source_markdown or ""
    else:
        text = rendered.data.decode()
    assert not set(text) & {"\x0b", "\x0c", "\x1b", "\x01", "￾", "￿"}


def test_source_filenames_lose_the_characters_xml_cant_hold() -> None:
    from app.schemas import Source

    shared = Source(
        index=1, document_id=7, filename="odd￿ name.pdf", page=None, chunk_index=0, score=1.0, snippet=""
    )

    refs, unknown = document_tool.source_refs([1], [shared])

    assert (refs, unknown) == ([SourceRef([1], "odd name.pdf", None)], [])


# ---------------- MD ----------------


def test_md_is_rewritten_from_the_tree_and_loads_nothing() -> None:
    content = (
        "Intro with ![x](https://a.example/?d=1) and <img src=https://a.example/p.png> <script>alert(1)</script>\n\n"
        "[a link](<https://example.com/a b>) and [js](javascript:alert(1))\n\n"
        "```\ncode with ``` inside\n```\n"
    )

    data = filegen.render(request("md", content=content, sources=[SourceRef([2], "notes_v2*.pdf", None)])).data

    text = data.decode("utf-8")
    assert text.startswith("# Quarterly report\n\nIntro with \\[Image: x\\] and")
    assert text.endswith("\n") and not text.endswith("\n\n") and "\r" not in text
    assert "https://a.example/?d=1" not in text
    assert re.search(r"(?<!\\)<[a-zA-Z/]", text) is None  # no tag starts: every < is escaped
    assert "\\<img src=https://a.example/p.png\\>" in text
    assert "[a link](https://example.com/a%20b)" in text and "\\[js\\](javascript:alert(1))" in text
    assert "````\ncode with ``` inside\n````" in text
    assert text.endswith("## Sources\n\n- \\[2\\] notes\\_v2\\*.pdf\n")


def test_md_round_trip_is_stable() -> None:
    samples = [
        SAMPLE,
        '**a*b*c** **bold *both* bold** *it **bi*** a**"q"**b ~~s~~ `` ` `` and ` a `',
        "Text with * stars _ under [brackets] <tags> | pipes ~ tilde \\ backslash &amp; AT&T !not",
        "1. one\n2. two\n   - nested\n\n3) other\n\n- a\n\n* c\n\n5. five\n\n- ---\n- # heading in item",
        "> quote\n> > nested\n>\n> - list in quote\n\nLine one  \nline two\\\n# not heading\n\n2026. year",
        "| A | B |\n|:-:|--:|\n| 1 | **2** `a\\|b` |\n| x \\| y | |",
    ]
    for sample in samples:
        once = write(parse(sample))
        assert write(parse(once)) == once, sample


def test_sources_list_every_cited_number_of_a_document_page_once() -> None:
    sources = [
        SourceRef([2, 5], "report.pdf", 3),
        SourceRef([4], "report.pdf", None),
    ]

    text = filegen.document_markdown(request("md", content="Body [2][5][4].", sources=sources))

    assert text.endswith("## Sources\n\n- \\[2\\], \\[5\\] report.pdf, page 3\n- \\[4\\] report.pdf\n")


def test_source_refs_group_passages_by_document_and_page() -> None:
    from app.schemas import Source

    def source(index: int, doc: int, page: int | None) -> Source:
        return Source(
            index=index, document_id=doc, filename=f"doc{doc}.pdf", page=page, chunk_index=0, score=1.0, snippet=""
        )

    sources = [source(1, 1, 1), source(2, 2, 4), source(3, 1, 1), source(4, 2, 5), source(5, 2, 4)]

    refs, unknown = document_tool.source_refs([5, 2, 9, 3], sources)

    assert refs == [SourceRef([2, 5], "doc2.pdf", 4), SourceRef([3], "doc1.pdf", 1)]
    assert unknown == [9]


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("Q3 report: North/South", "Q3 report- North-South.pdf"),
        ('a<b>c"d\\e|f?g*h', "a-b-c-d-e-f-g-h.pdf"),
        ("  ..Hidden.. ", "Hidden.pdf"),
        ("con", "con_.pdf"),
        ("LPT9.notes", "LPT9_.notes.pdf"),
        ("", "document.pdf"),
        ("‮\x00\t", "document.pdf"),
        ("tabs\tand\nlines", "tabs and lines.pdf"),
        ("x" * 200, "x" * 120 + ".pdf"),
    ],
)
def test_file_names_are_safe(title: str, expected: str) -> None:
    assert file_name(title, ".pdf") == expected


def test_render_request_survives_json() -> None:
    original = request("xlsx", sheets=[SheetData("S", ["a"], [[1, "x", None, True]])], sources=[SourceRef([1], "f", 2)])

    assert RenderRequest.from_json(original.to_json()) == original


def test_mixed_document_renders_quickly() -> None:
    paragraph = "Mixed text with **bold**, *italic*, हिन्दी शब्द and a citation [3]. " * 8
    content = "# Mixed\n\n" + "\n\n".join(f"## Part {n}\n\n{paragraph}\n\n- one\n- two" for n in range(40))
    content = content[:30_000]

    started = time.perf_counter()
    rendered = filegen.render(request(content=content))

    assert time.perf_counter() - started < 10
    assert rendered.pages is not None and rendered.pages > 3


# ---------------- The child process ----------------


async def test_rendering_in_the_child_matches_rendering_in_process() -> None:
    in_process = filegen.render(request())

    through_child = await document_tool.render(request())

    assert through_child.pages == in_process.pages
    assert len(PdfReader(io.BytesIO(through_child.data)).pages) == in_process.pages
    assert through_child.source_markdown == in_process.source_markdown


async def test_render_errors_come_back_from_the_child() -> None:
    with pytest.raises(document_tool.ToolInputError, match="this content has Arabic text"):
        await document_tool.render(request(content="مرحبا بالعالم " * 20))


async def test_a_slow_child_is_killed_at_the_deadline_and_frees_its_slot(
    file_child: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "generated_file_timeout_seconds", 1)
    monkeypatch.setattr(settings, "generated_file_concurrency", 1)
    real_child = document_tool._child_command
    child: FileChild = file_child("sleep")

    started = time.monotonic()
    with pytest.raises(document_tool.ToolInputError, match="took longer than 1 seconds"):
        await document_tool.render(request())
    assert time.monotonic() - started < 5
    [pid] = child.pids
    assert not _alive(pid)

    monkeypatch.setattr(document_tool, "_child_command", real_child)
    rendered = await document_tool.render(request())  # the slot was released
    assert rendered.data.startswith(b"%PDF")


async def test_waiting_for_a_render_slot_counts_toward_the_deadline(
    file_child: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "generated_file_timeout_seconds", 1)
    monkeypatch.setattr(settings, "generated_file_concurrency", 1)
    file_child("sleep")

    results = await asyncio.gather(
        document_tool.render(request()), document_tool.render(request()), return_exceptions=True
    )

    assert all(isinstance(result, document_tool.ToolInputError) for result in results)


async def test_a_child_out_of_memory_says_the_file_is_too_large(file_child: Any) -> None:
    file_child("crash", crash="MemoryError")

    with pytest.raises(document_tool.ToolInputError, match="too large or complex to create"):
        await document_tool.render(request())


async def test_a_child_without_a_result_is_a_server_error_logged_without_its_output(
    file_child: Any, caplog: pytest.LogCaptureFixture
) -> None:
    file_child("silent")

    with caplog.at_level(logging.WARNING), pytest.raises(document_tool.ToolInputError, match="server error"):
        await document_tool.render(request())

    assert "exit code 3" in caplog.text
    assert "SECRET STDERR" not in caplog.text


def _alive(pid: int) -> bool:
    import os

    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    return True
