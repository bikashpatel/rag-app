"""Generated images in app.services.filegen: validating and re-encoding them (images.normalize, in-process and through
the child), the JPEG the renderers embed, image references in the content, and PDF, DOCX, XLSX and MD files with
figures, parsed back. Builds its own images with Pillow; needs nothing from conftest."""

import io
import json
import struct
import subprocess
import sys
import time
import xml.etree.ElementTree as ElementTree
import zipfile
import zlib
from datetime import date
from pathlib import Path
from typing import Any

import docx
import openpyxl
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageCms, ImageDraw, PngImagePlugin
from pypdf import PdfReader
from pypdf.generic import ContentStream

from app.services import filegen
from app.services.filegen import ImageRef, RenderError, RenderRequest, image_markdown, image_refs, images
from app.services.filegen.markdown import Figure, Heading, Paragraph, Span, Table, parse, texts, walk
from app.services.filegen.mdwriter import write

CREATOR = "openai/gpt-image-2"
AI_SOURCE = "http://cv.iptc.org/newscodes/digitalsourcetype/trainedAlgorithmicMedia"
NS = {
    "x": "adobe:ns:meta/",
    "rdf": "http://www.w3.org/1999/02/22-rdf-syntax-ns#",
    "Iptc4xmpExt": "http://iptc.org/std/Iptc4xmpExt/2008-02-29/",
    "xmp": "http://ns.adobe.com/xap/1.0/",
}
NOT_AN_IMAGE = "The image model returned data that isn't a PNG, JPEG or WebP image."
TEXT_REFERENCES = "Images can only be embedded in pdf and docx files; this file has text references to them."
MM = 25.4 / 72  # mm per point
EMU_PER_INCH = 914400
EXTENSIONS = {"png": ".png", "jpeg": ".jpg", "webp": ".webp"}
PILLOW_NAMES = {"png": "PNG", "jpeg": "JPEG", "webp": "WEBP"}
SECRETS = (b"SecretCam", b"secret text", b"secret comment", b"secretxmp", b"acsp")  # acsp: every ICC profile


def picture(size: tuple[int, int] = (96, 64), fmt: str = "PNG", *, alpha: bool = False, **save: Any) -> bytes:
    """A few colored shapes; with `alpha`, on a transparent background."""
    image = Image.new("RGBA" if alpha else "RGB", size, (0, 0, 0, 0) if alpha else (235, 235, 235))
    draw = ImageDraw.Draw(image)
    width, height = size
    draw.ellipse((width * 0.2, height * 0.2, width * 0.8, height * 0.8), fill=(220, 110, 30, 255))
    draw.rectangle((0, 0, width // 4, height // 4), fill=(30, 90, 200, 255))
    buffer = io.BytesIO()
    image.save(buffer, fmt, **save)
    return buffer.getvalue()


def png_header(width: int, height: int) -> bytes:
    """A PNG signature, IHDR and IEND: what Pillow reads to learn the size, with no pixel data."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IEND", b"")


def write_file(path: Path, data: bytes) -> Path:
    path.write_bytes(data)
    return path


def xmp_fields(xmp: str | bytes) -> dict[str, str]:
    root = ElementTree.fromstring(xmp)
    description = root.find("rdf:RDF/rdf:Description", NS)
    assert description is not None
    return {
        "source": description.get(f"{{{NS['Iptc4xmpExt']}}}DigitalSourceType", ""),
        "tool": description.get(f"{{{NS['xmp']}}}CreatorTool", ""),
    }


def ref(image_id: int = 1, path: str = "", filename: str = "Fox logo.png", alt: str = "A fox", **size: int) -> ImageRef:
    return ImageRef(image_id, path, filename, alt, size.get("width", 1536), size.get("height", 1024))


def request(fmt: str, content: str, refs: list[ImageRef], title: str = "Report") -> RenderRequest:
    return RenderRequest(title, fmt, content, [], [], refs)  # type: ignore[arg-type]


def child_command() -> list[str]:
    return [sys.executable, "-I", str(Path(filegen.__file__).with_name("child.py"))]


def run_child(job: dict[str, Any], cwd: Path) -> dict[str, Any]:
    """The child as document_tool runs it: `python -I`, an empty environment, the work directory as cwd."""
    done = subprocess.run(
        child_command(), input=json.dumps(job).encode(), capture_output=True, cwd=cwd, env={}, timeout=120, check=False
    )
    lines = [line for line in done.stdout.decode().splitlines() if line.strip()]
    result = json.loads(lines[-1])
    assert isinstance(result, dict)
    return result


def drawn_images(data: bytes) -> list[tuple[int, float, float, float, float]]:
    """(page index, x, top, width, height) in mm (top measured from the page's top) of every image drawn."""
    reader = PdfReader(io.BytesIO(data))
    found = []
    for number, page in enumerate(reader.pages):
        page_height = float(page.mediabox.height)
        matrix: list[float] = []
        for operands, operator in ContentStream(page.get_contents(), reader).operations:
            if operator == b"cm":
                matrix = [float(value) for value in operands]
            elif operator == b"Do" and matrix:
                width, _, _, height, x, y = matrix
                found.append((number, x * MM, (page_height - y - height) * MM, width * MM, height * MM))
    return found


def pdf_xobjects(data: bytes) -> list[dict[str, Any]]:
    reader = PdfReader(io.BytesIO(data))
    seen: dict[int, dict[str, Any]] = {}
    for page in reader.pages:
        for item in (page["/Resources"].get("/XObject") or {}).values():
            image = item.get_object()
            seen[item.idnum] = {
                "filter": image["/Filter"],
                "width": int(image["/Width"]),
                "height": int(image["/Height"]),
            }
    return list(seen.values())


def alt_texts(data: bytes) -> list[str]:
    reader = PdfReader(io.BytesIO(data))
    found: list[str] = []

    def visit(node: Any) -> None:
        node = node.get_object() if hasattr(node, "get_object") else node
        if isinstance(node, list):
            for item in node:
                visit(item)
        elif isinstance(node, dict):
            if "/Alt" in node:
                found.append(str(node["/Alt"]))
            if "/K" in node:
                visit(node["/K"])

    visit(reader.trailer["/Root"]["/StructTreeRoot"])
    return found


def pdf_text(data: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)


# ---------------- normalize ----------------


@pytest.mark.parametrize("fmt", ["png", "jpeg", "webp"])
def test_normalize_keeps_the_format_and_writes_both_renditions(fmt: str, tmp_path: Path) -> None:
    source = write_file(tmp_path / "in", picture((1536, 1024), PILLOW_NAMES[fmt]))
    out = tmp_path / "out"
    out.mkdir()

    result = images.normalize(source, out, creator=CREATOR)

    assert result == images.NormalizedImage(fmt, 1536, 1024)  # type: ignore[arg-type]
    assert {path.name for path in out.iterdir()} == {f"image{EXTENSIONS[fmt]}", "preview.webp", "thumb.webp"}
    with Image.open(out / f"image{EXTENSIONS[fmt]}") as stored:
        assert (stored.format, stored.size, stored.mode) == (PILLOW_NAMES[fmt], (1536, 1024), "RGB")
    for name, size in (("preview.webp", (1024, 683)), ("thumb.webp", (256, 171))):
        with Image.open(out / name) as rendition:
            assert (rendition.format, rendition.size) == ("WEBP", size)
            assert "xmp" not in rendition.info and "exif" not in rendition.info


def test_small_images_are_never_enlarged(tmp_path: Path) -> None:
    source = write_file(tmp_path / "in", picture((64, 48)))

    images.normalize(source, tmp_path, creator=CREATOR)

    for name in ("image.png", "preview.webp", "thumb.webp"):
        with Image.open(tmp_path / name) as image:
            assert image.size == (64, 48)


@pytest.mark.parametrize("fmt", ["png", "jpeg", "webp"])
def test_source_metadata_is_dropped_and_an_ai_marker_added(fmt: str, tmp_path: Path) -> None:
    original = Image.new("RGB", (64, 32), (200, 30, 30))
    ImageDraw.Draw(original).rectangle((0, 0, 15, 15), fill=(0, 0, 255))  # top left
    exif = original.getexif()
    exif[0x010F] = "SecretCam"  # Make
    exif[0x0112] = 6  # Orientation: rotate 90° clockwise to display
    save: dict[str, Any] = {
        "exif": exif.tobytes(),
        "icc_profile": ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes(),
        "xmp": b"<x:xmpmeta xmlns:x='adobe:ns:meta/'>secretxmp</x:xmpmeta>",
    }
    if fmt == "png":
        info = PngImagePlugin.PngInfo()
        info.add_text("Comment", "secret text")
        info.add_itxt("XML:com.adobe.xmp", "<x:xmpmeta xmlns:x='adobe:ns:meta/'>secretxmp</x:xmpmeta>")
        save["pnginfo"] = info
    if fmt == "jpeg":
        save["comment"] = b"secret comment"
    buffer = io.BytesIO()
    original.save(buffer, PILLOW_NAMES[fmt], **save)
    source = write_file(tmp_path / "in", buffer.getvalue())

    result = images.normalize(source, tmp_path, creator='Model <"x"> & co' + "\x01" + "y" * 300)

    assert (result.width, result.height) == (32, 64)  # upright, then the orientation tag is gone
    stored = (tmp_path / f"image{EXTENSIONS[fmt]}").read_bytes()
    assert not [secret for secret in SECRETS if secret in stored]
    with Image.open(io.BytesIO(stored)) as image:
        assert not {"exif", "icc_profile", "comment", "Comment"} & set(image.info)
        fields = xmp_fields(image.info["xmp"])
        assert image.getpixel((28, 2))[2] > 200  # the blue corner moved to the top right
    assert fields == {"source": AI_SOURCE, "tool": 'Model <"x"> & co' + "y" * 184}  # 200 characters
    if fmt == "png":
        assert b"iTXtXML:com.adobe.xmp" in stored
    for name in ("preview.webp", "thumb.webp"):
        rendition = (tmp_path / name).read_bytes()
        assert not [secret for secret in SECRETS if secret in rendition] and AI_SOURCE.encode() not in rendition


@pytest.mark.parametrize(("fmt", "mode"), [("png", "RGBA"), ("webp", "RGBA"), ("jpeg", "RGB")])
def test_transparency_is_kept_where_the_format_has_it(fmt: str, mode: str, tmp_path: Path) -> None:
    source = write_file(tmp_path / "in", picture((96, 64), "PNG", alpha=True))
    if fmt != "png":  # the same pixels in the other format (JPEG flattens them on white)
        with Image.open(source) as image:
            buffer = io.BytesIO()
            (image if fmt == "webp" else image.convert("RGB")).save(buffer, PILLOW_NAMES[fmt], quality=95)
        source.write_bytes(buffer.getvalue())

    images.normalize(source, tmp_path, creator=CREATOR)

    with Image.open(tmp_path / f"image{EXTENSIONS[fmt]}") as stored:
        assert stored.mode == mode
        if mode == "RGBA":
            assert stored.getpixel((95, 63))[3] == 0
    with Image.open(tmp_path / "preview.webp") as preview:
        assert preview.mode == mode


def test_a_palette_image_with_transparency_becomes_rgba(tmp_path: Path) -> None:
    image = Image.new("P", (32, 32), 0)
    image.putpalette([255, 255, 255, 200, 0, 0] + [0] * 762)
    ImageDraw.Draw(image).rectangle((8, 8, 23, 23), fill=1)
    buffer = io.BytesIO()
    image.save(buffer, "PNG", transparency=0)
    source = write_file(tmp_path / "in", buffer.getvalue())

    images.normalize(source, tmp_path, creator=CREATOR)

    with Image.open(tmp_path / "image.png") as stored:
        assert stored.mode == "RGBA"
        assert stored.getpixel((0, 0))[3] == 0 and stored.getpixel((16, 16)) == (200, 0, 0, 255)


def test_an_animated_webp_keeps_its_first_frame(tmp_path: Path) -> None:
    frames = [Image.new("RGB", (48, 48), color) for color in ((220, 0, 0), (0, 0, 220))]
    buffer = io.BytesIO()
    frames[0].save(buffer, "WEBP", save_all=True, append_images=frames[1:], duration=100, lossless=True)
    source = write_file(tmp_path / "in", buffer.getvalue())
    with Image.open(source) as animated:
        assert animated.n_frames == 2

    images.normalize(source, tmp_path, creator=CREATOR)

    with Image.open(tmp_path / "image.webp") as stored:
        assert getattr(stored, "n_frames", 1) == 1
        red, _, blue = stored.convert("RGB").getpixel((24, 24))
        assert red > 180 and blue < 40


def _gif() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (32, 32), "red").save(buffer, "GIF")
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("data", "message"),
    [
        (_gif(), NOT_AN_IMAGE),
        (b'<svg xmlns="http://www.w3.org/2000/svg" width="64" height="64"><script>x()</script></svg>', NOT_AN_IMAGE),
        (picture((32, 32), "BMP"), NOT_AN_IMAGE),
        (b"not an image at all" * 10, NOT_AN_IMAGE),
        (picture((256, 256))[:300], NOT_AN_IMAGE),  # truncated
        (png_header(20000, 20000), NOT_AN_IMAGE),  # Pillow refuses a decompression bomb from the header
        (png_header(4097, 16), "The image is 4097×16 pixels, larger than this app accepts."),
        (png_header(16, 4097), "The image is 16×4097 pixels, larger than this app accepts."),
        (picture((10, 10)), "The image is 10×10 pixels, smaller than this app accepts."),
        (picture((15, 64)), "The image is 15×64 pixels, smaller than this app accepts."),
    ],
    ids=["gif", "svg", "bmp", "garbage", "truncated", "bomb", "wide", "tall", "tiny", "narrow"],
)
def test_data_that_isnt_a_usable_image_is_refused(data: bytes, message: str, tmp_path: Path) -> None:
    source = write_file(tmp_path / "in", data)
    out = tmp_path / "out"
    out.mkdir()

    started = time.perf_counter()
    with pytest.raises(images.ImageError) as raised:
        images.normalize(source, out, creator=CREATOR)

    assert str(raised.value) == message
    assert isinstance(raised.value, RenderError)
    assert time.perf_counter() - started < 1
    assert list(out.iterdir()) == []


@pytest.mark.filterwarnings("ignore::PIL.Image.DecompressionBombWarning")
def test_too_many_pixels_are_refused_before_decoding(tmp_path: Path) -> None:
    source = write_file(tmp_path / "in", png_header(5000, 4000))  # over the cap, under Pillow's own refusal

    with pytest.raises(images.ImageError, match="5000×4000 pixels, larger than this app accepts"):
        images.normalize(source, tmp_path, creator=CREATOR)


def test_the_largest_accepted_edges_and_bytes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = write_file(tmp_path / "in", picture((4096, 16)))
    assert images.normalize(source, tmp_path, creator=CREATOR) == images.NormalizedImage("png", 4096, 16)

    monkeypatch.setattr(images, "MAX_IMAGE_BYTES", source.stat().st_size - 1)
    with pytest.raises(images.ImageError, match="larger than 32 MB"):
        images.normalize(source, tmp_path, creator=CREATOR)


# ---------------- The child ----------------


def test_the_child_normalizes_an_image(tmp_path: Path) -> None:
    source = write_file(tmp_path / "in", picture((1536, 1024), "WEBP", quality=90))
    out = tmp_path / "out"
    out.mkdir()
    job = {"kind": "image", "in": str(source), "out_dir": str(out), "creator": CREATOR, "cpu_seconds": 30}

    result = run_child({**job, "max_bytes": 64 * 1024 * 1024}, tmp_path)

    assert result == {"ok": True, "format": "webp", "width": 1536, "height": 1024}
    with Image.open(out / "image.webp") as stored:
        assert xmp_fields(stored.info["xmp"]) == {"source": AI_SOURCE, "tool": CREATOR}
    assert (out / "preview.webp").is_file() and (out / "thumb.webp").is_file()


@pytest.mark.parametrize(
    ("data", "expected"),
    [(_gif(), {"error": NOT_AN_IMAGE}), (png_header(20000, 20000), {"error": NOT_AN_IMAGE})],
    ids=["gif", "bomb"],
)
def test_the_child_reports_an_unusable_image_as_an_error(data: bytes, expected: dict[str, str], tmp_path: Path) -> None:
    source = write_file(tmp_path / "in", data)
    job = {"kind": "image", "in": str(source), "out_dir": str(tmp_path), "creator": CREATOR}

    assert run_child({**job, "cpu_seconds": 30, "max_bytes": 64 * 1024 * 1024}, tmp_path) == expected


def test_the_child_reports_a_missing_input_as_a_crash(tmp_path: Path) -> None:
    job = {"kind": "image", "in": str(tmp_path / "missing"), "out_dir": str(tmp_path), "creator": CREATOR}

    assert run_child({**job, "cpu_seconds": 30, "max_bytes": 1024}, tmp_path) == {"crash": "FileNotFoundError"}


def test_the_child_embeds_images_from_its_work_directory_only(tmp_path: Path) -> None:
    work = tmp_path / "work"
    work.mkdir()
    write_file(work / "img-1.png", picture((300, 200), alpha=True))
    write_file(tmp_path / "secret.png", picture((300, 200)))
    refs = [ref(1, "img-1.png", alt="A logo"), ref(2, "../secret.png", "Secret.png", "Secret")]
    content = "![Logo](image:1)\n\n![Secret](image:2)"
    job = {"request": request("pdf", content, refs).to_json(), "out": str(work / "out")}

    result = run_child({**job, "cpu_seconds": 60, "max_bytes": 64 * 1024 * 1024}, work)

    assert result == {
        "ok": True,
        "pages": 1,
        "notes": ["Embedded 1 image.", "1 image couldn't be read, so the file has a text reference to it instead."],
    }
    data = (work / "out").read_bytes()
    assert [image["filter"] for image in pdf_xobjects(data)] == ["/DCTDecode"]
    assert "[Image: Secret] (Secret.png)" in pdf_text(data)


# ---------------- embeddable_jpeg ----------------


def test_embeddable_jpeg_flattens_shrinks_and_drops_metadata(tmp_path: Path) -> None:
    source = write_file(tmp_path / "in.png", picture((3000, 1500), alpha=True))

    data, width, height = images.embeddable_jpeg(source)

    assert (width, height) == (2000, 1000)
    with Image.open(io.BytesIO(data)) as image:
        assert (image.format, image.mode, image.size) == ("JPEG", "RGB", (2000, 1000))
        assert min(image.getpixel((1990, 990))) >= 245  # transparent: white
        assert not {"exif", "icc_profile", "xmp", "comment"} & set(image.info)


def test_embeddable_jpeg_never_enlarges_and_honours_max_edge(tmp_path: Path) -> None:
    source = write_file(tmp_path / "in.webp", picture((120, 60), "WEBP"))

    assert images.embeddable_jpeg(source)[1:] == (120, 60)
    assert images.embeddable_jpeg(source, max_edge=40)[1:] == (40, 20)


@pytest.mark.parametrize(
    "path", ["", "..", "../img-1.png", "/etc/hosts", "sub/img-1.png", "sub\\img-1.png", ".img-1.png", "img..png"]
)
def test_embedded_image_paths_must_be_bare_names(path: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    opened: list[Path] = []
    monkeypatch.setattr(images, "embeddable_jpeg", lambda target, max_edge=2000: opened.append(target))

    with pytest.raises(images.ImageError, match="plain name"):
        images.load_embeddable(ref(1, path))

    assert opened == []


def test_a_linked_or_missing_image_file_is_refused(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "real.png", picture())
    (tmp_path / "img-1.png").symlink_to(tmp_path / "real.png")

    for name in ("img-1.png", "img-2.png"):
        with pytest.raises(images.ImageError, match="missing"):
            images.load_embeddable(ref(1, name))
    assert images.load_embeddable(ref(1, "real.png"))[1:] == (96, 64)


# ---------------- References in the content ----------------


def test_image_refs_lists_the_references_a_file_can_embed() -> None:
    content = (
        "![a](image:3) text ![b](image:1)\n\n![again](image:3)\n\n# Heading ![h](image:7)\n\n"
        "| x |\n|--|\n| ![t](image:8) |\n\n- item ![l](image:5)\n\n> ![q](<image:6>)\n\n"
        "![s](image: 9) ![w](https://example.com/x.png) ![d](image:12345678901) `![c](image:10)`\n\n"
        "```\n![f](image:11)\n```\n\n[![linked](image:2)](https://example.com)"
    )

    assert image_refs(content) == [3, 1, 5, 6, 2]


@pytest.mark.parametrize(
    "caption",
    ["Fox logo", "Logo [v2]", "a\\b *c* _d_ `e` <f> |g| ~~h~~", "AT&amp;T &#65; & ! ![x]", "ελληνικά हिन्दी"],
)
def test_image_markdown_parses_back_as_that_image_with_that_caption(caption: str) -> None:
    refs = {812: ref(812)}

    assert parse(image_markdown(caption, 812), refs) == [Figure(812, caption, refs[812])]
    assert image_refs(image_markdown(caption, 812)) == [812]


def test_image_markdown_escapes_and_keeps_one_line() -> None:
    assert image_markdown("Logo [v2]", 812) == "![Logo \\[v2\\]](image:812)"
    assert image_markdown("  two\nlines\t", 5) == "![two lines](image:5)"
    assert image_markdown("", 5) == "![](image:5)"


def test_image_source_markdown_is_the_indexed_text_of_an_image() -> None:
    text = filegen.image_source_markdown(
        title="Fox logo concept",
        alt_text="A minimalist orange fox head.\n",
        prompt="Minimalist flat vector logo of a fox head, *orange* and white.\n# Not a heading\n\nUse [brackets].",
        revised_prompt="A fox.",
        model="GPT Image 2",
        created=date(2026, 10, 16),
    )

    assert text == (
        "# Fox logo concept\n\n"
        "AI-generated image, made in chat with GPT Image 2 on 2026-10-16.\n\n"
        "A minimalist orange fox head.\n\n"
        "Prompt: Minimalist flat vector logo of a fox head, \\*orange\\* and white.\\\n\\# Not a heading\n\n"
        "Use \\[brackets\\].\n\n"
        "Revised prompt: A fox.\n"
    )
    assert [block for block in parse(text) if isinstance(block, Heading)] == [Heading(1, [Span("Fox logo concept")])]


def test_image_source_markdown_without_a_revised_prompt() -> None:
    text = filegen.image_source_markdown(
        title="Logo", alt_text="", prompt="A logo", revised_prompt=None, model="GPT Image 2", created=date(2026, 1, 2)
    )

    assert text == "# Logo\n\nAI-generated image, made in chat with GPT Image 2 on 2026-01-02.\n\nPrompt: A logo\n"


# ---------------- Parsing ----------------


def test_a_reference_splits_its_paragraph_into_text_figure_text() -> None:
    refs = {1: ref(1), 2: ref(2, filename="Poster.png")}

    blocks = parse("Before **bold** ![Logo](image:1)\nafter  \n![](<image:2>)![x](image:9)", refs)

    assert blocks == [
        Paragraph([Span("Before "), Span("bold", bold=True)]),
        Figure(1, "Logo", refs[1]),
        Paragraph([Span("after")]),
        Figure(2, "Poster", refs[2]),
        Paragraph([Span("[Image: x]")]),
    ]


def test_references_in_headings_and_cells_and_to_unknown_ids_stay_text() -> None:
    refs = {1: ref(1)}

    blocks = parse("# Title ![](image:1)\n\n| a |\n|--|\n| ![Cell](image:1) |\n\n![x](image:2) ![](image:3)", refs)

    assert blocks[0] == Heading(1, [Span("Title [Image: Fox logo]")])
    assert isinstance(blocks[1], Table) and blocks[1].rows == [[[Span("[Image: Cell]")]]]
    assert blocks[2] == Paragraph([Span("[Image: x] [Image]")])
    assert parse("![Logo](image:1)") == [Paragraph([Span("[Image: Logo]")])]  # no images given


def test_captions_are_plain_text_and_figures_nest_in_lists_and_quotes() -> None:
    refs = {1: ref(1)}

    blocks = parse("- ![**Bold** \\[v2\\] `c`](image:1)\n\n> Quoted ![q](image:1)", refs)

    figures = [block for block in walk(blocks) if isinstance(block, Figure)]
    assert [figure.caption for figure in figures] == ["Bold [v2] c", "q"]
    assert ("Bold [v2] c", False) in list(texts(blocks))


@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_a_file_embeds_at_most_ten_images(fmt: str) -> None:
    content = "\n\n".join("![x](image:1)" for _ in range(11))

    with pytest.raises(RenderError, match="at most 10 images"):
        filegen.render(request(fmt, content, [ref(1, "img-1.png")]))


def test_render_request_with_images_survives_json() -> None:
    original = request("pdf", "x", [ref(1, "img-1.png"), ref(2, "", "b.webp", "B", width=10, height=20)])

    assert RenderRequest.from_json(json.loads(json.dumps(original.to_json()))) == original
    older = {key: value for key, value in original.to_json().items() if key != "images"}
    assert RenderRequest.from_json(older).images == []


# ---------------- Text references: MD, XLSX and the PDF's indexed source ----------------


def test_md_files_get_text_references() -> None:
    content = "Intro\n\n![Fox [v2]](image:1)\n\n- ![](image:1)\n\n![x](image:9)"

    rendered = filegen.render(request("md", content, [ref(1, filename="Fox *logo*.png")]))

    text = rendered.data.decode()
    assert text == (
        "# Report\n\nIntro\n\n\\[Image: Fox \\[v2\\]\\] (Fox \\*logo\\*.png)\n\n"
        "- \\[Image: Fox \\*logo\\*\\] (Fox \\*logo\\*.png)\n\n\\[Image: x\\]\n"
    )
    assert rendered.notes == [TEXT_REFERENCES]
    assert write(parse(text)) == text
    assert filegen.render(request("md", "![x](image:9)", [ref(1)])).notes == []


def test_xlsx_notes_get_text_references(tmp_path: Path) -> None:
    rendered = filegen.render(
        RenderRequest("Sheet", "xlsx", "See ![Logo](image:1)", [filegen.SheetData("S", ["a"], [["1"]])], [], [ref(1)])
    )

    workbook = openpyxl.load_workbook(io.BytesIO(rendered.data))
    assert [row[0] for row in workbook["Notes"].iter_rows(values_only=True)] == [
        "See",
        None,
        "\\[Image: Logo\\] (Fox logo.png)",
    ]
    assert rendered.notes == [TEXT_REFERENCES]
    assert not [name for name in zipfile.ZipFile(io.BytesIO(rendered.data)).namelist() if "media" in name]


def test_the_pdf_source_markdown_has_text_references(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-1.png", picture())
    pdf_request = request("pdf", "![Logo](image:1)", [ref(1, "img-1.png")])

    rendered = filegen.render(pdf_request)

    expected = "# Report\n\n\\[Image: Logo\\] (Fox logo.png)\n"
    assert rendered.source_markdown == filegen.document_markdown(pdf_request) == expected


# ---------------- PDF ----------------


def test_a_pdf_embeds_figures_scaled_to_the_text_width_with_captions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-1.png", picture((3000, 2000), alpha=True))
    write_file(tmp_path / "img-2.webp", picture((1024, 1536), "WEBP"))
    refs = [
        ref(1, "img-1.png", "Wide.png", "An orange circle", width=3000, height=2000),
        ref(2, "img-2.webp", "Tall.webp", "A tall poster", width=1024, height=1536),
    ]

    rendered = filegen.render(request("pdf", "# Two\n\n![Wide one](image:1)\n\n![Tall one](image:2)", refs))

    assert rendered.notes == ["Embedded 2 images."]
    assert rendered.pages == 2
    assert sorted((image["filter"], image["width"], image["height"]) for image in pdf_xobjects(rendered.data)) == [
        ("/DCTDecode", 1024, 1536),
        ("/DCTDecode", 2000, 1333),
    ]
    assert alt_texts(rendered.data) == ["An orange circle", "A tall poster"]
    (page1, x1, top1, w1, h1), (page2, x2, top2, w2, h2) = drawn_images(rendered.data)
    assert (page1, round(x1, 1), round(w1, 1), round(h1, 1)) == (0, 20.0, 170.0, 113.3)  # the text width
    tallest = 0.6 * (297 - 18 - 18)
    assert (page2, round(top2, 1), round(h2, 1)) == (1, 18.0, round(tallest, 1))  # on the next page, capped
    assert round(w2, 1) == round(tallest * 1024 / 1536, 1) and round(x2 + w2 / 2, 1) == 105.0  # centered
    text = pdf_text(rendered.data)
    assert "Wide one" in text and "Tall one" in text and "image:" not in text
    assert "/MarkInfo" not in PdfReader(io.BytesIO(rendered.data)).trailer["/Root"]
    for marker in (b"/JavaScript", b"/OpenAction", b"/URI", b"/Launch", b"/EmbeddedFile", b"/MarkInfo", b"/Marked"):
        assert marker not in rendered.data


def test_pdf_captions_are_italic_grey_and_figures_in_lists_are_indented(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-1.png", picture((400, 100)))

    rendered = filegen.render(request("pdf", "- Item\n- ![Caption here](image:1)", [ref(1, "img-1.png")]))

    [(_, x, _, width, _)] = drawn_images(rendered.data)
    assert round(x, 1) == 27.0 and round(width, 1) == 163.0  # the list's 7 mm indent
    reader = PdfReader(io.BytesIO(rendered.data))
    fonts = [str(font.get_object()["/BaseFont"]) for font in reader.pages[0]["/Resources"]["/Font"].values()]
    assert any("Italic" in font for font in fonts)
    assert b"0.2941 0.3333 0.3882 rg" in zlib.decompress(reader.pages[0]["/Contents"].get_object()._data)


def test_a_pdf_keeps_a_figure_that_cant_be_read_as_a_text_reference(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-2.png", b"not a png")
    refs = [ref(1, "img-1.png"), ref(2, "img-2.png", "Broken.png")]

    rendered = filegen.render(request("pdf", "![Missing](image:1)\n\n![Broken](image:2)", refs))

    assert rendered.notes == ["2 images couldn't be read, so the file has text references to them instead."]
    assert pdf_xobjects(rendered.data) == []
    text = pdf_text(rendered.data)
    assert "[Image: Missing] (Fox logo.png)" in text and "[Image: Broken] (Broken.png)" in text


def test_the_same_image_twice_is_stored_once(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-1.png", picture((200, 100)))

    rendered = filegen.render(request("pdf", "![One](image:1)\n\n![Two](image:1)", [ref(1, "img-1.png")]))

    assert rendered.notes == ["Embedded 2 images."]
    assert len(pdf_xobjects(rendered.data)) == 1 and len(drawn_images(rendered.data)) == 2


@pytest.mark.parametrize("with_figure", [False, True])
def test_a_pdf_with_or_without_figures_never_claims_to_be_tagged(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, with_figure: bool
) -> None:
    """fpdf2 tags only the figure; a Tagged PDF claim would hide the untagged text from screen readers."""
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-1.png", picture((300, 200)))
    content = "# T\n\nBody text paragraph." + ("\n\n![Fox](image:1)" if with_figure else "")

    rendered = filegen.render(request("pdf", content, [ref(1, "img-1.png", alt="a fox")]))

    root = PdfReader(io.BytesIO(rendered.data)).trailer["/Root"]
    assert "/MarkInfo" not in root
    assert ("/StructTreeRoot" in root) is with_figure  # the figure's /Alt stays
    if with_figure:
        assert alt_texts(rendered.data) == ["a fox"]
    assert "Body text paragraph." in pdf_text(rendered.data)


FILLER = "This paragraph fills the page so that the next heading lands below the middle of it. " * 3


def page_texts(data: bytes) -> list[str]:
    return [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]


@pytest.mark.parametrize("headings", ["## Concept 2", "### Concept 2\n\n#### Hand-drawn"])
def test_a_pdf_heading_moves_to_the_next_page_with_its_figure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, headings: str
) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-1.png", picture((1536, 1024)))
    refs = [ref(1, "img-1.png")]
    start = "# Report\n\n" + "\n\n".join([FILLER] * 9) + "\n\n" + headings

    without_figure = filegen.render(request("pdf", start + "\n\nShort text.", refs))
    assert len(page_texts(without_figure.data)) == 1  # the headings alone fit on page 1

    rendered = filegen.render(request("pdf", start + "\n\n![Fox](image:1)\n\nThe end.", refs))

    first, second = page_texts(rendered.data)
    assert "Concept 2" not in first and "Concept 2" in second and "Fox" in second
    if "Hand-drawn" in headings:
        assert "Hand-drawn" in second
    [(page, _, top, _, _)] = drawn_images(rendered.data)
    assert page == 1 and top > 18.0  # below the headings
    assert rendered.notes == ["Embedded 1 image."]  # the look-ahead doesn't count the image


def test_a_pdf_heading_stays_when_its_figure_fits_or_cant_be_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-1.png", picture((1024, 1536)))
    write_file(tmp_path / "img-2.png", picture((1536, 1024)))
    refs = [ref(1, "img-1.png"), ref(2, "img-2.png"), ref(3, "img-3.png", "Missing.png")]

    at_the_top = filegen.render(request("pdf", "# Poster\n\n![Tall](image:1)", refs))
    assert at_the_top.pages == 1  # no blank first page

    fits = filegen.render(request("pdf", "# R\n\n" + "\n\n".join([FILLER] * 3) + "\n\n## C\n\n![W](image:2)", refs))
    assert fits.pages == 1 and [page for page, *_ in drawn_images(fits.data)] == [0]

    unreadable = "# R\n\n" + "\n\n".join([FILLER] * 9) + "\n\n## Concept 3\n\n![Gone](image:3)"
    rendered = filegen.render(request("pdf", unreadable, refs))
    assert rendered.pages == 1 and "Concept 3" in page_texts(rendered.data)[0]


def test_an_embedded_images_file_name_adds_no_placeholder_note(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-5.png", picture((300, 200)))
    content = "# Trip report\n\nOur trip.\n\n![Mount Fuji at dawn](image:5)"

    rendered = filegen.render(request("pdf", content, [ref(5, "img-5.png", "富士山 ★ sunrise.png")]))

    assert rendered.notes == ["Embedded 1 image."]


def test_a_text_reference_counts_the_characters_it_draws_once(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)  # img-5.png doesn't exist: the text reference is drawn
    content = f"# Trip report\n\n{FILLER}\n\n![Mount Fuji 山](image:5)"

    rendered = filegen.render(request("pdf", content, [ref(5, "img-5.png", "Fuji 山.png")]))

    assert "[Image: Mount Fuji □] (Fuji □.png)" in pdf_text(rendered.data)
    assert rendered.notes == [
        "1 image couldn't be read, so the file has a text reference to it instead.",
        "2 characters in Han couldn't be shown in the PDF and appear as □. A docx or md file keeps them.",
    ]


# ---------------- DOCX ----------------


def test_a_docx_embeds_centered_pictures_with_alt_text_and_captions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-1.png", picture((1536, 1024), alpha=True))
    write_file(tmp_path / "img-2.jpg", picture((1024, 1536), "JPEG"))
    refs = [ref(1, "img-1.png", alt="An orange circle"), ref(2, "img-2.jpg", "Tall.jpg", "A tall poster")]
    content = "Intro\n\n![Wide one](image:1)\n\n- Item\n\n  ![](image:2)\n\n| a |\n|--|\n| ![Cell](image:1) |"

    rendered = filegen.render(request("docx", content, refs))

    assert rendered.notes == ["Embedded 2 images."]
    document = docx.Document(io.BytesIO(rendered.data))
    shapes = document.inline_shapes
    assert [shape._inline.docPr.get("descr") for shape in shapes] == ["An orange circle", "A tall poster"]
    sizes = [(round(shape.width / EMU_PER_INCH, 2), round(shape.height / EMU_PER_INCH, 2)) for shape in shapes]
    assert sizes == [(6.0, 4.0), (3.6, 5.4)]  # the text width; 60% of the 9 in usable height
    holders = [paragraph for paragraph in document.paragraphs if paragraph._p.xpath(".//pic:pic")]
    assert all(holder.alignment == WD_ALIGN_PARAGRAPH.CENTER for holder in holders)
    assert all(holder.paragraph_format.keep_with_next for holder in holders)
    assert holders[1].paragraph_format.left_indent == round(0.25 * EMU_PER_INCH)  # in the list
    captions = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Caption"]
    assert [caption.text for caption in captions] == ["Wide one", "Tall"]
    assert all(caption.alignment == WD_ALIGN_PARAGRAPH.CENTER for caption in captions)
    assert document.tables[0].cell(1, 0).text == "[Image: Cell]"
    media = [name for name in zipfile.ZipFile(io.BytesIO(rendered.data)).namelist() if name.startswith("word/media/")]
    assert len(media) == 2 and all(name.endswith((".jpg", ".jpeg")) for name in media)


def test_a_docx_keeps_a_figure_that_cant_be_read_as_a_text_reference() -> None:
    rendered = filegen.render(request("docx", "![Missing](image:1)", [ref(1, "../img-1.png")]))

    document = docx.Document(io.BytesIO(rendered.data))
    assert len(document.inline_shapes) == 0
    assert [paragraph.text for paragraph in document.paragraphs][-1] == "[Image: Missing] (Fox logo.png)"
    assert rendered.notes == ["1 image couldn't be read, so the file has a text reference to it instead."]


def test_a_right_to_left_caption_is_marked_right_to_left(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    write_file(tmp_path / "img-1.png", picture())

    rendered = filegen.render(request("docx", "![شعار الثعلب](image:1)", [ref(1, "img-1.png")]))

    [caption] = [p for p in docx.Document(io.BytesIO(rendered.data)).paragraphs if p.style.name == "Caption"]
    assert caption.text == "شعار الثعلب" and caption._p.pPr.xpath("./w:bidi")


def test_a_pdf_text_reference_in_another_script_is_drawn_with_its_font() -> None:
    refs = [ref(1, "img-1.png", "लोगो चित्र.png")]  # the file is missing: the reference is drawn instead

    rendered = filegen.render(request("pdf", "A Latin document.\n\n![Logo](image:1)", refs))

    assert b"NotoSansDevanagari" in rendered.data
    assert rendered.notes == ["1 image couldn't be read, so the file has a text reference to it instead."]
