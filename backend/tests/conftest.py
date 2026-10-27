"""Backend test harness: isolated app state per test, a mocked LiteLLM, and shared helpers.

Every test gets a fresh SQLite database, upload directory, FakeRedis server and in-memory Qdrant
client, and LiteLLM is always replaced by `FakeLLM`, so nothing touches the network. Test modules
import the plain helpers from here, e.g. `from conftest import chat, ingest, pdf_bytes, upload`.
"""

import asyncio
import base64
import copy
import csv
import functools
import hashlib
import io
import json
import math
import os
import re
import shutil
import sys
import tempfile
import zlib
from collections.abc import AsyncIterator, Awaitable, Callable, Iterator, Mapping, Sequence
from dataclasses import dataclass, field, replace
from datetime import datetime, timedelta
from html import escape
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Literal

# app.config reads the environment when it is first imported, so configure it before any app import.
_ROOT = Path(tempfile.mkdtemp(prefix="rag-tests-"))
ADMIN_USERNAME = "admin"
ADMIN_PASSWORD = "admin-password-123"
USER_PASSWORD = "user-password-123"
os.environ.update(
    {
        "SECRET_KEY": "test-secret-key-with-at-least-32-characters",
        "ADMIN_USERNAME": ADMIN_USERNAME,
        "ADMIN_PASSWORD": ADMIN_PASSWORD,
        # Each test swaps in its own database (see isolated_state); this one is never opened.
        "DATABASE_URL": f"sqlite+aiosqlite:///{_ROOT / 'unused.db'}",
        # Unreachable on purpose: every test installs a FakeRedis.
        "REDIS_URL": "redis://127.0.0.1:1/0",
        "QDRANT_URL": ":memory:",
        "UPLOAD_DIR": str(_ROOT / "uploads"),
        "COOKIE_SECURE": "false",
        # Non-empty provider keys make bootstrap enable the seeded models. LiteLLM is mocked anyway.
        "OPENAI_API_KEY": "test-openai-key",
        "ANTHROPIC_API_KEY": "test-anthropic-key",
        # Use LiteLLM's bundled cost map instead of downloading it at import.
        "LITELLM_LOCAL_MODEL_COST_MAP": "True",
        # In its default DEV mode, importing LiteLLM loads the nearest .env (the repo's own, found from the venv),
        # which would hand the owner's real settings (OAuth client, token key, limits) to the tests.
        "LITELLM_MODE": "PRODUCTION",
    }
)

import httpx  # noqa: E402
import litellm  # noqa: E402
import pytest  # noqa: E402
from cryptography.fernet import Fernet  # noqa: E402
from docx import Document as new_docx  # noqa: E402
from fake_google import CLIENT_ID, CLIENT_SECRET, FakeGoogle, fake_google_app  # noqa: E402
from fake_razorpay import API_HOST, FakeRazorpay, fake_razorpay_app  # noqa: E402
from fakeredis import FakeServer  # noqa: E402
from fakeredis.aioredis import FakeRedis  # noqa: E402
from fastapi import FastAPI  # noqa: E402
from litellm.types.utils import ImageUsage, ImageUsageInputTokensDetails  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402
from sqlalchemy import make_url, select  # noqa: E402
from sqlalchemy import text as sql  # noqa: E402
from sqlalchemy.ext.asyncio import AsyncEngine  # noqa: E402

from app import db  # noqa: E402
from app.config import settings  # noqa: E402
from app.main import app as fastapi_app  # noqa: E402
from app.models import Document, Plan, PlanPeriod, User, utcnow  # noqa: E402
from app.plugins import connections, intent, registry  # noqa: E402
from app.plugins import http as plugin_http  # noqa: E402
from app.redis_client import INGEST_QUEUE, close_redis, get_redis, set_redis  # noqa: E402
from app.services import document_tool, image_tool, ingestion, llm, pdf_images, vision  # noqa: E402
from app.services.billing import entitlements, periods  # noqa: E402
from app.services.ingestion import process_document  # noqa: E402
from app.services.vectorstore import close_qdrant  # noqa: E402

SseEvent = tuple[str, dict[str, Any]]


def pytest_sessionfinish(session: pytest.Session, exitstatus: int) -> None:
    shutil.rmtree(_ROOT, ignore_errors=True)


# ---------------- Isolated state ----------------

# The database of a test is a new SQLite file, or the TEST_POSTGRES_URL database for tests marked `postgres` (skipped
# without it) and, with TEST_DATABASE=postgres, for every test. The Postgres database's public schema is dropped and
# re-created before each test, so give each pytest process its own database. Two guards keep that away from real data:
# the database's name must say it is a throwaway one, and a schema holding tables is emptied only when these tests
# marked it as theirs.
POSTGRES_URL = os.environ.get("TEST_POSTGRES_URL", "")
ON_POSTGRES = os.environ.get("TEST_DATABASE") == "postgres"
if ON_POSTGRES and not POSTGRES_URL:
    raise RuntimeError("TEST_DATABASE=postgres needs TEST_POSTGRES_URL")
if POSTGRES_URL and not re.search(r"test|scratch", make_url(POSTGRES_URL).database or "", re.IGNORECASE):
    raise RuntimeError(
        "TEST_POSTGRES_URL must name a throwaway database (its name containing 'test' or 'scratch'): every test drops "
        "its public schema"
    )
_SCRATCH_MARK = "rag-app tests: emptied before every test"
sqlite_only = pytest.mark.skipif(ON_POSTGRES, reason="reads the SQLite test database directly")


def pytest_collection_modifyitems(items: list[pytest.Item]) -> None:
    if not POSTGRES_URL:
        skip = pytest.mark.skip(reason="TEST_POSTGRES_URL isn't set")
        for item in items:
            if item.get_closest_marker("postgres"):
                item.add_marker(skip)


async def _empty_postgres_engine() -> AsyncEngine:
    engine = db.make_engine(POSTGRES_URL)
    async with engine.begin() as connection:
        tables = await connection.scalar(sql("SELECT count(*) FROM pg_tables WHERE schemaname = 'public'"))
        mark = await connection.scalar(sql("SELECT obj_description('public'::regnamespace, 'pg_namespace')"))
        if tables and mark != _SCRATCH_MARK:
            pytest.exit(
                f"Refusing to empty the TEST_POSTGRES_URL database: its public schema holds {tables} tables these "
                "tests didn't create. Point TEST_POSTGRES_URL at a new, empty database.",
                returncode=4,
            )
        await connection.execute(sql("DROP SCHEMA public CASCADE"))
        await connection.execute(sql("CREATE SCHEMA public"))
        await connection.execute(sql(f"COMMENT ON SCHEMA public IS '{_SCRATCH_MARK}'"))
    return engine


@pytest.fixture(autouse=True)
async def isolated_state(monkeypatch: pytest.MonkeyPatch, request: pytest.FixtureRequest) -> AsyncIterator[Path]:
    """Fresh database, upload directory, Redis and Qdrant for every test; yields the test's scratch dir.

    Tables exist only once the `app` fixture has run its lifespan. Override settings per test with
    `monkeypatch.setattr(settings, "<field>", value)` (every module shares that one object).
    """
    workdir = Path(tempfile.mkdtemp(dir=_ROOT))
    if ON_POSTGRES or request.node.get_closest_marker("postgres"):
        engine = await _empty_postgres_engine()
    else:
        engine = db.make_engine(f"sqlite+aiosqlite:///{workdir / 'app.db'}")
    monkeypatch.setattr(db, "engine", engine)
    db.SessionLocal.configure(bind=engine)
    entitlements.invalidate_unlock()  # the plans cached from another test's database
    monkeypatch.setattr(settings, "upload_dir", str(workdir / "uploads"))
    set_redis(FakeRedis(server=FakeServer(), decode_responses=True))
    await close_qdrant()
    yield workdir
    await close_redis()
    await close_qdrant()
    await engine.dispose()
    shutil.rmtree(workdir, ignore_errors=True)


# ---------------- LiteLLM mock ----------------


def fake_embedding(text: str, dim: int = 64) -> list[float]:
    """Deterministic unit vector (signed hashed bag of words): texts that share words score higher."""
    vector = [0.0] * dim
    for word in re.findall(r"\w+", text.lower()):
        digest = hashlib.sha256(word.encode()).digest()
        vector[int.from_bytes(digest[:4], "big") % dim] += 1.0 if digest[4] % 2 else -1.0
    if not any(vector):
        vector[0] = 1.0
    norm = math.sqrt(sum(x * x for x in vector))
    return [x / norm for x in vector]


def default_vision_reply(page: int, mode: str) -> str:
    """What FakeLLM's vision model answers by default, per page and mode ("transcribe", "describe" or "test")."""
    if mode == "transcribe":
        return f"Scanned text of page {page}: invoice total 1234."
    if mode == "describe":
        return f"[Chart] Revenue by quarter on page {page}, rising every quarter."
    return "TEST 4712"


@dataclass
class ScriptedCall:
    """One scripted streaming chat call, in the chunk shapes LiteLLM 1.103 produces (see FakeLLM.script).

    `blocks` in stream order: ("thinking", (text deltas, signature)) is a signed thinking block (one chunk per delta,
    then one with the whole signed block), ("redacted", data) a redacted one, ("text", str) one text chunk, and
    ("tool", (name, arguments as JSON text or a dict)) a tool call whose arguments arrive `fragment` characters per
    chunk. `finish_reason` defaults to "tool_calls" with tool blocks, else "stop"; "" sends no finish chunk. `usage`
    is reported after the finish (None: not reported). `error` is raised after the chunks, or by the call itself
    (before any output) with `before_output`.
    """

    blocks: list[tuple[str, Any]] = field(default_factory=list)
    shape: Literal["anthropic", "openai"] = "anthropic"
    finish_reason: str | None = None
    usage: tuple[int, int] | None = (100, 20)
    fragment: int = 17
    error: BaseException | None = None
    before_output: bool = False


def tool_call(name: str = "create_document", **arguments: Any) -> tuple[str, tuple[str, Any]]:
    """A ("tool", ...) block for ScriptedCall.blocks."""
    return ("tool", (name, arguments))


def _tool_item(index: int, call_id: str | None, name: str | None, arguments: str) -> Any:
    function = SimpleNamespace(name=name, arguments=arguments)
    return SimpleNamespace(index=index, id=call_id, type="function" if call_id else None, function=function)


def _scripted_chunks(call: ScriptedCall) -> list[Any]:
    chunks: list[Any] = []
    tools = 0
    if call.shape == "openai":
        chunks.append(_stream_chunk(content=""))
    for kind, value in call.blocks:
        if kind == "thinking":
            deltas, signature = value
            for text in deltas:
                block = {"type": "thinking", "thinking": text}
                chunks.append(_stream_chunk(content=None, reasoning_content=text, thinking_blocks=[block]))
            whole = {"type": "thinking", "thinking": "".join(deltas), "signature": signature}
            chunks.append(_stream_chunk(content=None, thinking_blocks=[whole]))
        elif kind == "redacted":
            chunks.append(_stream_chunk(content=None, thinking_blocks=[{"type": "redacted_thinking", "data": value}]))
        elif kind == "text":
            chunks.append(_stream_chunk(content=value))
        elif kind == "tool":
            name, arguments = value
            raw = arguments if isinstance(arguments, str) else json.dumps(arguments)
            index = tools
            tools += 1
            if call.shape == "openai":
                whole = raw if index else ""  # later calls arrive whole in their start chunk
                chunks.append(_stream_chunk(content=None, tool_calls=[_tool_item(index, f"call_{index}", name, whole)]))
                if index:
                    continue
            else:
                chunks.append(_stream_chunk(content="", tool_calls=[_tool_item(index, f"toolu_{index}", name, "")]))
            for start in range(0, len(raw), max(1, call.fragment)):
                fragment = raw[start : start + call.fragment]
                chunks.append(_stream_chunk(content=None, tool_calls=[_tool_item(index, None, None, fragment)]))
        else:
            raise AssertionError(f"Unknown scripted block {kind!r}")
    finish = call.finish_reason
    if finish is None:
        finish = "tool_calls" if tools else "stop"
    if finish:
        chunks.append(_stream_chunk(content=None, finish_reason=finish))
    if call.usage is not None:
        prompt_tokens, completion_tokens = call.usage
        total = prompt_tokens + completion_tokens
        usage = SimpleNamespace(prompt_tokens=prompt_tokens, completion_tokens=completion_tokens, total_tokens=total)
        chunks.append(SimpleNamespace(choices=[], usage=usage))
    return chunks


def fake_image(width: int = 64, height: int = 64, fmt: str = "png") -> bytes:
    """A real image of `width` x `height` drawn by Pillow (a few colored shapes, so re-encoding and renditions do real
    work), as PNG, JPEG or WebP bytes."""
    image = Image.new("RGB", (width, height), (245, 240, 230))
    draw = ImageDraw.Draw(image)
    draw.rectangle((width // 8, height // 8, width // 2, height // 2), fill=(220, 90, 30))
    draw.ellipse((width // 3, height // 3, width - 4, height - 4), fill=(40, 90, 200))
    draw.line((0, height - 1, width - 1, 0), fill=(20, 20, 20), width=2)
    buffer = io.BytesIO()
    image.save(buffer, {"png": "PNG", "jpeg": "JPEG", "webp": "WEBP"}[fmt])
    return buffer.getvalue()


def refusal(code: str = "moderation_blocked") -> Exception:
    """A provider's safety refusal as LiteLLM raises it (test_llm_images.py pins these shapes): GPT image models'
    `moderation_blocked` is a plain BadRequestError with the error object as `body`; `content_policy_violation` is
    LiteLLM's ContentPolicyViolationError."""
    if code == "content_policy_violation":
        return litellm.ContentPolicyViolationError(
            message="OpenAIException - Your request was rejected as a result of our safety system.",
            model="gpt-image-2",
            llm_provider="openai",
            body={"code": "content_policy_violation", "type": "invalid_request_error"},
        )
    return litellm.BadRequestError(
        message="OpenAIException - Your request was rejected by the safety system. If you believe this is an error, "
        "contact us at help.openai.com and include the request ID req_123.",
        model="gpt-image-2",
        llm_provider="openai",
        body={"code": code, "type": "image_generation_user_error", "message": "Your request was rejected"},
    )


ALL_ACTIONS = '{"file": true, "image": true, "draft": true, "send": true, "import": true}'


def intent_answer(*actions: str) -> str:
    """A request-check reply asking for exactly `actions` (e.g. "draft")."""
    return json.dumps({action: action in actions for action in intent.ACTIONS})


@dataclass
class FakeLLM:
    """Stand-in for `litellm.acompletion` and `litellm.aembedding`; change its attributes inside a test.

    Streaming chat: one delta per `reasoning` item (thinking: `delta.content` is None and
    `delta.reasoning_content` holds the text), then one delta per `chunks` item (sleeping `chunk_delay`
    seconds before each), then a usage-only chunk reporting `usage` = (prompt_tokens, completion_tokens);
    with `usage=None` nothing is reported and the app estimates. `fail_after=n` raises `error` after n
    content chunks; `fail_after=0` makes the call itself raise. `stall_after=n` makes the stream wait forever
    before its n-th content chunk. While `script` holds ScriptedCall items, each streaming call takes the first one
    and streams it instead (its chunks also wait `chunk_delay`). Non-streaming calls (the admin model
    test) answer `reply`. Embeddings return `fake_embedding(text, embedding_dim)` per input with
    `prompt_tokens` = word count (no usage when `embedding_usage` is False); `embedding_error` makes every
    embedding call raise it. Keyword arguments of every call are kept in `chat_calls` / `embedding_calls`
    (e.g. `chat_calls[-1]["messages"]` is the prompt the model got).

    Vision calls (non-streaming, with an image part) go to `vision_calls` instead. The page and mode come from
    the instruction ("test" and page 0 for the admin Test). The reply is `vision_reply`: a string, or a function
    of (page, mode), which may raise to make that call fail (e.g. `litellm.RateLimitError`). The response
    reports `vision_usage` (None: no usage) and `vision_finish_reason`, after `vision_delay` seconds.
    `vision_max_in_flight` is the peak number of vision calls running at once.

    Image calls (`litellm.aimage_generation`) go to `image_calls`. Each waits on `image_gate` (an asyncio.Event) or
    `image_delay` seconds, raises `image_error` when set (an exception, or a function of the call number returning
    one or None), and otherwise returns an ImageResponse with a real image drawn by Pillow in the requested
    output_format, at 1/16 of the requested size (64x64, 96x64 or 64x96), unless `image_reply` gives the bytes.
    `image_url_only` returns only a URL; `image_usage` = (input_tokens, output_tokens) or None (not reported);
    `image_revised_prompt` is the item's revised_prompt.

    Request checks (app.plugins.intent: non-streaming calls whose system prompt is intent.SYSTEM_PROMPT) go to
    `intent_calls` instead of `chat_calls`. They answer `intent_reply` (a string, or a function of the user message
    holding the tagged blocks), default every action true, with `intent_finish_reason` and `intent_usage`; with
    `intent_error` set they raise it.
    """

    chunks: list[str] = field(default_factory=lambda: ["Hello", " from", " the", " fake", " model."])
    usage: tuple[int, int] | None = (42, 5)
    reasoning: list[str] = field(default_factory=list)
    chunk_delay: float = 0.0
    fail_after: int | None = None
    stall_after: int | None = None
    script: list[ScriptedCall] = field(default_factory=list)
    error: BaseException = field(default_factory=lambda: RuntimeError("The provider stream broke"))
    reply: str = "OK"
    embedding_dim: int = 64
    embedding_usage: bool = True
    embedding_error: BaseException | None = None
    chat_calls: list[dict[str, Any]] = field(default_factory=list)
    embedding_calls: list[dict[str, Any]] = field(default_factory=list)
    vision_reply: str | Callable[[int, str], str] = default_vision_reply
    vision_usage: tuple[int, int] | None = (1500, 200)
    vision_finish_reason: str = "stop"
    vision_delay: float = 0.0
    vision_calls: list[dict[str, Any]] = field(default_factory=list)
    vision_max_in_flight: int = 0
    _vision_in_flight: int = 0
    image_calls: list[dict[str, Any]] = field(default_factory=list)
    image_delay: float = 0.0
    image_gate: asyncio.Event | None = None
    image_error: BaseException | Callable[[int], BaseException | None] | None = None
    image_reply: bytes | None = None
    image_url_only: bool = False
    image_usage: tuple[int, int] | None = (50, 272)
    image_revised_prompt: str | None = None
    intent_reply: str | Callable[[str], str] = ALL_ACTIONS
    intent_finish_reason: str = "stop"
    intent_usage: tuple[int, int] | None = (300, 20)
    intent_error: BaseException | None = None
    intent_calls: list[dict[str, Any]] = field(default_factory=list)

    async def acompletion(self, **kwargs: Any) -> Any:
        if not kwargs.get("stream") and vision_part(kwargs, "image_url") is not None:
            return await self._vision(kwargs)
        messages = kwargs.get("messages") or []
        if not kwargs.get("stream") and messages and messages[0].get("content") == intent.SYSTEM_PROMPT:
            return self._intent(kwargs)
        # A snapshot: the app keeps extending (and may strip) the same message list across the calls of an answer.
        self.chat_calls.append({**kwargs, "messages": copy.deepcopy(kwargs.get("messages"))})
        if kwargs.get("stream") and self.script:
            scripted = self.script.pop(0)
            if scripted.error is not None and scripted.before_output:
                raise scripted.error
            return self._scripted(scripted)
        if self.fail_after == 0:
            raise self.error
        if not kwargs.get("stream"):
            return SimpleNamespace(
                choices=[SimpleNamespace(message=SimpleNamespace(content=self.reply), finish_reason="stop")],
                usage=SimpleNamespace(prompt_tokens=5, completion_tokens=1, total_tokens=6),
            )
        return self._stream()

    def _intent(self, kwargs: dict[str, Any]) -> Any:
        self.intent_calls.append({**kwargs, "messages": copy.deepcopy(kwargs.get("messages"))})
        if self.intent_error is not None:
            raise self.intent_error
        latest = kwargs["messages"][-1]["content"]
        reply = self.intent_reply(latest) if callable(self.intent_reply) else self.intent_reply
        usage = None
        if self.intent_usage is not None:
            prompt_tokens, completion_tokens = self.intent_usage
            usage = SimpleNamespace(
                prompt_tokens=prompt_tokens,
                completion_tokens=completion_tokens,
                total_tokens=prompt_tokens + completion_tokens,
            )
        return SimpleNamespace(
            choices=[
                SimpleNamespace(message=SimpleNamespace(content=reply), finish_reason=self.intent_finish_reason)
            ],
            usage=usage,
        )

    async def _vision(self, kwargs: dict[str, Any]) -> Any:
        self.vision_calls.append(kwargs)
        page, mode = vision_page_and_mode(kwargs)
        self._vision_in_flight += 1
        self.vision_max_in_flight = max(self.vision_max_in_flight, self._vision_in_flight)
        try:
            if self.vision_delay:
                await asyncio.sleep(self.vision_delay)
            reply = self.vision_reply(page, mode) if callable(self.vision_reply) else self.vision_reply
        finally:
            self._vision_in_flight -= 1
        usage = None
        if self.vision_usage is not None:
            prompt_tokens, completion_tokens = self.vision_usage
            usage = SimpleNamespace(
                prompt_tokens=prompt_tokens,
                completion_tokens=completion_tokens,
                total_tokens=prompt_tokens + completion_tokens,
            )
        return SimpleNamespace(
            choices=[
                SimpleNamespace(message=SimpleNamespace(content=reply), finish_reason=self.vision_finish_reason)
            ],
            usage=usage,
        )

    async def _stream(self) -> AsyncIterator[Any]:
        for text in self.reasoning:
            yield _stream_chunk(content=None, reasoning_content=text)
        for sent, text in enumerate(self.chunks):
            if self.fail_after is not None and sent >= self.fail_after:
                raise self.error
            if self.stall_after is not None and sent >= self.stall_after:
                await asyncio.Event().wait()
            if self.chunk_delay:
                await asyncio.sleep(self.chunk_delay)
            yield _stream_chunk(content=text)
        if self.fail_after is not None and self.fail_after >= len(self.chunks):
            raise self.error
        if self.usage is not None:
            prompt_tokens, completion_tokens = self.usage
            yield SimpleNamespace(
                choices=[],
                usage=SimpleNamespace(
                    prompt_tokens=prompt_tokens,
                    completion_tokens=completion_tokens,
                    total_tokens=prompt_tokens + completion_tokens,
                ),
            )

    async def _scripted(self, call: ScriptedCall) -> AsyncIterator[Any]:
        for chunk in _scripted_chunks(call):
            if self.chunk_delay:
                await asyncio.sleep(self.chunk_delay)
            yield chunk
        if call.error is not None:
            raise call.error

    async def aimage_generation(self, **kwargs: Any) -> Any:
        self.image_calls.append(kwargs)
        number = len(self.image_calls)
        if self.image_gate is not None:
            await self.image_gate.wait()
        elif self.image_delay:
            await asyncio.sleep(self.image_delay)
        error = self.image_error(number) if callable(self.image_error) else self.image_error
        if error is not None:
            raise error
        size = kwargs.get("size") or "1024x1024"
        width, height = (int(value) for value in size.split("x"))
        output_format = kwargs.get("output_format") or "png"
        data = self.image_reply
        if data is None:
            data = fake_image(width // 16, height // 16, output_format)
        if self.image_url_only:
            item: dict[str, Any] = {"url": "https://images.example.invalid/generated.png"}
        else:
            item = {"b64_json": base64.b64encode(data).decode(), "revised_prompt": self.image_revised_prompt}
        usage = None
        if self.image_usage is not None:
            input_tokens, output_tokens = self.image_usage
            usage = ImageUsage(
                input_tokens=input_tokens,
                input_tokens_details=ImageUsageInputTokensDetails(image_tokens=0, text_tokens=input_tokens),
                output_tokens=output_tokens,
                total_tokens=input_tokens + output_tokens,
            )
        return litellm.ImageResponse(
            created=1_760_000_000,
            data=[item],
            usage=usage,
            size=size,
            quality=kwargs.get("quality"),
            output_format=output_format,
        )

    async def aembedding(self, **kwargs: Any) -> Any:
        self.embedding_calls.append(kwargs)
        if self.embedding_error is not None:
            raise self.embedding_error
        texts = kwargs["input"]
        tokens = sum(len(text.split()) for text in texts)
        return SimpleNamespace(
            data=[
                {"object": "embedding", "index": index, "embedding": fake_embedding(text, self.embedding_dim)}
                for index, text in enumerate(texts)
            ],
            usage=SimpleNamespace(prompt_tokens=tokens, total_tokens=tokens) if self.embedding_usage else None,
        )


def vision_part(call: Mapping[str, Any], part_type: str) -> Any:
    """The first `part_type` ("image_url" or "text") part of a call's user message, or None."""
    for message in call.get("messages") or []:
        content = message.get("content")
        if message.get("role") == "user" and isinstance(content, list):
            for part in content:
                if isinstance(part, dict) and part.get("type") == part_type:
                    return part
    return None


def vision_page_and_mode(call: Mapping[str, Any]) -> tuple[int, str]:
    """(page, mode) of a vision call from its instruction: mode "transcribe", "describe" or "test" (page 0)."""
    text = (vision_part(call, "text") or {}).get("text", "")
    if text == vision.TEST_INSTRUCTION:
        return 0, "test"
    found = re.search(r"\bPage (\d+)\b", text)
    page = int(found.group(1)) if found else 0
    for mode in ("transcribe", "describe"):
        if text == vision.instruction(page, mode):  # type: ignore[arg-type]
            return page, mode
    raise AssertionError(f"Unexpected vision instruction: {text[:100]!r}")


def vision_pages(fake: "FakeLLM", mode: str | None = None) -> list[int]:
    """The pages of FakeLLM's vision calls so far, in call order (only `mode` calls if given)."""
    calls = (vision_page_and_mode(call) for call in fake.vision_calls)
    return [page for page, call_mode in calls if mode is None or call_mode == mode]


def _stream_chunk(
    *,
    content: str | None,
    reasoning_content: str | None = None,
    thinking_blocks: list[dict[str, Any]] | None = None,
    tool_calls: list[Any] | None = None,
    finish_reason: str | None = None,
) -> Any:
    delta = SimpleNamespace(
        role="assistant",
        content=content,
        reasoning_content=reasoning_content,
        thinking_blocks=thinking_blocks,
        tool_calls=tool_calls,
    )
    return SimpleNamespace(choices=[SimpleNamespace(index=0, delta=delta, finish_reason=finish_reason)], usage=None)


@pytest.fixture(autouse=True)
def fake_llm(monkeypatch: pytest.MonkeyPatch) -> FakeLLM:
    """The FakeLLM installed for this test (always active; request it to configure or inspect it). Also forgets
    which models rejected tools and LiteLLM's function-calling answers, which the app keeps for its lifetime, and
    which image models were marked unusable.

    CHAT_IMAGE_TOOLS is off by default, so the tests written before images keep their file-only prompts and tool
    lists; request `image_tools` to turn it on."""
    fake = FakeLLM()
    monkeypatch.setattr(litellm, "acompletion", fake.acompletion)
    monkeypatch.setattr(litellm, "aembedding", fake.aembedding)
    monkeypatch.setattr(litellm, "aimage_generation", fake.aimage_generation)
    monkeypatch.setattr(llm, "_tools_rejected", set())
    monkeypatch.setattr(llm, "_tool_support", {})
    monkeypatch.setattr(llm, "_unpriced_image_models", set())
    monkeypatch.setattr(image_tool, "_unusable", {})
    monkeypatch.setattr(image_tool, "_warned_broken", set())
    monkeypatch.setattr(settings, "chat_image_tools", False)
    monkeypatch.setattr(intent, "_length_warned", set())
    return fake


@pytest.fixture(autouse=True)
def plugin_registry(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    """An empty plugin registry, so the tests written before plugins keep their prompts, tools and queries; plugin
    tests install theirs with `registry.use_plugins([...])` (or the `plugins` fixture's helpers). Also forgets the
    connection state the app keeps for its lifetime."""
    monkeypatch.setattr(connections, "_client_rejected", set())
    monkeypatch.setattr(connections, "_unreadable_logged", set())
    with registry.use_plugins([]):
        yield
    plugin_http.set_transport(None)


@pytest.fixture(autouse=True)
def fake_plugin_in_plans(monkeypatch: pytest.MonkeyPatch) -> None:
    """While the framework's test plugin (fake_plugin.py, id "fake") is registered, it is in every plan that includes
    Gmail, which it stands in for: the plugin tests written before plans keep a user who may use it. Plan gating is
    tested with Gmail itself."""
    load, load_many = entitlements.load, entitlements.load_many

    def widened(ent: entitlements.Entitlements) -> entitlements.Entitlements:
        if "gmail" in ent.plugins and registry.get_plugin("fake") is not None:
            return replace(ent, plugins=ent.plugins | {"fake"})
        return ent

    async def load_widened(db: Any, user: User, *, now: datetime | None = None) -> entitlements.Entitlements:
        return widened(await load(db, user, now=now))

    async def load_many_widened(
        db: Any, users: Sequence[User], *, now: datetime | None = None
    ) -> dict[int, entitlements.Entitlements]:
        return {user_id: widened(ent) for user_id, ent in (await load_many(db, users, now=now)).items()}

    monkeypatch.setattr(entitlements, "load", load_widened)
    monkeypatch.setattr(entitlements, "load_many", load_many_widened)


TEST_TOKEN_KEY = Fernet.generate_key().decode()
TEST_ORIGIN = "http://testserver"


@dataclass
class PluginHarness:
    """The fake Google the `plugins` fixture installed. `app` is its FastAPI app: add Gmail routes with
    `app.include_router(...)`."""

    google: FakeGoogle
    app: FastAPI

    def consent_client(self) -> httpx.AsyncClient:
        return httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app), base_url="https://accounts.google.com")

    async def start(
        self, client: httpx.AsyncClient, plugin_id: str, *, capabilities: list[str] | None = None
    ) -> httpx.Response:
        """POST connect from TEST_ORIGIN; returns the response (200: `authorize_url`, and the nonce cookie set)."""
        body = {"capabilities": capabilities} if capabilities is not None else None
        return await client.post(f"/api/plugins/{plugin_id}/connect", json=body, headers={"Origin": TEST_ORIGIN})

    async def consent(self, authorize_url: str) -> str:
        """The fake's consent: the URL it redirects the browser to (the app's callback)."""
        async with self.consent_client() as google:
            response = await google.get(authorize_url)
        assert response.status_code == 302, response.text
        return response.headers["location"]

    async def connect(
        self,
        client: httpx.AsyncClient,
        plugin_id: str,
        *,
        account: str | None = None,
        capabilities: list[str] | None = None,
    ) -> httpx.Response:
        """The whole flow as a browser: connect, consent at the fake as `account`, follow the redirect to the callback.
        Returns the callback's response (303 to `/?plugin=<id>&oauth=<result>`)."""
        if account is not None:
            self.google.login_as = account
        started = await self.start(client, plugin_id, capabilities=capabilities)
        assert started.status_code == 200, started.text
        location = await self.consent(started.json()["authorize_url"])
        return await client.get(location)


def oauth_result(response: httpx.Response) -> str:
    """The `oauth` result of a callback's redirect."""
    assert response.status_code == 303, response.text
    return httpx.URL(response.headers["location"]).params["oauth"]


@pytest.fixture
def plugins(monkeypatch: pytest.MonkeyPatch) -> PluginHarness:
    """A fake Google behind the plugins' HTTP client, the fake OAuth client in the settings, a test
    TOKEN_ENCRYPTION_KEY and PUBLIC_BASE_URLS=http://testserver. Install a plugin with registry.use_plugins."""
    google = FakeGoogle()
    google.add_client(
        CLIENT_ID,
        CLIENT_SECRET,
        [f"{TEST_ORIGIN}/api/plugins/{plugin_id}/oauth/callback" for plugin_id in ("fake", "gmail")],
    )
    google.accounts.update({"bob@example.test": "100000000000000000002", "carol@example.test": "100000000000000000003"})
    app = fake_google_app(google)
    monkeypatch.setattr(settings, "google_oauth_client_id", CLIENT_ID)
    monkeypatch.setattr(settings, "google_oauth_client_secret", CLIENT_SECRET)
    monkeypatch.setattr(settings, "token_encryption_key", TEST_TOKEN_KEY)
    monkeypatch.setattr(settings, "public_base_urls", TEST_ORIGIN)
    monkeypatch.setattr(connections, "_sleep", _no_sleep)
    plugin_http.set_transport(httpx.ASGITransport(app=app))
    return PluginHarness(google, app)


async def _no_sleep(seconds: float) -> None:
    return None


@pytest.fixture
def image_tools(monkeypatch: pytest.MonkeyPatch, fake_llm: FakeLLM) -> FakeLLM:
    """CHAT_IMAGE_TOOLS on, as in production: the seeded GPT Image 2 is the active image model (the test environment
    sets OPENAI_API_KEY). Returns the FakeLLM."""
    monkeypatch.setattr(settings, "chat_image_tools", True)
    return fake_llm


@dataclass
class FileChild:
    """The fake file render child installed by the `file_child` fixture; `pids` lists the children started."""

    directory: Path

    @property
    def pids(self) -> list[int]:
        path = self.directory / "starts"
        return [int(line) for line in path.read_text().splitlines()] if path.exists() else []


_FAKE_FILE_CHILD = """
import json, os, sys, time
config = json.loads(sys.argv[1])
with open(os.path.join(config["directory"], "starts"), "a") as log:
    log.write(str(os.getpid()) + "\\n")
mode = config["mode"]
if mode == "sleep":
    time.sleep(600)
elif mode == "crash":
    print(json.dumps({"crash": config.get("crash", "MemoryError")}))
elif mode == "silent":
    sys.stderr.write("SECRET STDERR CONTENT\\n")
    sys.exit(3)
elif mode == "wait":
    while not os.path.exists(os.path.join(config["directory"], "go")):
        time.sleep(0.02)
"""


@pytest.fixture
def file_child(monkeypatch: pytest.MonkeyPatch, isolated_state: Path) -> Callable[..., FileChild]:
    """Factory: `file_child("sleep" | "crash" | "silent" | "wait", crash="MemoryError")` replaces the file render
    child with a script that sleeps 10 minutes, prints a crash result, exits without output (writing to stderr), or
    waits for a file named "go" in its directory and then exits without output."""

    def install(mode: str, *, crash: str = "MemoryError") -> FileChild:
        directory = Path(tempfile.mkdtemp(dir=isolated_state))
        script = directory / "fake_file_child.py"
        script.write_text(_FAKE_FILE_CHILD)
        config = json.dumps({"directory": str(directory), "mode": mode, "crash": crash})
        monkeypatch.setattr(document_tool, "_child_command", lambda: [sys.executable, "-I", str(script), config])
        return FileChild(directory)

    return install


@pytest.fixture(autouse=True)
def ingestion_sleeps(monkeypatch: pytest.MonkeyPatch) -> list[float]:
    """The waits ingestion asked for before retrying a dropped connection or a timeout (none actually happen)."""
    sleeps: list[float] = []

    async def record(seconds: float) -> None:
        sleeps.append(seconds)

    monkeypatch.setattr(ingestion, "_sleep", record)
    return sleeps


@pytest.fixture(autouse=True)
def vision_sleeps(monkeypatch: pytest.MonkeyPatch) -> list[float]:
    """The waits the vision reader asked for before retrying a call (none actually happen)."""
    sleeps: list[float] = []

    async def record(seconds: float) -> None:
        sleeps.append(seconds)

    monkeypatch.setattr(vision, "_RETRY_DELAYS", (0.0, 0.0, 0.0, 0.0))
    monkeypatch.setattr(vision, "_sleep", record)
    return sleeps


@dataclass
class RenderChild:
    """The fake render child installed by the `render_child` fixture. `pids` lists the children started, in order;
    `environment` holds the environment variable names the last one saw."""

    directory: Path

    @property
    def pids(self) -> list[int]:
        path = self.directory / "starts"
        return [int(line) for line in path.read_text().splitlines()] if path.exists() else []

    @property
    def environment(self) -> list[str]:
        return json.loads((self.directory / "environment.json").read_text())


_FAKE_RENDER_CHILD = """
import importlib.util, io, json, os, sys, time
config = json.loads(sys.argv[1])
with open(os.path.join(config["directory"], "starts"), "a") as log:
    log.write(str(os.getpid()) + "\\n")
with open(os.path.join(config["directory"], "environment.json"), "w") as out:
    json.dump(sorted(os.environ), out)
spec = importlib.util.spec_from_file_location("real_child", config["real"])
child = importlib.util.module_from_spec(spec)
spec.loader.exec_module(child)
render = child._render

def fake_render(pdf, job, request):
    page = int(job["page"])
    if page in config["abort"]:
        os.abort()
    if page in config["hang"]:
        time.sleep(600)
    if page in config["fail"]:
        return {"page": page, "error": "InjectedError"}
    return render(pdf, job, request)

child._render = fake_render
sys.exit(child.main())
"""


@pytest.fixture
def render_child(
    monkeypatch: pytest.MonkeyPatch, isolated_state: Path
) -> Callable[..., RenderChild]:
    """Factory: `render_child(abort=[2], hang=[3], fail=[4])` replaces the PDF render child with one that renders
    like the real one, except that it calls os.abort() on the `abort` pages, sleeps 10 minutes on the `hang`
    pages and reports an error for the `fail` pages. Lower `pdf_images._RENDER_PAGE_SECONDS` for `hang`."""

    def install(*, abort: Sequence[int] = (), hang: Sequence[int] = (), fail: Sequence[int] = ()) -> RenderChild:
        directory = Path(tempfile.mkdtemp(dir=isolated_state))
        script = directory / "fake_render_child.py"
        script.write_text(_FAKE_RENDER_CHILD)
        config = {
            "directory": str(directory),
            "real": str(Path(pdf_images.__file__).with_name("pdf_render_child.py")),
            "abort": list(abort),
            "hang": list(hang),
            "fail": list(fail),
        }
        command = [sys.executable, "-I", str(script), json.dumps(config)]
        monkeypatch.setattr(pdf_images, "_child_command", lambda: command)
        return RenderChild(directory)

    return install


@pytest.fixture
def child_starts(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    """Counts the real render children started: one entry per start."""
    starts: list[int] = []
    command = pdf_images._child_command

    def counted() -> list[str]:
        starts.append(1)
        return command()

    monkeypatch.setattr(pdf_images, "_child_command", counted)
    return starts


# ---------------- App and clients ----------------


class ApiClient(httpx.AsyncClient):
    """In-process HTTP client for the app with its own cookie jar.

    `user_id`, `username` and `password` identify the account it is logged in as (None while anonymous).
    """

    def __init__(self, app: FastAPI) -> None:
        super().__init__(transport=httpx.ASGITransport(app=app), base_url="http://testserver")
        self.user_id: int | None = None
        self.username: str | None = None
        self.password: str | None = None

    async def login(self, username: str, password: str) -> httpx.Response:
        """POST /api/auth/login. On 200 the session cookie is stored and the identity attributes are set."""
        response = await self.post("/api/auth/login", json={"username": username, "password": password})
        if response.status_code == 200:
            me = response.json()
            self.user_id, self.username, self.password = me["id"], me["username"], password
        return response


@pytest.fixture
async def app(isolated_state: Path) -> AsyncIterator[FastAPI]:
    """The FastAPI app with its lifespan running: tables created, admin synced, models seeded."""
    async with fastapi_app.router.lifespan_context(fastapi_app):
        yield fastapi_app


@pytest.fixture
async def new_client(app: FastAPI) -> AsyncIterator[Callable[[], ApiClient]]:
    """Factory for extra anonymous clients (e.g. a second session of the same user); closed at teardown."""
    clients: list[ApiClient] = []

    def factory() -> ApiClient:
        client = ApiClient(app)
        clients.append(client)
        return client

    yield factory
    for client in clients:
        await client.aclose()


@pytest.fixture
def client(new_client: Callable[[], ApiClient]) -> ApiClient:
    """An anonymous client."""
    return new_client()


@pytest.fixture
async def admin_client(new_client: Callable[[], ApiClient]) -> ApiClient:
    """A client logged in as the env admin (ADMIN_USERNAME / ADMIN_PASSWORD)."""
    admin = new_client()
    response = await admin.login(ADMIN_USERNAME, ADMIN_PASSWORD)
    assert response.status_code == 200, response.text
    return admin


@pytest.fixture
def create_user(admin_client: ApiClient, new_client: Callable[[], ApiClient]) -> Callable[..., Awaitable[ApiClient]]:
    """Factory: `await create_user("alice", **fields)` creates a user through POST /api/admin/users and
    returns a client logged in as them.

    `fields` are any AdminUserCreate fields (restrict_models, allowed_model_ids, is_active, ...) or limits, sent as
    overrides (an unlimited chat allowance confirmed); omitted limits are the plan's. The user gets an open-ended
    complimentary plan, `plan` ("professional" unless given; None: no plan), so its allowance window is the
    calendar month and limits passed here are overrides of the plan's. The name defaults to user1, user2,
    ... and the password to USER_PASSWORD. With `login=False` the client stays anonymous (needed for
    `is_active=False`); its `user_id`/`username`/`password` are set either way. Asserts the user was created.
    """
    created = 0

    async def factory(
        username: str | None = None,
        password: str = USER_PASSWORD,
        *,
        login: bool = True,
        plan: str | None = "professional",
        **fields: Any,
    ) -> ApiClient:
        nonlocal created
        created += 1
        name = username or f"user{created}"
        overrides = {key: fields.pop(key) for key in list(fields) if key in periods.LIMIT_NAMES}
        if overrides:
            fields["overrides"] = overrides
        if overrides.get("monthly_token_limit") == 0:
            fields.setdefault("confirm_unlimited", True)
        response = await admin_client.post("/api/admin/users", json={"username": name, "password": password, **fields})
        assert response.status_code == 201, response.text
        if plan is not None:
            await grant_plan(response.json()["id"], plan, days=None)
        user = new_client()
        if login:
            login_response = await user.login(name, password)
            assert login_response.status_code == 200, login_response.text
        account = response.json()
        user.user_id, user.username, user.password = account["id"], account["username"], password
        return user

    return factory


# ---------------- Plans ----------------


async def grant_plan(
    user_id: int, code: str = "professional", days: int | None = 30, start: Literal["now", "after_current"] = "now"
) -> None:
    """A complimentary plan, as the admin grants it: dated (`days`, in rows of at most 30 days) or open-ended (None)."""
    async with db.SessionLocal() as session:
        user = await session.get(User, user_id)
        plan = await session.scalar(select(Plan).where(Plan.code == code))
        assert user is not None and plan is not None, (user_id, code)
        await periods.grant(session, user, plan, days=days, reason="Test grant", actor=ADMIN_USERNAME, start=start)
        await session.commit()


async def end_plan(user_id: int) -> None:
    """The user's plan has ended: their periods move into the past (the last ending a minute ago) and the open-ended
    plan ends. Expired with periods, else no plan at all."""
    async with db.SessionLocal() as session:
        user = await session.get(User, user_id)
        assert user is not None, user_id
        rows = list(await session.scalars(select(PlanPeriod).where(PlanPeriod.user_id == user_id)))
        if rows:
            shift = max(row.ends_at for row in rows) - utcnow() + timedelta(minutes=1)
            for row in rows:
                row.starts_at, row.ends_at = row.starts_at - shift, row.ends_at - shift
        user.comp_plan_id = user.comp_reason = user.comp_granted_by = user.comp_granted_at = None
        await session.commit()


async def make_tester(user_id: int) -> None:
    """The user may check out while test keys run (Admin → Users, "Tester")."""
    async with db.SessionLocal() as session:
        user = await session.get(User, user_id)
        assert user is not None, user_id
        user.billing_tester = True
        await session.commit()


@pytest.fixture
def fake_razorpay(monkeypatch: pytest.MonkeyPatch) -> FakeRazorpay:
    """Razorpay test keys and a webhook secret in the settings, and every HTTPS request to api.razorpay.com (the
    default RAZORPAY_API_BASE) answered by the fake (tests/fake_razorpay.py) through a monkeypatched httpx transport;
    other hosts are untouched. Drive the payer and the Dashboard through the returned FakeRazorpay."""
    fake = FakeRazorpay()
    monkeypatch.setattr(settings, "razorpay_key_id", fake.key_id)
    monkeypatch.setattr(settings, "razorpay_key_secret", fake.key_secret)
    monkeypatch.setattr(settings, "razorpay_webhook_secret", fake.webhook_secret)
    asgi = httpx.ASGITransport(app=fake_razorpay_app(fake))
    network = httpx.AsyncHTTPTransport.handle_async_request

    async def handle(transport: httpx.AsyncHTTPTransport, request: httpx.Request) -> httpx.Response:
        if request.url.host != API_HOST:
            return await network(transport, request)
        return fake.injected(request) or await asgi.handle_async_request(request)

    monkeypatch.setattr(httpx.AsyncHTTPTransport, "handle_async_request", handle)
    return fake


# ---------------- Documents and ingestion ----------------


async def upload(
    client: httpx.AsyncClient,
    filename: str,
    data: bytes,
    scope: str = "private",
    content_type: str | None = None,
) -> httpx.Response:
    """POST /api/documents as multipart (`file` + `scope`). Returns the raw response (201 on success)."""
    file = (filename, data) if content_type is None else (filename, data, content_type)
    return await client.post("/api/documents", files={"file": file}, data={"scope": scope})


async def process(doc_id: int) -> Document | None:
    """Take the document off the ingestion queue and process it in-process, like the worker does.

    Returns its row afterwards: a detached but fully readable ORM object (`status`, `error`,
    `num_chunks`, `embedding_model_id`, `owner`, ...), or None if the document no longer exists.
    """
    await get_redis().lrem(INGEST_QUEUE, 0, str(doc_id))
    await process_document(doc_id)
    async with db.SessionLocal() as session:
        return await session.get(Document, doc_id)


async def process_queue() -> list[int]:
    """Pop every id from the ingestion queue and process it, like the worker. Returns the ids in order."""
    redis = get_redis()
    processed: list[int] = []
    while (item := await redis.lpop(INGEST_QUEUE)) is not None:
        processed.append(int(item))
        await process_document(int(item))
    return processed


async def ingest(client: ApiClient, filename: str, data: bytes, scope: str = "private") -> dict[str, Any]:
    """Upload (asserting 201), process, and return the document's DocumentItem JSON as `client` now sees it."""
    response = await upload(client, filename, data, scope)
    assert response.status_code == 201, response.text
    doc_id = response.json()["id"]
    await process(doc_id)
    listing = await client.get("/api/documents")
    assert listing.status_code == 200, listing.text
    return next(item for item in listing.json() if item["id"] == doc_id)


# ---------------- Chat (SSE) ----------------


def sse_events(response: httpx.Response) -> list[SseEvent]:
    """Parse a text/event-stream body into (event name, JSON data) pairs, in order; comment lines are ignored."""
    events: list[SseEvent] = []
    for block in response.text.split("\n\n"):
        if not block.strip():
            continue
        name, data = "message", []
        for line in block.split("\n"):
            if line.startswith("event:"):
                name = line.removeprefix("event:").strip()
            elif line.startswith("data:"):
                data.append(line.removeprefix("data:").removeprefix(" "))
        if data:  # comment-only blocks (": keepalive") carry no event
            events.append((name, json.loads("\n".join(data))))
    return events


async def chat(
    client: httpx.AsyncClient,
    message: str,
    *,
    conversation_id: int | None = None,
    model_id: int | None = None,
    document_ids: list[int] | None = None,
) -> tuple[httpx.Response, list[SseEvent]]:
    """POST /api/chat/stream and read the whole stream. Returns the response and its parsed events
    (an empty list unless the status is 200)."""
    body = {"conversation_id": conversation_id, "message": message, "model_id": model_id, "document_ids": document_ids}
    response = await client.post("/api/chat/stream", json=body)
    return response, sse_events(response) if response.status_code == 200 else []


# ---------------- File builders ----------------


def _pdf_string(line: str) -> str:
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_bytes(*pages: str) -> bytes:
    """A PDF with one page per argument and a real text layer (Helvetica, one text line per input line).

    An empty or blank page has no text layer at all, so `pdf_bytes("")` acts as a scanned PDF. Text must
    be Latin-1 (other characters become "?").
    """
    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    kids = []
    for number, text in enumerate(pages or ("",)):
        page_id, content_id = 4 + 2 * number, 5 + 2 * number
        kids.append(f"{page_id} 0 R")
        shown = " ".join(f"({_pdf_string(line)}) Tj T*" for line in text.splitlines())
        stream = f"BT /F1 12 Tf 14 TL 72 760 Td {shown} ET".encode("latin-1", "replace") if text.strip() else b""
        objects[content_id] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
        objects[page_id] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()
    return _pdf_file(objects)


def _pdf_file(objects: Mapping[int, bytes]) -> bytes:
    """A PDF file from numbered object bodies (object 1 must be the catalog)."""
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = {}
    for object_id in sorted(objects):
        offsets[object_id] = len(out)
        out += b"%d 0 obj\n%s\nendobj\n" % (object_id, objects[object_id])
    xref = len(out)
    size = max(objects) + 1
    out += b"xref\n0 %d\n0000000000 65535 f \n" % size
    out += b"".join(
        b"%010d 00000 n \n" % offsets[object_id] if object_id in offsets else b"0000000000 65535 f \n"
        for object_id in range(1, size)
    )
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref)
    return bytes(out)


def pdf_stream(data: bytes, entries: str = "") -> bytes:
    """A stream object body: `entries` are extra dictionary entries, e.g. "/Type /XObject /Subtype /Form ..."."""
    return b"<< /Length %d %s >>\nstream\n%s\nendstream" % (len(data), entries.encode(), data)


@functools.cache
def chart_jpeg() -> bytes:
    """A 400x250 bar chart drawn by Pillow, as JPEG."""
    image = Image.new("RGB", (400, 250), "white")
    draw = ImageDraw.Draw(image)
    for index, height in enumerate((60, 130, 170, 210)):
        draw.rectangle((40 + index * 85, 230 - height, 100 + index * 85, 230), fill=(40, 90, 200))
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=80)
    return buffer.getvalue()


def jpeg_xobject(jpeg: bytes | None = None) -> bytes:
    """An image XObject body holding `jpeg` (default: chart_jpeg())."""
    data = jpeg or chart_jpeg()
    with Image.open(io.BytesIO(data)) as image:
        width, height = image.size
    return pdf_stream(
        data,
        f"/Type /XObject /Subtype /Image /Width {width} /Height {height} /ColorSpace /DeviceRGB "
        "/BitsPerComponent 8 /Filter /DCTDecode",
    )


def _text_layer(text: str) -> str:
    shown = " ".join(f"({_pdf_string(line)}) Tj T*" for line in text.splitlines())
    return f"BT /F1 12 Tf 14 TL 72 760 Td {shown} ET" if text.strip() else ""


def image_pdf_bytes(
    *pages: tuple[str, float], rotate: int = 0, in_form: bool = False, shared_image: bool = True
) -> bytes:
    """One US Letter page per (text, area): a Helvetica text layer holding `text` (none when blank) and, when
    `area` > 0, chart_jpeg() stretched over the bottom `area` share of the page.

    `in_form` draws the image through a form XObject with a scaling /Matrix (the placed area is the same).
    `shared_image=False` gives every page its own image XObject; with one shared image, the vision reader ignores
    it as a repeated image (letterhead) on 3 or more pages, or on every page. Text pages need at least 40
    non-whitespace characters to count as having text.
    """
    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        4: jpeg_xobject(),
    }
    kids = []
    for number, (text, area) in enumerate(pages or (("", 0.0),)):
        page_id, content_id, image_id, form_id = (10 + 4 * number + offset for offset in range(4))
        kids.append(f"{page_id} 0 R")
        image_ref = 4
        if not shared_image:
            objects[image_id] = jpeg_xobject()
            image_ref = image_id
        parts, xobjects = [], f"/Im1 {image_ref} 0 R"
        if area > 0:
            if in_form:
                form = f"q 1224 0 0 {1584 * area:.3f} 0 0 cm /Im1 Do Q".encode()
                objects[form_id] = pdf_stream(
                    form,
                    "/Type /XObject /Subtype /Form /BBox [0 0 1224 1584] /Matrix [0.5 0 0 0.5 0 0] "
                    f"/Resources << /XObject << /Im1 {image_ref} 0 R >> >>",
                )
                parts.append("q /Fm1 Do Q")
                xobjects = f"/Fm1 {form_id} 0 R"
            else:
                parts.append(f"q 612 0 0 {792 * area:.3f} 0 0 cm /Im1 Do Q")
        parts.append(_text_layer(text))
        objects[content_id] = pdf_stream(" ".join(parts).encode("latin-1", "replace"))
        rotation = f"/Rotate {rotate} " if rotate else ""
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] {rotation}"
            f"/Resources << /Font << /F1 3 0 R >> /XObject << {xobjects} >> >> /Contents {content_id} 0 R >>"
        ).encode()
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()
    return _pdf_file(objects)


def raw_pdf_bytes(
    *contents: bytes,
    objects: Mapping[int, bytes] | None = None,
    xobjects: str = "",
    media_box: str = "0 0 612 792",
    rotate: int = 0,
    compress: bool = False,
) -> bytes:
    """One page per content stream, all with the same resources: Helvetica as /F1 and the `xobjects` entries
    (e.g. "/S0 100 0 R /S1 101 0 R"). `objects` adds numbered objects; use numbers from 100 (see pdf_stream,
    jpeg_xobject). `compress` stores the contents with FlateDecode. For strips, invisible text (`3 Tr`),
    overlays, inline images and crafted pages."""
    all_objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        **(objects or {}),
    }
    kids = []
    rotation = f"/Rotate {rotate} " if rotate else ""
    for number, content in enumerate(contents or (b"",)):
        page_id, content_id = 10 + 2 * number, 11 + 2 * number
        kids.append(f"{page_id} 0 R")
        all_objects[content_id] = (
            pdf_stream(zlib.compress(content, 9), "/Filter /FlateDecode") if compress else pdf_stream(content)
        )
        all_objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [{media_box}] {rotation}"
            f"/Resources << /Font << /F1 3 0 R >> /XObject << {xobjects} >> >> /Contents {content_id} 0 R >>"
        ).encode()
    all_objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()
    return _pdf_file(all_objects)


def page_scan(text: str, size: tuple[int, int] = (1240, 1754), mode: str = "RGB") -> Image.Image:
    """A white page image (A4 at 150 DPI) with `text` in black, one line per input line."""
    image = Image.new(mode, size, "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=36)
    for index, line in enumerate(text.splitlines()):
        draw.text((100, 120 + 50 * index), line, fill="black", font=font)
    return image


def scanned_pdf_bytes(*texts: str, mode: str = "RGB") -> bytes:
    """An image-only PDF written by Pillow: one scanned page per text, no text layer."""
    images = [page_scan(text, mode=mode) for text in texts or ("",)]
    buffer = io.BytesIO()
    images[0].save(buffer, "PDF", resolution=150, save_all=True, append_images=images[1:])
    return buffer.getvalue()


def docx_bytes(*paragraphs: str, table: Sequence[Sequence[str]] | None = None) -> bytes:
    """A .docx with these paragraphs, then `table` (rows of cell texts) if given."""
    document = new_docx()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=max(len(row) for row in table))
        for grid_row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(grid_row.cells, values, strict=False):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes(*slides: str, notes: Sequence[str] = ()) -> bytes:
    """A .pptx with one slide per argument (a text box holding the text); `notes[i]` adds speaker notes."""
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for number, text in enumerate(slides):
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5)).text_frame.text = text
        if number < len(notes) and notes[number]:
            slide.notes_slide.notes_text_frame.text = notes[number]
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(rows: Sequence[Sequence[Any]], title: str = "Sheet1") -> bytes:
    """A one-sheet .xlsx named `title` holding `rows` (cell values: str, int, float, datetime, None)."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def html_bytes(text: str, title: str = "Test page") -> bytes:
    """An HTML page: `title`, a <script> the parser must drop ("SCRIPT_MARKER"), and one <p> per line of `text`."""
    paragraphs = "".join(f"<p>{escape(line)}</p>" for line in text.splitlines())
    return (
        f"<!doctype html><html><head><title>{escape(title)}</title>"
        "<script>var marker = 'SCRIPT_MARKER';</script></head>"
        f"<body>{paragraphs}</body></html>"
    ).encode()


def md_bytes(text: str) -> bytes:
    """A Markdown document: a heading followed by `text`."""
    return f"# Notes\n\n{text}\n".encode()


def txt_bytes(text: str) -> bytes:
    """`text` as UTF-8."""
    return text.encode()


def csv_bytes(rows: Sequence[Sequence[Any]]) -> bytes:
    """CSV of `rows`."""
    buffer = io.StringIO()
    csv.writer(buffer).writerows(rows)
    return buffer.getvalue().encode()


def json_bytes(value: Any) -> bytes:
    """`value` as compact JSON (the parser pretty-prints it)."""
    return json.dumps(value).encode()


def sample_file(extension: str, text: str) -> bytes:
    """A valid file of any supported extension (e.g. ".pdf") whose extracted text contains `text`.

    Keep `text` Latin-1 for ".pdf".
    """
    builders: dict[str, Callable[[], bytes]] = {
        ".pdf": lambda: pdf_bytes(text),
        ".docx": lambda: docx_bytes(text, table=[["Column A", "Column B"], ["cell one", "cell two"]]),
        ".pptx": lambda: pptx_bytes(text),
        ".xlsx": lambda: xlsx_bytes([["Text"], [text]]),
        ".html": lambda: html_bytes(text),
        ".htm": lambda: html_bytes(text),
        ".md": lambda: md_bytes(text),
        ".markdown": lambda: md_bytes(text),
        ".txt": lambda: txt_bytes(text),
        ".csv": lambda: csv_bytes([["text"], [text]]),
        ".json": lambda: json_bytes({"text": text}),
    }
    return builders[extension.lower()]()
