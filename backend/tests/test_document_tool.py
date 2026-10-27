"""Files from chat through POST /api/chat/stream: the create_document loop, its events, replay, limits, quotas,
recoveries, usage, Stop, provenance and history (docs/briefs/document-generation.md, section 14.5)."""

import asyncio
import hashlib
import io
import json
import time
from collections.abc import AsyncIterator, Awaitable, Callable
from typing import Any

import docx
import litellm
import openpyxl
import pytest
from conftest import (
    ApiClient,
    FakeLLM,
    FileChild,
    ScriptedCall,
    SseEvent,
    chat,
    process,
    sse_events,
    tool_call,
    upload,
)
from pypdf import PdfReader
from sqlalchemy import select

from app.config import settings
from app.db import SessionLocal
from app.models import Document, GeneratedDocument, Message, MessageFile, UsageEvent
from app.redis_client import INGEST_QUEUE, get_redis
from app.routers import chat as chat_router
from app.routers.chat import _stream_answer, _Turn
from app.services import document_tool, image_tool, llm, rag, storage
from app.services import usage as usage_service
from app.services.billing import entitlements

CreateUser = Callable[..., Awaitable[ApiClient]]
StartTurn = Callable[..., Awaitable[_Turn]]

REPORT = "# Q3 report\n\nRevenue grew **12%** in the north [1].\n\n- North: up\n- South: flat\n"
DECLINED = "The model declined to continue, so the file wasn't created."
INCOMPLETE = "The model's response ended before the file was complete, so it wasn't created."
CALL_FAILED = "The model call failed before the file was finished."
BUDGET_SPENT = (
    "This period's chat allowance ran out while this answer was being written, so it stops here. Files created above "
    "are in your documents."
)
STOPPED = "Stopped before completion"


def pdf_call(title: str = "Q3 report", content: str = REPORT, fmt: str = "pdf", **extra: Any) -> tuple[str, Any]:
    return tool_call(title=title, format=fmt, content=content, **extra)


def names(events: list[SseEvent]) -> list[str]:
    return [name for name, _ in events]


def file_events(events: list[SseEvent]) -> list[dict[str, Any]]:
    return [data for name, data in events if name == "file"]


def tool_messages(call: dict[str, Any]) -> list[dict[str, Any]]:
    return [json.loads(message["content"]) for message in call["messages"] if message["role"] == "tool"]


def system_prompt(call: dict[str, Any]) -> str:
    return call["messages"][0]["content"]


async def rows(model: Any, *conditions: Any) -> list[Any]:
    async with SessionLocal() as session:
        return list(await session.scalars(select(model).where(*conditions)))


async def chat_usage(user_id: int | None) -> list[UsageEvent]:
    return await rows(UsageEvent, UsageEvent.user_id == user_id, UsageEvent.kind == "chat")


async def answer_of(conversation_id: int) -> Message:
    messages = await rows(Message, Message.conversation_id == conversation_id)
    return max(messages, key=lambda message: message.id)


def stored_files() -> list[str]:
    return sorted(path.name for path in storage.upload_dir().iterdir())


async def model_id(client: ApiClient, display_name: str) -> int:
    models = (await client.get("/api/models")).json()
    return next(model["id"] for model in models if model["display_name"] == display_name)


async def can_create_files(client: ApiClient) -> dict[str, bool]:
    return {model["display_name"]: model["can_create_files"] for model in (await client.get("/api/models")).json()}


async def _no_events() -> AsyncIterator[str]:
    for event in ():
        yield event


@pytest.fixture
def start_turn(monkeypatch: pytest.MonkeyPatch) -> StartTurn:
    """POST /api/chat/stream without running its stream; returns the turn so a test can drive `_stream_answer`."""
    turns: list[_Turn] = []

    def capture(turn: _Turn) -> AsyncIterator[str]:
        turns.append(turn)
        return _no_events()

    monkeypatch.setattr(chat_router, "_stream_answer", capture)

    async def start(client: ApiClient, message: str, **fields: Any) -> _Turn:
        response, events = await chat(client, message, **fields)
        assert response.status_code == 200, response.text
        return turns[-1]

    return start


@pytest.fixture(params=["files", "files+images"])
def tool_mode(request: pytest.FixtureRequest, monkeypatch: pytest.MonkeyPatch) -> str:
    """Runs a scenario as written (images off, as in every test here) and in the production configuration, with
    create_image offered next to create_document (docs/briefs/image-generation.md, I19)."""
    if request.param == "files+images":
        monkeypatch.setattr(settings, "chat_image_tools", True)
    return str(request.param)


def offered_tools(tool_mode: str) -> list[str]:
    return ["create_document", "create_image"] if tool_mode == "files+images" else ["create_document"]


def tool_names(call: dict[str, Any]) -> list[str]:
    return [tool["function"]["name"] for tool in call.get("tools") or []]


def parse(raw: list[str]) -> list[SseEvent]:
    import httpx

    return sse_events(httpx.Response(200, text="".join(raw)))


# ---------------- One file, end to end ----------------


async def test_a_pdf_is_created_stored_linked_and_indexed(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch, tool_mode: str
) -> None:
    monkeypatch.setattr(settings, "generated_files_per_answer", 1)  # so the second call is the last one (files only)
    alice = await create_user("alice")
    fake_llm.script = [
        ScriptedCall(
            [("thinking", (["I will ", "write ", "it."], "SIG-1")), ("text", "Creating it."), pdf_call()],
            usage=(1000, 400),
        ),
        ScriptedCall([("text", "Done: Q3 report.pdf is in your documents.")], usage=(1500, 30)),
    ]

    response, events = await chat(alice, "Make a PDF report of Q3", document_ids=[])

    assert names(events) == ["meta", "sources", "delta", "file", "file", "file", "delta", "delta", "done"]
    writing, validated, created = file_events(events)
    assert writing == {"call_id": "f1", "kind": "file", "status": "writing", "title": None, "format": None, "chars": 0}
    assert (validated["call_id"], validated["status"], validated["title"], validated["format"]) == (
        "f1",
        "writing",
        "Q3 report",
        "pdf",
    )
    assert validated["chars"] > len(REPORT)
    assert (created["call_id"], created["status"], created["already_existed"]) == ("f1", "created", False)
    file = created["file"]
    doc_id = file["document_id"]
    [doc] = await rows(Document, Document.id == doc_id)
    assert file == {
        "document_id": doc_id,
        "filename": "Q3 report.pdf",
        "format": "pdf",
        "size_bytes": doc.size_bytes,
        "status": "pending",
        "image": None,
    }
    assert created["document"]["id"] == doc_id and created["document"]["generated"] is True
    assert created["document"]["created_at"].endswith("Z")
    assert [data["content"] for name, data in events if name == "delta"] == [
        "Creating it.",
        "\n\n",
        "Done: Q3 report.pdf is in your documents.",
    ]

    first, second = fake_llm.chat_calls
    assert (first["tool_choice"], first["parallel_tool_calls"], first["max_tokens"]) == ("auto", False, 48_000)
    if tool_mode == "files":
        assert second["tool_choice"] == "none" and "parallel_tool_calls" not in second
        assert second["max_tokens"] == 16_000
    else:  # images can still be made, so the second call still offers the tools
        assert (second["tool_choice"], second["parallel_tool_calls"], second["max_tokens"]) == ("auto", False, 48_000)
    assert tool_names(first) == offered_tools(tool_mode)
    assert first["tools"] == second["tools"] and first["tools"][0]["eager_input_streaming"] is True
    assert system_prompt(first) == system_prompt(second)
    assert "You can create files with the create_document tool" in system_prompt(first)
    assert "about 16,000 words" in system_prompt(first)
    replay = second["messages"][len(first["messages"]) :]
    assistant, tool = replay
    assert assistant["role"] == "assistant" and assistant["content"] == "Creating it."
    assert assistant["thinking_blocks"] == [{"type": "thinking", "thinking": "I will write it.", "signature": "SIG-1"}]
    [echoed] = assistant["tool_calls"]
    assert echoed["id"] == "toolu_0" and json.loads(echoed["function"]["arguments"])["title"] == "Q3 report"
    assert tool["role"] == "tool" and tool["tool_call_id"] == "toolu_0"
    result = json.loads(tool["content"])
    assert result == {
        "ok": True,
        "filename": "Q3 report.pdf",
        "format": "pdf",
        "size_bytes": doc.size_bytes,
        "pages": 1,
        "already_existed": False,
        "note": "Saved to the user's private documents. Tell them the file name; they can download it from the file "
        "card. It becomes searchable once indexing finishes.",
    }

    done = events[-1][1]
    expected_cost = (2500 * 4.0 + 430 * 20.0) / 1_000_000
    assert done["usage"] == {
        "prompt_tokens": 2500,
        "completion_tokens": 430,
        "total_tokens": 2930,
        "cost_usd": pytest.approx(expected_cost),
    }
    [usage] = await chat_usage(alice.user_id)
    assert (usage.prompt_tokens, usage.completion_tokens) == (2500, 430)
    answer = await answer_of(events[0][1]["conversation_id"])
    assert answer.content == "Creating it.\n\nDone: Q3 report.pdf is in your documents."
    assert (answer.prompt_tokens, answer.completion_tokens, answer.error) == (2500, 430, None)
    detail = (await alice.get(f"/api/conversations/{events[0][1]['conversation_id']}")).json()
    assert detail["messages"][0]["files"] == [] and detail["messages"][1]["files"] == [file]

    assert (doc.owner_id, doc.scope, doc.filename, doc.content_type) == (
        alice.user_id,
        "private",
        "Q3 report.pdf",
        "application/pdf",
    )
    data = storage.stored_path(doc.stored_name).read_bytes()
    assert (len(data), hashlib.sha256(data).hexdigest()) == (doc.size_bytes, doc.sha256)
    assert len(PdfReader(io.BytesIO(data)).pages) == 1
    source = storage.stored_path(storage.source_name(doc.stored_name)).read_text(encoding="utf-8")
    assert source.startswith("# Q3 report\n\nRevenue grew **12%** in the north \\[1\\].")
    assert len(await rows(GeneratedDocument, GeneratedDocument.document_id == doc_id)) == 1
    listed = next(item for item in (await alice.get("/api/documents")).json() if item["id"] == doc_id)
    assert listed["generated"] is True and listed["filename"] == "Q3 report.pdf"
    assert await get_redis().lrange(INGEST_QUEUE, 0, -1) == [str(doc_id)]

    indexed = await process(doc_id)
    assert indexed is not None and (indexed.status, indexed.num_pages) == ("ready", 1)
    assert fake_llm.vision_calls == []  # the active vision model never sees a generated PDF

    fake_llm.script = []
    await chat(alice, "How much did Q3 revenue grow in the north?")
    prompt = system_prompt(fake_llm.chat_calls[-1])
    assert "(source: Q3 report.pdf, a file the assistant created on " in prompt
    assert "(ref 1)" in prompt and "\\[1\\]" not in prompt
    assert rag._GENERATED_PASSAGES_NOTE in prompt


async def test_interleaved_thinking_is_replayed_as_ordered_segments(create_user: CreateUser, fake_llm: FakeLLM) -> None:
    alice = await create_user("alice")
    fake_llm.script = [
        ScriptedCall(
            [
                ("thinking", (["Plan."], "SIG-1")),
                ("text", "Sure, one moment."),
                ("thinking", ([""], "SIG-2")),
                pdf_call(fmt="md"),
            ]
        ),
        ScriptedCall([("text", "Done.")]),
    ]

    await chat(alice, "Export it as markdown", document_ids=[])

    first, second = fake_llm.chat_calls
    segments = [message for message in second["messages"][len(first["messages"]) :] if message["role"] == "assistant"]
    assert [
        (
            message["content"],
            [block["signature"] for block in message.get("thinking_blocks", [])],
            "tool_calls" in message,
        )
        for message in segments
    ] == [("Sure, one moment.", ["SIG-1"], False), (None, ["SIG-2"], True)]


async def test_openai_model_creates_two_files_in_successive_calls(
    create_user: CreateUser, fake_llm: FakeLLM, tool_mode: str
) -> None:
    alice = await create_user("alice")
    mini = await model_id(alice, "GPT-5.4 mini")
    fake_llm.script = [
        ScriptedCall([pdf_call("Letter", "Dear team,\n\nThanks.", fmt="docx")], shape="openai"),
        ScriptedCall([pdf_call("Notes", "# Notes\n\n- one", fmt="md")], shape="openai"),
        ScriptedCall([("text", "Both files are in your documents.")], shape="openai"),
    ]

    response, events = await chat(alice, "A Word letter and markdown notes, please", model_id=mini, document_ids=[])

    created = [data for data in file_events(events) if data["status"] == "created"]
    assert [data["file"]["filename"] for data in created] == ["Letter.docx", "Notes.md"]
    assert [data["call_id"] for data in created] == ["f1", "f2"]
    calls = fake_llm.chat_calls
    assert len(calls) == 3 and "eager_input_streaming" not in calls[0]["tools"][0]
    assert all(tool_names(call) == offered_tools(tool_mode) for call in calls)
    assert [result["filename"] for result in tool_messages(calls[2])] == ["Letter.docx", "Notes.md"]
    assert [len(tool_messages(call)) for call in calls] == [0, 1, 2]
    answer = await answer_of(events[0][1]["conversation_id"])
    linked = await rows(MessageFile, MessageFile.message_id == answer.id)
    assert [(row.filename, row.format) for row in linked] == [("Letter.docx", "docx"), ("Notes.md", "md")]
    assert answer.content == "Both files are in your documents."


async def test_every_format_opens(create_user: CreateUser, fake_llm: FakeLLM) -> None:
    alice = await create_user("alice")
    sheets = [{"name": "Figures", "columns": ["Region", "Revenue"], "rows": [["North", "1200.5"], ["South", "980"]]}]
    fake_llm.script = [
        ScriptedCall([pdf_call("Report", fmt="pdf")]),
        ScriptedCall([pdf_call("Letter", "Dear team,\n\n**Thanks** for Q3.", fmt="docx")]),
        ScriptedCall([pdf_call("Export", "# Export\n\nPlain *text*.", fmt="md")]),
        ScriptedCall([("text", "Done.")]),
        ScriptedCall([tool_call(title="Figures", format="xlsx", sheets=sheets)]),
        ScriptedCall([("text", "Done.")]),
    ]

    await chat(alice, "Three files please", document_ids=[])
    await chat(alice, "And a spreadsheet", document_ids=[])

    documents = {doc.filename: doc for doc in await rows(Document, Document.owner_id == alice.user_id)}
    assert set(documents) == {"Report.pdf", "Letter.docx", "Export.md", "Figures.xlsx"}

    def read(name: str) -> bytes:
        return storage.stored_path(documents[name].stored_name).read_bytes()

    assert "Revenue grew 12%" in PdfReader(io.BytesIO(read("Report.pdf"))).pages[0].extract_text()
    letter = docx.Document(io.BytesIO(read("Letter.docx")))
    assert [paragraph.text for paragraph in letter.paragraphs] == ["Letter", "Dear team,", "Thanks for Q3."]
    assert read("Export.md").decode() == "# Export\n\nPlain *text*.\n"
    workbook = openpyxl.load_workbook(io.BytesIO(read("Figures.xlsx")))
    assert [[cell.value for cell in row] for row in workbook["Figures"].iter_rows()] == [
        ["Region", "Revenue"],
        ["North", 1200.5],
        ["South", 980],
    ]
    content_types = {doc.filename: doc.content_type for doc in documents.values()}
    assert content_types["Export.md"] == "text/markdown"
    assert content_types["Figures.xlsx"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------- When the tool is offered ----------------


async def test_the_tool_is_not_offered_when_switched_off(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "chat_file_tools", False)
    alice = await create_user("alice")

    await chat(alice, "Make a PDF", document_ids=[])

    [call] = fake_llm.chat_calls
    assert not {"tools", "tool_choice", "parallel_tool_calls"} & set(call)
    assert "create_document" not in system_prompt(call) and "App note" not in system_prompt(call)
    assert not any((await can_create_files(alice)).values())


async def test_models_without_function_calling_get_a_note_instead(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(litellm, "supports_function_calling", lambda model: False)
    alice = await create_user("alice")

    await chat(alice, "Make a PDF", document_ids=[])

    [call] = fake_llm.chat_calls
    assert "tools" not in call
    assert document_tool.NO_FILE_TOOLS_NOTE in system_prompt(call)
    assert not any((await can_create_files(alice)).values())


@pytest.mark.parametrize(
    ("limits", "reason"),
    [
        ({"max_documents": 1}, "they have reached their limit of 1 documents"),
        ({"max_storage_mb": 1}, "their storage is almost full (0.8 of 1 MB used)"),
    ],
)
async def test_users_without_room_for_a_file_are_told_why(
    create_user: CreateUser, fake_llm: FakeLLM, limits: dict[str, int], reason: str
) -> None:
    alice = await create_user("alice", **limits)
    assert (await upload(alice, "notes.txt", b"x" * 820_000)).status_code == 201

    await chat(alice, "Make a PDF", document_ids=[])

    [call] = fake_llm.chat_calls
    assert "tools" not in call
    assert f"The user can't get new files right now because {reason}." in system_prompt(call)
    assert (await can_create_files(alice))["Claude Opus 5.5"] is True  # quotas aren't part of the model list


async def test_the_admin_is_offered_the_tool_whatever_the_limits(admin_client: ApiClient, fake_llm: FakeLLM) -> None:
    await chat(admin_client, "Make a PDF", document_ids=[])

    assert fake_llm.chat_calls[0]["tools"][0]["function"]["name"] == "create_document"


async def test_a_rejected_tool_is_retried_without_it_and_remembered(
    create_user: CreateUser, fake_llm: FakeLLM, tool_mode: str
) -> None:
    alice = await create_user("alice")
    rejected = litellm.BadRequestError(message="tools.0: unsupported schema", model="x", llm_provider="anthropic")
    fake_llm.script = [ScriptedCall(error=rejected, before_output=True), ScriptedCall([("text", "Plain answer.")])]

    response, events = await chat(alice, "Hello?", document_ids=[])

    assert names(events)[-1] == "done"
    first, retry = fake_llm.chat_calls
    assert "tools" in first and "tools" not in retry
    # Only the tool changes, so only the tool is blamed.
    assert first["max_tokens"] == retry["max_tokens"] == settings.generated_file_max_output_tokens
    assert tool_names(first) == offered_tools(tool_mode)
    assert document_tool.NO_FILE_TOOLS_NOTE in system_prompt(retry)
    assert "create_document tool" not in system_prompt(retry) and "create_image tool" not in system_prompt(retry)
    assert (image_tool.NO_IMAGE_TOOLS_NOTE in system_prompt(retry)) is (tool_mode == "files+images")
    [usage] = await chat_usage(alice.user_id)
    assert (usage.prompt_tokens, usage.completion_tokens) == (100, 20)  # the rejected call isn't billed
    assert (await answer_of(events[0][1]["conversation_id"])).content == "Plain answer."
    assert (await can_create_files(alice))["Claude Opus 5.5"] is False
    images = {model["display_name"]: model["can_create_images"] for model in (await alice.get("/api/models")).json()}
    assert images["Claude Opus 5.5"] is False and images["Claude Sonnet 5.5"] is (tool_mode == "files+images")

    await chat(alice, "Again?", document_ids=[])
    assert "tools" not in fake_llm.chat_calls[-1]


def offered(call: dict[str, Any]) -> tuple[bool, int | None]:
    """Whether the call offered the tool, and its max_tokens."""
    return "tools" in call, call.get("max_tokens")


FILE_MAX = 48_000  # GENERATED_FILE_MAX_OUTPUT_TOKENS: the seeded models' output caps aren't in LiteLLM's map
USUAL_MAX = 16_000  # DEFAULT_MAX_OUTPUT_TOKENS


async def test_when_every_retry_fails_the_answer_ends_and_nothing_is_remembered(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")

    def rejected() -> Exception:
        return litellm.BadRequestError(message="bad schema", model="x", llm_provider="anthropic")

    fake_llm.script = [ScriptedCall(error=rejected(), before_output=True) for _ in range(4)]

    response, events = await chat(alice, "Hello?", document_ids=[])

    assert events[-1][0] == "error" and "bad schema" in events[-1][1]["message"]
    # Without the tool; then, as that failed too, with the tool and the usual max_tokens; then without the tool again.
    assert [offered(call) for call in fake_llm.chat_calls] == [
        (True, FILE_MAX),
        (False, FILE_MAX),
        (True, USUAL_MAX),
        (False, USUAL_MAX),
    ]
    assert (await can_create_files(alice))["Claude Opus 5.5"] is True
    assert await chat_usage(alice.user_id) == []  # rejected calls cost nothing


async def test_a_rejected_max_tokens_keeps_the_tool_and_retries_with_the_usual_max_tokens(
    create_user: CreateUser, fake_llm: FakeLLM, caplog: pytest.LogCaptureFixture, tool_mode: str
) -> None:
    alice = await create_user("alice")
    too_many = litellm.BadRequestError(
        message="max_tokens: 48000 > 32000, which is the maximum allowed number of output tokens for claude-x",
        model="x",
        llm_provider="anthropic",
    )
    fake_llm.script = [
        ScriptedCall(error=too_many, before_output=True),
        ScriptedCall([pdf_call(fmt="md")]),
        ScriptedCall([("text", "Done.")]),
    ]

    response, events = await chat(alice, "Make a markdown report", document_ids=[])

    assert events[-1][0] == "done"
    assert [offered(call) for call in fake_llm.chat_calls] == [(True, FILE_MAX), (True, USUAL_MAX), (True, USUAL_MAX)]
    rejected, retried, last = fake_llm.chat_calls
    images = tool_mode == "files+images"
    assert system_prompt(retried).count(document_tool.file_instructions(USUAL_MAX, images)) == 1
    assert document_tool.file_instructions(USUAL_MAX) != document_tool.file_instructions(FILE_MAX)  # fewer words
    assert tool_names(retried) == offered_tools(tool_mode)
    assert ("You can generate images with the create_image tool" in system_prompt(retried)) is images
    assert system_prompt(last) == system_prompt(retried)
    assert tool_messages(last)[0]["ok"] is True
    assert (await can_create_files(alice))["Claude Opus 5.5"] is True  # the tool wasn't to blame
    assert "Set the model's Max output tokens".lower() in caplog.text.lower()

    fake_llm.script = []
    await chat(alice, "Hello again", document_ids=[])
    assert offered(fake_llm.chat_calls[-1]) == (True, FILE_MAX)  # each answer starts with the configured budget


async def test_a_retry_without_the_tool_that_fails_too_is_retried_with_the_tool_and_the_usual_max_tokens(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    """A provider wording the size check doesn't recognize: dropping the tool doesn't help, lowering max_tokens does."""
    alice = await create_user("alice")

    def rejected() -> Exception:
        return litellm.BadRequestError(message="Invalid request: limit exceeded", model="x", llm_provider="anthropic")

    fake_llm.script = [
        ScriptedCall(error=rejected(), before_output=True),
        ScriptedCall(error=rejected(), before_output=True),
        ScriptedCall([("text", "Plain answer.")]),
    ]

    response, events = await chat(alice, "Hello?", document_ids=[])

    assert events[-1][0] == "done"
    assert [offered(call) for call in fake_llm.chat_calls] == [(True, FILE_MAX), (False, FILE_MAX), (True, USUAL_MAX)]
    assert (await can_create_files(alice))["Claude Opus 5.5"] is True


async def test_a_tool_rejected_at_the_usual_max_tokens_too_is_remembered(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    too_many = litellm.BadRequestError(message="max_tokens is too large", model="x", llm_provider="anthropic")
    schema = litellm.BadRequestError(message="tools.0: unsupported schema", model="x", llm_provider="anthropic")
    fake_llm.script = [
        ScriptedCall(error=too_many, before_output=True),
        ScriptedCall(error=schema, before_output=True),
        ScriptedCall([("text", "Plain answer.")]),
    ]

    response, events = await chat(alice, "Hello?", document_ids=[])

    assert events[-1][0] == "done"
    assert [offered(call) for call in fake_llm.chat_calls] == [(True, FILE_MAX), (True, USUAL_MAX), (False, USUAL_MAX)]
    assert (await can_create_files(alice))["Claude Opus 5.5"] is False


async def test_a_context_window_error_is_retried_once_with_the_usual_max_tokens(
    create_user: CreateUser, fake_llm: FakeLLM, tool_mode: str
) -> None:
    alice = await create_user("alice")

    def too_long() -> Exception:
        return litellm.ContextWindowExceededError(message="prompt is too long", model="x", llm_provider="anthropic")

    fake_llm.script = [ScriptedCall(error=too_long(), before_output=True), ScriptedCall([("text", "Fits now.")])]
    response, events = await chat(alice, "Hello?", document_ids=[])

    assert events[-1][0] == "done"
    assert [offered(call) for call in fake_llm.chat_calls] == [(True, FILE_MAX), (True, USUAL_MAX)]
    assert all(tool_names(call) == offered_tools(tool_mode) for call in fake_llm.chat_calls)
    assert (await can_create_files(alice))["Claude Opus 5.5"] is True

    fake_llm.chat_calls.clear()
    fake_llm.script = [ScriptedCall(error=too_long(), before_output=True) for _ in range(2)]
    response, events = await chat(alice, "Hello?", document_ids=[])

    assert events[-1][0] == "error" and len(fake_llm.chat_calls) == 2  # never retried without the tool
    assert (await can_create_files(alice))["Claude Opus 5.5"] is True


async def test_a_context_window_error_is_not_retried_when_max_tokens_cant_go_lower(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "generated_file_max_output_tokens", settings.default_max_output_tokens)
    alice = await create_user("alice")
    too_long = litellm.ContextWindowExceededError(message="prompt is too long", model="x", llm_provider="anthropic")
    fake_llm.script = [ScriptedCall(error=too_long, before_output=True)]

    response, events = await chat(alice, "Hello?", document_ids=[])

    assert events[-1][0] == "error" and len(fake_llm.chat_calls) == 1


async def test_a_rejected_thinking_signature_is_stripped_and_retried_once(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    signature = litellm.BadRequestError(
        message="messages.1.content.0: Invalid `signature` in `thinking` block", model="x", llm_provider="anthropic"
    )
    fake_llm.script = [
        ScriptedCall([("thinking", (["Plan."], "SIG-1")), pdf_call(fmt="md")]),
        ScriptedCall(error=signature, before_output=True),
        ScriptedCall([("text", "Done.")]),
    ]

    response, events = await chat(alice, "Export it", document_ids=[])

    assert events[-1][0] == "done"
    first, rejected, retried = fake_llm.chat_calls
    assert any("thinking_blocks" in message for message in rejected["messages"])
    assert not any("thinking_blocks" in message for message in retried["messages"])
    assert tool_messages(retried)[0]["ok"] is True


async def test_a_second_signature_error_ends_the_answer(create_user: CreateUser, fake_llm: FakeLLM) -> None:
    alice = await create_user("alice")

    def signature() -> Exception:
        return litellm.BadRequestError(message="Invalid `signature` in `thinking` block", model="x", llm_provider="x")

    fake_llm.script = [
        ScriptedCall([("thinking", (["Plan."], "SIG-1")), pdf_call(fmt="md")]),
        ScriptedCall(error=signature(), before_output=True),
        ScriptedCall(error=signature(), before_output=True),
    ]

    response, events = await chat(alice, "Export it", document_ids=[])

    assert events[-1][0] == "error" and len(fake_llm.chat_calls) == 3
    answer = await answer_of(events[0][1]["conversation_id"])
    assert len(await rows(MessageFile, MessageFile.message_id == answer.id)) == 1  # the file stays linked


# ---------------- Limits and validation ----------------


async def test_an_answer_creates_at_most_the_file_limit(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch, tool_mode: str
) -> None:
    monkeypatch.setattr(settings, "generated_files_per_answer", 2)
    alice = await create_user("alice")
    fake_llm.script = [ScriptedCall([pdf_call(f"File {n}", f"Body {n}", fmt="md")]) for n in range(5)]

    response, events = await chat(alice, "Make many files", document_ids=[])

    created = [data for data in file_events(events) if data["status"] == "created"]
    assert [data["file"]["filename"] for data in created] == ["File 0.md", "File 1.md"]
    if tool_mode == "files":
        assert [call.get("tool_choice") for call in fake_llm.chat_calls] == ["auto", "auto", "none"]
        assert len(fake_llm.script) == 2  # no more calls than the limit allowed
    else:  # create_image is still offered, so further files are attempted and each gets the limit error
        assert [call.get("tool_choice") for call in fake_llm.chat_calls] == ["auto"] * 6
        assert fake_llm.script == [] and len(await rows(Document)) == 2
        assert {result["error"] for result in tool_messages(fake_llm.chat_calls[-1])[2:]} == {
            "Only 2 files can be created per answer; this one wasn't created. Tell the user."
        }


async def test_extra_calls_in_one_response_get_the_limit_error(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "generated_files_per_answer", 1)
    alice = await create_user("alice")
    fake_llm.script = [
        ScriptedCall([pdf_call("A", "a", fmt="md"), pdf_call("B", "b", fmt="md")]),
        ScriptedCall([("text", "Only one.")]),
    ]

    response, events = await chat(alice, "Two files", document_ids=[])

    statuses = [(data["call_id"], data["status"]) for data in file_events(events) if data["status"] != "writing"]
    assert statuses == [("f1", "created"), ("f2", "failed")]
    results = tool_messages(fake_llm.chat_calls[1])
    assert results[1] == {
        "ok": False,
        "error": "Only 1 files can be created per answer; this one wasn't created. Tell the user.",
    }


async def test_the_attempt_limit_counts_failed_calls(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "generated_files_per_answer", 1)
    alice = await create_user("alice")
    fake_llm.script = [ScriptedCall([tool_call(title="x", format="pdf")]) for _ in range(3)]  # no content: errors

    response, events = await chat(alice, "Try", document_ids=[])

    calls = fake_llm.chat_calls
    assert [call.get("tool_choice") for call in calls] == ["auto", "auto", "none"]
    assert [result["error"] for result in tool_messages(calls[-1])] == ["content is required for pdf, docx and md"] * 2
    assert all(data["status"] != "created" for data in file_events(events))


@pytest.mark.parametrize(
    ("block", "error"),
    [
        (
            ("tool", ("create_document", '{"title": "x", "format": "pdf", "content": "a"b')),
            "The arguments weren't valid JSON.",
        ),
        (tool_call(title="x", format="pdf"), "content is required for pdf, docx and md"),
        (tool_call(title="x", format="pptx", content="a"), "format must be one of pdf, docx, xlsx, md"),
        (tool_call(title="x", format="xlsx", content="notes"), "sheets is required for xlsx"),
        (tool_call(title=" ", format="md", content="a"), "title is required"),
        (tool_call("search_web", query="x"), "Unknown tool: search_web. The only tool is create_document."),
    ],
    ids=["invalid json", "no content", "unknown format", "xlsx without sheets", "no title", "unknown tool"],
)
async def test_invalid_calls_become_tool_errors_and_the_answer_goes_on(
    create_user: CreateUser, fake_llm: FakeLLM, block: tuple[str, Any], error: str
) -> None:
    alice = await create_user("alice")
    fake_llm.script = [ScriptedCall([block]), ScriptedCall([("text", "Sorry about that.")])]

    response, events = await chat(alice, "Make a file", document_ids=[])

    assert events[-1][0] == "done"
    [result] = tool_messages(fake_llm.chat_calls[1])
    assert result == {"ok": False, "error": error}
    failed = [data for data in file_events(events) if data["status"] == "failed"]
    if block[1][0] == "create_document":
        assert [data["error"] for data in failed] == [error]
    else:
        assert file_events(events) == []  # another tool's call is never announced
    echoed = fake_llm.chat_calls[1]["messages"][-2]["tool_calls"][0]["function"]["arguments"]
    if error == "The arguments weren't valid JSON.":
        assert echoed == "{}"
    assert await rows(Document) == []


async def test_content_over_the_limit_is_a_tool_error(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "generated_file_max_chars", 1000)
    alice = await create_user("alice")
    fake_llm.script = [ScriptedCall([pdf_call(content="x" * 1500)]), ScriptedCall([("text", "Too long.")])]

    await chat(alice, "Make a file", document_ids=[])

    [result] = tool_messages(fake_llm.chat_calls[1])
    assert result["error"] == (
        "The content is 1,500 characters; the limit is 1,000. Make the document shorter or split it into several files."
    )


# ---------------- Finish reasons ----------------


async def test_a_call_cut_off_by_the_output_limit_is_a_tool_error_and_the_loop_goes_on(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    cut = ("tool", ("create_document", '{"title": "Long report", "format": "pdf", "content": "# Long\\n\\nMany words'))
    fake_llm.script = [
        ScriptedCall([cut], finish_reason="length"),
        ScriptedCall([pdf_call("Short report", "# Short\n\nFewer words.", fmt="md")]),
        ScriptedCall([("text", "Here is a shorter one.")]),
    ]

    response, events = await chat(alice, "A long report", document_ids=[])

    statuses = [(data["call_id"], data["status"]) for data in file_events(events) if data["status"] != "writing"]
    assert statuses == [("f1", "failed"), ("f2", "created")]
    failed = file_events(events)[1]
    assert failed["error"].startswith("The file was cut off at the model's output limit (about 16,000 words per file)")
    second = fake_llm.chat_calls[1]
    assert second["messages"][-2]["tool_calls"][0]["function"]["arguments"] == "{}"
    assert tool_messages(second)[0]["error"] == failed["error"]
    assert [doc.filename for doc in await rows(Document)] == ["Short report.md"]


@pytest.mark.parametrize(("finish", "error"), [("content_filter", DECLINED), ("", INCOMPLETE)])
async def test_a_refused_or_unfinished_response_ends_the_answer(
    create_user: CreateUser, fake_llm: FakeLLM, finish: str, error: str
) -> None:
    alice = await create_user("alice")
    partial = ("tool", ("create_document", '{"title": "Report", "format": "pdf", "content": "# Rep'))
    fake_llm.script = [ScriptedCall([("text", "Writing it."), partial], finish_reason=finish)]

    response, events = await chat(alice, "A report", document_ids=[])

    assert len(fake_llm.chat_calls) == 1
    assert [data["status"] for data in file_events(events)] == ["writing", "failed"]
    assert file_events(events)[1]["error"] == error
    answer = await answer_of(events[0][1]["conversation_id"])
    assert events[-1] == ("error", {"message": error, "assistant_message_id": answer.id})
    assert (answer.content, answer.error) == ("Writing it.", error)
    assert await rows(Document) == []


# ---------------- Quotas and duplicates ----------------


async def test_the_document_limit_at_store_time_is_a_tool_error(create_user: CreateUser, fake_llm: FakeLLM) -> None:
    alice = await create_user("alice", max_documents=1)
    fake_llm.script = [
        ScriptedCall([pdf_call("A", "first", fmt="md")]),
        ScriptedCall([pdf_call("B", "second", fmt="md")]),
        ScriptedCall([("text", "One of two.")]),
    ]

    await chat(alice, "Two files", document_ids=[])

    results = tool_messages(fake_llm.chat_calls[2])
    assert results[1] == {
        "ok": False,
        "error": "The user has reached their limit of 1 documents; the file wasn't created. They can delete documents "
        "to make room.",
    }
    [doc] = await rows(Document)
    assert stored_files() == [doc.stored_name]


async def test_storage_and_file_size_limits_are_tool_errors(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "generated_file_max_chars", 1_000_000)
    alice = await create_user("alice", max_storage_mb=2, max_upload_mb=1)
    devanagari = "क" * 400_000  # 1.2 MB of UTF-8
    big = json.dumps({"title": "Big", "format": "md", "content": devanagari}, ensure_ascii=False)
    fake_llm.script = [
        ScriptedCall([("tool", ("create_document", big))], fragment=100_000),
        ScriptedCall([("text", "Too big.")]),
    ]
    for name, data in (("a.txt", b"x" * 900_000), ("b.txt", b"y" * 800_000)):  # 1.7 of 2 MB used
        assert (await upload(alice, name, data)).status_code == 201

    await chat(alice, "A big file", document_ids=[])

    [result] = tool_messages(fake_llm.chat_calls[1])
    assert result["error"] == (
        "The file is 1.1 MB, over the user's limit of 1 MB per file; it wasn't created. Make it shorter."
    )
    fake_llm.script = [
        ScriptedCall([pdf_call("Medium", "क" * 140_000, fmt="md")]),
        ScriptedCall([("text", "No room.")]),
    ]
    await chat(alice, "A medium file", document_ids=[])
    [result] = tool_messages(fake_llm.chat_calls[-1])
    assert result["error"] == (
        "This file would exceed the user's storage limit of 2 MB; it wasn't created. They can delete documents to "
        "free up space."
    )
    assert len(stored_files()) == 2  # only the uploads


async def test_the_admin_has_no_quotas_and_identical_files_are_linked_once(
    admin_client: ApiClient, fake_llm: FakeLLM
) -> None:
    fake_llm.script = [
        ScriptedCall([pdf_call("Notes", "# Notes\n\nSame text.", fmt="md")]),
        ScriptedCall([("text", "Saved.")]),
        ScriptedCall([pdf_call("Notes again", "# Notes\n\nSame text.", fmt="md")]),
        ScriptedCall([("text", "Already there.")]),
    ]

    await chat(admin_client, "Save notes", document_ids=[])
    response, events = await chat(admin_client, "Save the same notes", document_ids=[])

    [created] = [data for data in file_events(events) if data["status"] == "created"]
    assert created["already_existed"] is True and created["file"]["filename"] == "Notes.md"
    result = tool_messages(fake_llm.chat_calls[-1])[0]
    assert result["already_existed"] is True and result["filename"] == "Notes.md"
    assert result["note"].startswith("An identical file was already in the user's private documents")
    [doc] = await rows(Document)
    assert stored_files() == [doc.stored_name]
    answer = await answer_of(events[0][1]["conversation_id"])
    [linked] = await rows(MessageFile, MessageFile.message_id == answer.id)
    assert (linked.document_id, linked.filename) == (doc.id, "Notes.md")


# ---------------- Errors, usage and budget ----------------


async def test_a_provider_error_after_a_file_keeps_the_file_and_bills_both_calls(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    fake_llm.script = [
        ScriptedCall([pdf_call("First", "one", fmt="md")], usage=(1000, 300)),
        ScriptedCall(
            [("text", "Now the second"), pdf_call("Second", "two", fmt="md")],
            error=litellm.APIConnectionError(message="connection reset", model="x", llm_provider="anthropic"),
            finish_reason="",
            usage=(1200, 100),
        ),
    ]

    response, events = await chat(alice, "Two files", document_ids=[])

    assert events[-1][0] == "error" and events[-1][1]["message"] == "Could not connect to the model provider."
    second = [data for data in file_events(events) if data["call_id"] == "f2"]
    assert [data["status"] for data in second] == ["writing", "failed"] and second[-1]["error"] == CALL_FAILED
    answer = await answer_of(events[0][1]["conversation_id"])
    assert [row.filename for row in await rows(MessageFile, MessageFile.message_id == answer.id)] == ["First.md"]
    [usage] = await chat_usage(alice.user_id)
    assert (usage.prompt_tokens, usage.completion_tokens) == (2200, 400)


async def test_a_tool_only_call_is_billed_when_the_next_call_fails_before_output(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    fake_llm.script = [
        ScriptedCall([pdf_call("First", "one", fmt="md")], usage=(1000, 9000)),
        ScriptedCall(
            error=litellm.InternalServerError(message="overloaded", model="x", llm_provider="anthropic"),
            before_output=True,
        ),
    ]

    response, events = await chat(alice, "A file", document_ids=[])

    assert events[-1][0] == "error"
    [usage] = await chat_usage(alice.user_id)
    assert (usage.prompt_tokens, usage.completion_tokens) == (1000, 9000)
    answer = await answer_of(events[0][1]["conversation_id"])
    assert (answer.prompt_tokens, answer.completion_tokens) == (1000, 9000)


async def test_no_follow_up_call_starts_once_the_budget_is_spent(create_user: CreateUser, fake_llm: FakeLLM) -> None:
    alice = await create_user("alice", monthly_token_limit=5000)
    fake_llm.script = [ScriptedCall([pdf_call("Report", "body", fmt="md")], usage=(4000, 1500))]

    response, events = await chat(alice, "A file", document_ids=[])

    assert len(fake_llm.chat_calls) == 1
    answer = await answer_of(events[0][1]["conversation_id"])
    assert events[-1] == ("error", {"message": BUDGET_SPENT, "assistant_message_id": answer.id})
    assert [row.filename for row in await rows(MessageFile, MessageFile.message_id == answer.id)] == ["Report.md"]
    [usage] = await chat_usage(alice.user_id)
    assert usage.total_tokens == 5500


async def test_refreshing_a_slot_never_recreates_a_released_one(create_user: CreateUser) -> None:
    alice = await create_user("alice", monthly_token_limit=100_000)
    async with SessionLocal() as session:
        from app.models import User

        user = await session.get(User, alice.user_id)
        assert user is not None
        ent = await entitlements.load(session, user)
    redis = get_redis()
    slot = await usage_service.acquire_chat_slot(user, ent)
    assert slot is not None
    key = f"chat:inflight:{user.id}"
    await redis.zadd(key, {slot: 1.0})

    await usage_service.refresh_chat_slot(user.id, slot)
    assert await redis.zscore(key, slot) > time.time() - 60
    await usage_service.release_chat_slot(user.id, slot)
    await usage_service.refresh_chat_slot(user.id, slot)
    assert await redis.zscore(key, slot) is None


# ---------------- Progress and keepalive ----------------


async def test_progress_and_keepalives_while_a_file_is_written(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch, tool_mode: str
) -> None:
    monkeypatch.setattr(document_tool, "_PROGRESS_SECONDS", 0.05)
    monkeypatch.setattr(chat_router, "_KEEPALIVE_SECONDS", 0.05)
    real_render = document_tool._render

    async def slow_render(*args: Any) -> Any:
        # A warm render child can finish a small md file within the keepalive interval, so make it slower.
        await asyncio.sleep(0.3)
        return await real_render(*args)

    monkeypatch.setattr(document_tool, "_render", slow_render)
    alice = await create_user("alice")
    fake_llm.chunk_delay = 0.02
    content = "# Report\n\n" + "Revenue grew in every region this quarter. " * 30
    fake_llm.script = [ScriptedCall([pdf_call("Slow report", content, fmt="md")], fragment=60), ScriptedCall([])]

    response, events = await chat(alice, "A report", document_ids=[])

    text = response.text
    rendering = text[text.rindex('"status": "writing"') : text.index('"status": "created"')]
    assert ": keepalive\n\n" in rendering  # the render child takes longer than the keepalive interval
    writing = [data for data in file_events(events) if data["status"] == "writing"]
    progress = writing[1:-1]
    assert len(progress) >= 3
    chars = [data["chars"] for data in progress]
    assert chars == sorted(chars) and chars[-1] > chars[0]
    assert all(data["title"] == "Slow report" and data["format"] == "md" for data in progress[1:])
    assert writing[-1]["chars"] >= chars[-1]
    assert events[-1][0] == "done"


async def test_keepalives_while_the_model_is_silent(
    create_user: CreateUser, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(chat_router, "_KEEPALIVE_SECONDS", 0.05)
    alice = await create_user("alice")
    fake_llm.chunks = ["Thinking", " done."]
    fake_llm.chunk_delay = 0.3

    response, events = await chat(alice, "Hello?", document_ids=[])

    between = response.text[response.text.index('"Thinking"') : response.text.index('" done."')]
    assert between.count(": keepalive\n\n") >= 2
    assert names(events) == ["meta", "sources", "delta", "delta", "done"]


def test_the_title_in_progress_appears_once_its_quote_closed() -> None:
    session = document_tool.FileSession(None, [], None, ent=None)  # type: ignore[arg-type]
    session.started(0)
    step = llm.ChatStep(tool_calls=[llm.ToolCall(0, "t", "create_document", "")])
    for arguments, title, file_format in [
        ('{"title": "Q3 rep', None, None),
        ('{"title": "Q3 report\\u00e9", "format": "Word"', "Q3 reporté", "docx"),
    ]:
        session._calls[0].reported = 0.0
        step.tool_calls[0].arguments = arguments
        event = session.progress(0, step)
        assert event is not None and (event["title"], event["format"]) == (title, file_format)
    assert session.progress(0, step) is None  # throttled


# ---------------- Stop ----------------


async def test_stop_during_rendering_kills_the_child_and_creates_nothing(
    create_user: CreateUser, start_turn: StartTurn, fake_llm: FakeLLM, file_child: Callable[..., FileChild]
) -> None:
    alice = await create_user("alice")
    child = file_child("wait")
    fake_llm.script = [ScriptedCall([pdf_call()])]
    turn = await start_turn(alice, "A report", document_ids=[])
    stream = _stream_answer(turn)
    received = []
    while '"title": "Q3 report"' not in "".join(received):
        received.append(await anext(stream))

    rendering = asyncio.create_task(anext(stream))
    deadline = time.monotonic() + 10
    while not child.pids:
        assert time.monotonic() < deadline
        await asyncio.sleep(0.02)
    rendering.cancel()
    with pytest.raises(asyncio.CancelledError):
        await rendering

    [pid] = child.pids
    assert not _alive(pid)
    answer = await answer_of(turn.conversation_id)
    assert answer.error == STOPPED
    assert await rows(MessageFile) == [] and await rows(Document) == []
    assert stored_files() == []


async def test_stop_while_storing_links_the_document_exactly_once(
    create_user: CreateUser, start_turn: StartTurn, fake_llm: FakeLLM, monkeypatch: pytest.MonkeyPatch, tool_mode: str
) -> None:
    alice = await create_user("alice")
    gate, reached = asyncio.Event(), asyncio.Event()
    real_save = document_tool._save

    async def gated_save(*args: Any) -> Any:
        reached.set()
        await gate.wait()
        return await real_save(*args)

    monkeypatch.setattr(document_tool, "_save", gated_save)
    fake_llm.script = [ScriptedCall([pdf_call(fmt="md")])]
    turn = await start_turn(alice, "A file", document_ids=[])
    stream = _stream_answer(turn)
    received = []
    while '"title": "Q3 report"' not in "".join(received):
        received.append(await anext(stream))

    storing = asyncio.create_task(anext(stream))
    await asyncio.wait_for(reached.wait(), 30)
    storing.cancel()
    gate.set()
    with pytest.raises(asyncio.CancelledError):
        await storing
    await asyncio.gather(*list(chat_router._pending_saves))

    answer = await answer_of(turn.conversation_id)
    assert answer.error == STOPPED
    [doc] = await rows(Document)
    [linked] = await rows(MessageFile, MessageFile.message_id == answer.id)
    assert linked.document_id == doc.id


async def test_stop_during_the_follow_up_call_keeps_the_first_file(
    create_user: CreateUser, start_turn: StartTurn, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    fake_llm.script = [ScriptedCall([pdf_call(fmt="md")])]
    fake_llm.stall_after = 1  # the follow-up call says "Hello", then waits
    turn = await start_turn(alice, "A file", document_ids=[])
    stream = _stream_answer(turn)
    received = []
    while '"content": "Hello"' not in "".join(received):
        received.append(await anext(stream))

    await stream.aclose()

    answer = await answer_of(turn.conversation_id)
    assert (answer.content, answer.error) == ("Hello", STOPPED)
    assert [row.filename for row in await rows(MessageFile, MessageFile.message_id == answer.id)] == ["Q3 report.md"]
    [usage] = await chat_usage(alice.user_id)
    assert usage.prompt_tokens > 100  # both calls billed


# ---------------- Deletes, provenance and history ----------------


async def test_a_conversation_deleted_mid_stream_leaves_its_files_in_the_documents(
    create_user: CreateUser, start_turn: StartTurn, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    fake_llm.script = [ScriptedCall([pdf_call(fmt="md")]), ScriptedCall([("text", "Done.")])]
    turn = await start_turn(alice, "A file", document_ids=[])
    stream = _stream_answer(turn)
    received = []
    while '"status": "created"' not in "".join(received):
        received.append(await anext(stream))

    assert (await alice.delete(f"/api/conversations/{turn.conversation_id}")).status_code == 204
    received += [event async for event in stream]

    assert parse(received)[-1][0] == "error"
    [doc] = await rows(Document)
    assert await rows(MessageFile) == [] and len(await rows(GeneratedDocument)) == 1
    assert next(item for item in (await alice.get("/api/documents")).json() if item["id"] == doc.id)["generated"]


async def test_deleting_the_document_or_the_conversation_keeps_the_rest(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    fake_llm.script = [
        ScriptedCall([pdf_call("Kept", "kept", fmt="md")]),
        ScriptedCall([pdf_call("Removed", "removed")]),
        ScriptedCall([("text", "Two files.")]),
    ]
    response, events = await chat(alice, "Two files", document_ids=[])
    conversation_id = events[0][1]["conversation_id"]
    kept, removed = (data["file"] for data in file_events(events) if data["status"] == "created")
    [removed_doc] = await rows(Document, Document.id == removed["document_id"])
    source = storage.stored_path(storage.source_name(removed_doc.stored_name))
    assert source.exists()

    assert (await alice.delete(f"/api/documents/{removed['document_id']}")).status_code == 204

    assert not source.exists() and not storage.stored_path(removed_doc.stored_name).exists()
    files = (await alice.get(f"/api/conversations/{conversation_id}")).json()["messages"][1]["files"]
    assert files == [
        {**kept, "status": "pending"},
        {**removed, "document_id": None, "status": None},
    ]
    assert [row.document_id for row in await rows(GeneratedDocument)] == [kept["document_id"]]

    assert (await alice.delete(f"/api/conversations/{conversation_id}")).status_code == 204
    assert await rows(MessageFile) == []
    assert [row.document_id for row in await rows(GeneratedDocument)] == [kept["document_id"]]


async def test_the_next_turn_knows_the_files_an_answer_created(
    create_user: CreateUser, fake_llm: FakeLLM, tool_mode: str
) -> None:
    alice = await create_user("alice")
    fake_llm.script = [
        ScriptedCall([pdf_call("Q3 report"), pdf_call("Figures", "| a |\n|---|\n| 1 |", fmt="md")]),
        ScriptedCall([("text", "")]),
    ]
    response, events = await chat(alice, "A PDF and a markdown table", document_ids=[])

    await chat(
        alice, "Now make it a Word document too", conversation_id=events[0][1]["conversation_id"], document_ids=[]
    )

    history = fake_llm.chat_calls[-1]["messages"][1:-1]
    assert history[-1] == {
        "role": "assistant",
        "content": '(App note: this answer created these files, which are in the user\'s documents: "Q3 report.pdf", '
        '"Figures.md".)',
    }
    assert system_prompt(fake_llm.chat_calls[-1]).count('Notes in parentheses that start with "App note:"') == 1
    assert tool_names(fake_llm.chat_calls[-1]) == offered_tools(tool_mode)


async def test_the_history_note_says_which_files_were_deleted_since(create_user: CreateUser, fake_llm: FakeLLM) -> None:
    alice = await create_user("alice")
    fake_llm.script = [
        ScriptedCall([pdf_call("Q3 report"), pdf_call("Figures", "| a |\n|---|\n| 1 |", fmt="md")]),
        ScriptedCall([("text", "")]),
    ]
    response, events = await chat(alice, "A PDF and a markdown table", document_ids=[])
    conversation_id = events[0][1]["conversation_id"]
    report, figures = [data["file"]["document_id"] for data in file_events(events) if data["status"] == "created"]

    assert (await alice.delete(f"/api/documents/{report}")).status_code == 204
    await chat(alice, "Where is the report you made?", conversation_id=conversation_id, document_ids=[])

    history = fake_llm.chat_calls[-1]["messages"][1:-1]
    assert history[-1]["content"] == (
        '(App note: this answer created these files, which are in the user\'s documents: "Figures.md". It also '
        'created these files, which have since been deleted from the user\'s documents: "Q3 report.pdf".)'
    )

    assert (await alice.delete(f"/api/documents/{figures}")).status_code == 204
    await chat(alice, "And the table?", conversation_id=conversation_id, document_ids=[])

    history = fake_llm.chat_calls[-1]["messages"][1:-1]
    assert history[1]["content"] == (
        "(App note: this answer created these files, which have since been deleted from the user's documents: "
        '"Q3 report.pdf", "Figures.md".)'
    )


async def test_a_docx_whose_content_has_characters_xml_cant_hold_is_created(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    content = "# Memo\n\nA soft\x0bbreak, \x1b[1man escape\x1b[0m and U+FFFF \uffff."
    fake_llm.script = [
        ScriptedCall([pdf_call("Memo\uffff", content, fmt="docx")]),
        ScriptedCall([("text", "Done.")]),
    ]

    response, events = await chat(alice, "A Word memo", document_ids=[])

    assert events[-1][0] == "done"
    [result] = tool_messages(fake_llm.chat_calls[-1])
    assert (result["ok"], result["filename"]) == (True, "Memo.docx")
    [doc] = await rows(Document, Document.owner_id == alice.user_id)
    texts = [paragraph.text for paragraph in docx.Document(storage.stored_path(doc.stored_name)).paragraphs]
    assert texts == ["Memo", "A soft break, [1man escape[0m and U+FFFF ."]


def test_runs_of_citation_markers_in_generated_passages_are_neutralized() -> None:
    from datetime import datetime, timezone

    from app.schemas import Source

    def source(index: int, document_id: int) -> Source:
        return Source(
            index=index, document_id=document_id, filename=f"d{document_id}.docx", page=None, chunk_index=0, score=1.0,
            snippet="",
        )  # fmt: skip

    generated_text = (
        "Refunds take 14 days [1][3]. Shipping is free [2], [4]. All agree [1][2][3]. Code m[0][1] and "
        "[label][1] and a link [5](https://example.com) and escaped \\[6\\]\\[7\\]."
    )
    block = rag._context_block(
        [source(1, 10), source(2, 11)],
        [generated_text, "An upload: matrix[1][2] and see [1][3]."],
        {10: datetime(2026, 10, 15, tzinfo=timezone.utc)},
    )

    generated, upload = block.removeprefix("<context>\n").removesuffix("\n</context>").split("\n\n")
    assert generated.endswith(
        "Refunds take 14 days (ref 1)(ref 3). Shipping is free (ref 2), (ref 4). All agree (ref 1)(ref 2)(ref 3). "
        "Code m[0][1] and [label][1] and a link [5](https://example.com) and escaped (ref 6)(ref 7)."
    )
    assert upload.endswith("An upload: matrix[1][2] and see [1][3].")  # only generated files are rewritten


async def test_sources_list_the_cited_passages_and_unknown_numbers_are_noted(
    create_user: CreateUser, fake_llm: FakeLLM
) -> None:
    alice = await create_user("alice")
    first = await upload(alice, "orchids.txt", b"The orchid greenhouse is kept at 24 degrees all year.")
    await process(first.json()["id"])
    fake_llm.script = [
        ScriptedCall([pdf_call("Orchids", "Kept at 24 degrees [1].", fmt="md", source_numbers=[1, 7, 1])]),
        ScriptedCall([("text", "Done.")]),
    ]

    response, events = await chat(alice, "Export what you know about the orchid greenhouse")

    [result] = tool_messages(fake_llm.chat_calls[1])
    assert result["note"].endswith("These source numbers match no passage and were left out of the Sources: [7].")
    [doc] = await rows(Document, Document.filename == "Orchids.md")
    text = storage.stored_path(doc.stored_name).read_text()
    assert text.endswith("## Sources\n\n- \\[1\\] orchids.txt\n")


def _alive(pid: int) -> bool:
    import os

    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    return True
