"""Unit tests for parse_file: every format, text normalisation, and the safety limits."""

import io
import json
import time
import zipfile
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from conftest import docx_bytes, html_bytes, json_bytes, pdf_bytes, pptx_bytes, sample_file, txt_bytes, xlsx_bytes
from docx import Document as new_docx
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PageObject, PdfReader, PdfWriter
from pypdf.errors import DependencyError

from app.config import settings
from app.services import parsing
from app.services.parsing import (
    LEGACY_EXTENSIONS,
    SUPPORTED_EXTENSIONS,
    DocumentParseError,
    ParsedDocument,
    Segment,
    parse_file,
)

NO_PDF_TEXT = (
    "No extractable text: this PDF looks scanned or image-only. Ask the admin to enable a vision model to read it."
)
NO_TEXT = "The document contains no text."
TOO_COMPLEX = "The file is too large or complex to process."
UNSUPPORTED_ENCRYPTION = "This PDF uses encryption that isn't supported. Save an unencrypted copy and upload it again."

WriteFile = Callable[[str, bytes], Path]


@pytest.fixture
def write(isolated_state: Path) -> WriteFile:
    """Write bytes to a file in the test's scratch directory and return its path."""

    def writer(name: str, data: bytes) -> Path:
        path = isolated_state / name
        path.write_bytes(data)
        return path

    return writer


def parse(write: WriteFile, name: str, data: bytes) -> ParsedDocument:
    path = write(name, data)
    return parse_file(path, path.suffix)


def parse_error(write: WriteFile, name: str, data: bytes) -> str:
    with pytest.raises(DocumentParseError) as caught:
        parse(write, name, data)
    return str(caught.value)


def encrypted_pdf(text: str, *, user_password: str, algorithm: str = "RC4-128") -> bytes:
    writer = PdfWriter(clone_from=io.BytesIO(pdf_bytes(text)))
    writer.encrypt(user_password=user_password, owner_password="owner-secret", algorithm=algorithm)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def padded_zip(data: bytes, padding_bytes: int, name: str = "customXml/padding.bin") -> bytes:
    """`data` (an Office file) plus a highly compressible entry that inflates to `padding_bytes`."""
    buffer = io.BytesIO(data)
    with zipfile.ZipFile(buffer, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, b"\0" * padding_bytes)
    return buffer.getvalue()


def repeated_page_pdf(text: str, copies: int, *, own_resources: bool = False) -> bytes:
    """A PDF whose page tree lists one page object `copies` times (or `copies` page objects drawing one shared
    content stream, each with resources of its own)."""
    stream = f"BT /F1 12 Tf 72 760 Td ({text}) Tj ET".encode()
    page = "<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 3 0 R >> >> /Contents 4 0 R >>"
    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        4: b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
    }
    pages = range(5, 5 + copies) if own_resources else [5] * copies
    for page_id in sorted(set(pages)):
        objects[page_id] = page.encode()
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(f'{page} 0 R' for page in pages)}] /Count {copies} >>".encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for object_id in sorted(objects):
        offsets[object_id] = len(out)
        out += b"%d 0 obj\n%s\nendobj\n" % (object_id, objects[object_id])
    xref, size = len(out), max(objects) + 1
    out += b"xref\n0 %d\n0000000000 65535 f \n" % size
    out += b"".join(b"%010d 00000 n \n" % offsets[object_id] for object_id in range(1, size))
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref)
    return bytes(out)


def count_text_extractions(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    calls: list[int] = []
    extract = PageObject.extract_text

    def counting(self: PageObject, *args: object, **kwargs: object) -> str:
        calls.append(1)
        return extract(self, *args, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(PageObject, "extract_text", counting)
    return calls


def test_supported_and_legacy_extensions_match_the_spec() -> None:
    assert sorted(SUPPORTED_EXTENSIONS) == sorted(
        [".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".md", ".markdown", ".txt", ".csv", ".json"]
    )
    assert LEGACY_EXTENSIONS == {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}


@pytest.mark.parametrize("extension", SUPPORTED_EXTENSIONS)
def test_every_supported_format_yields_its_text(write: WriteFile, extension: str) -> None:
    parsed = parse(write, f"sample{extension}", sample_file(extension, "The heron nests by the quiet lake."))

    assert "The heron nests by the quiet lake." in "\n".join(segment.text for segment in parsed.segments)
    assert parsed.num_pages == (1 if extension in (".pdf", ".pptx") else None)


# ---------------- PDF ----------------


def test_pdf_has_one_segment_per_page_and_skips_blank_pages(write: WriteFile) -> None:
    parsed = parse(write, "report.pdf", pdf_bytes("First page line one\nline two (with parens)", "", "Third page"))

    assert parsed == ParsedDocument(
        segments=[Segment("First page line one\nline two (with parens)", 1), Segment("Third page", 3)],
        num_pages=3,
    )


def test_scanned_pdf_without_a_text_layer_needs_a_vision_model(write: WriteFile) -> None:
    assert parse_error(write, "scan.pdf", pdf_bytes("", "")) == NO_PDF_TEXT


def test_password_protected_pdf_is_refused(write: WriteFile) -> None:
    data = encrypted_pdf("Top secret quarterly figures.", user_password="open-sesame")

    assert parse_error(write, "locked.pdf", data) == "Password-protected PDFs aren't supported."


@pytest.mark.parametrize("algorithm", ["RC4-128", "AES-128", "AES-256"])
def test_pdf_restricted_only_by_an_owner_password_is_read(write: WriteFile, algorithm: str) -> None:
    data = encrypted_pdf("Permissions restricted but readable.", user_password="", algorithm=algorithm)

    assert parse(write, "restricted.pdf", data).segments == [Segment("Permissions restricted but readable.", 1)]


def test_aes_pdf_with_a_user_password_is_refused(write: WriteFile) -> None:
    data = encrypted_pdf("Top secret quarterly figures.", user_password="open-sesame", algorithm="AES-256")

    assert parse_error(write, "locked.pdf", data) == "Password-protected PDFs aren't supported."


@pytest.mark.parametrize("failing", ["open", "page"])
def test_missing_crypto_support_is_never_reported_as_corrupt_or_scanned(
    write: WriteFile, monkeypatch: pytest.MonkeyPatch, failing: str
) -> None:
    def unsupported(*args: object, **kwargs: object) -> None:
        raise DependencyError("cryptography>=3.1 is required for AES algorithm")

    if failing == "open":  # AES-256 checks the empty password while opening the file
        monkeypatch.setattr(parsing, "PdfReader", unsupported)
    else:  # AES-128 decrypts each page's stream while extracting it
        monkeypatch.setattr(PageObject, "extract_text", unsupported)

    assert parse_error(write, "aes.pdf", pdf_bytes("AES protected text.")) == UNSUPPORTED_ENCRYPTION


def test_pdf_with_unsupported_encryption_explains_what_to_do(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    def decrypt(self: PdfReader, password: str) -> None:
        raise DependencyError("cryptography>=3.1 is required for AES algorithm")

    monkeypatch.setattr(PdfReader, "decrypt", decrypt)
    data = encrypted_pdf("AES protected text.", user_password="")

    assert parse_error(write, "aes.pdf", data) == UNSUPPORTED_ENCRYPTION


def test_corrupt_pdf_is_reported_as_invalid(write: WriteFile) -> None:
    assert parse_error(write, "broken.pdf", b"%PDF-1.4\nthis is not really a pdf") == (
        "The file is corrupted or not a valid .pdf file."
    )


def test_pdf_with_too_many_pages_is_refused(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "max_pdf_pages", 2)

    assert parse(write, "two.pdf", pdf_bytes("One", "Two")).num_pages == 2
    assert parse_error(write, "three.pdf", pdf_bytes("One", "Two", "Three")) == (
        "The PDF has more than 2 pages, which isn't supported."
    )


def test_pages_drawing_the_same_content_are_read_once(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    calls = count_text_extractions(monkeypatch)

    parsed = parse(write, "repeated.pdf", repeated_page_pdf("Same page again", 50))

    assert parsed == ParsedDocument(segments=[Segment("Same page again", n) for n in range(1, 51)], num_pages=50)
    assert len(calls) == 1


def test_pages_with_resources_of_their_own_are_each_read(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    # The same stream can draw other text with other fonts or forms, so it is only reused with the same resources.
    calls = count_text_extractions(monkeypatch)

    parsed = parse(write, "shared-stream.pdf", repeated_page_pdf("Shared stream", 3, own_resources=True))

    assert [segment.text for segment in parsed.segments] == ["Shared stream"] * 3
    assert len(calls) == 3


def test_parsing_stops_at_the_deadline(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "parse_timeout_seconds", 60)
    clock = iter(range(0, 10_000, 45))  # every reading is 45 s later
    monkeypatch.setattr(parsing, "time", SimpleNamespace(monotonic=lambda: next(clock)))

    error = parse_error(write, "slow.pdf", pdf_bytes("One", "Two", "Three", "Four"))

    assert error == "The file is too large or complex to process (it took over 60 seconds)."


def test_a_page_that_cant_be_read_is_skipped_and_the_rest_kept(
    write: WriteFile, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    extract = PageObject.extract_text

    def flaky(page: PageObject, *args: object, **kwargs: object) -> str:
        text = extract(page, *args, **kwargs)
        if "Two" in text:
            raise KeyError("/Font")  # e.g. a page whose resources are broken
        return text

    monkeypatch.setattr(PageObject, "extract_text", flaky)

    parsed = parse(write, "report.pdf", pdf_bytes("Page One text", "Page Two text", "Page Three text"))

    assert [(segment.page, segment.text) for segment in parsed.segments] == [
        (1, "Page One text"),
        (3, "Page Three text"),
    ]
    assert parsed.num_pages == 3
    assert "Skipping unreadable page 2 of report.pdf" in caplog.text


# ---------------- Office formats ----------------


def test_docx_keeps_paragraphs_and_tables_in_body_order(write: WriteFile) -> None:
    document = new_docx()
    document.add_paragraph("Introduction paragraph")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.cell(0, 2).text = "Third"
    for column, value in enumerate(["Ada", "Engineer", "London"]):
        table.cell(1, column).text = value
    document.add_paragraph("Closing paragraph")
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse(write, "team.docx", buffer.getvalue())

    assert parsed == ParsedDocument(
        segments=[
            Segment("Introduction paragraph\nMerged heading | Third\nAda | Engineer | London\nClosing paragraph", None)
        ],
        num_pages=None,
    )


def test_pptx_has_one_segment_per_slide_with_groups_tables_and_notes(write: WriteFile) -> None:
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    first = presentation.slides.add_slide(blank)
    first.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Quarterly results"
    first.notes_slide.notes_text_frame.text = "Mention the record quarter"
    group = presentation.slides.add_slide(blank).shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Grouped roadmap item"
    presentation.slides.add_slide(blank)
    table = presentation.slides.add_slide(blank).shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2))
    for (row, column), value in zip(
        [(0, 0), (0, 1), (1, 0), (1, 1)], ["Region", "Revenue", "North", "42"], strict=True
    ):
        table.table.cell(row, column).text = value
    buffer = io.BytesIO()
    presentation.save(buffer)

    parsed = parse(write, "deck.pptx", buffer.getvalue())

    assert parsed == ParsedDocument(
        segments=[
            Segment("Quarterly results\nNotes: Mention the record quarter", 1),
            Segment("Grouped roadmap item", 2),
            Segment("Region | Revenue\nNorth | 42", 4),
        ],
        num_pages=4,
    )


def test_pptx_notes_helper_matches_the_conftest_builder(write: WriteFile) -> None:
    parsed = parse(write, "talk.pptx", pptx_bytes("Opening slide", "Second slide", notes=["Say hello first"]))

    assert parsed.segments == [Segment("Opening slide\nNotes: Say hello first", 1), Segment("Second slide", 2)]


def test_xlsx_has_one_segment_per_non_empty_sheet(write: WriteFile) -> None:
    workbook = Workbook()
    budget = workbook.active
    assert budget is not None
    budget.title = "Budget"
    for row in [
        ["Item", "Cost", None],
        [None, None],
        ["Laptop", 1200.0],
        ["Ratio", 2.5],
        ["Due", datetime(2026, 1, 2)],
        ["At", datetime(2026, 1, 2, 13, 45)],
        [None, "  spaced   out  "],
    ]:
        budget.append(row)
    workbook.create_sheet("Empty")
    staff = workbook.create_sheet("Staff")
    staff.append(["Name", "Team"])
    staff.append(["Ada", "Core"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    parsed = parse(write, "budget.xlsx", buffer.getvalue())

    assert parsed == ParsedDocument(
        segments=[
            Segment(
                "Sheet: Budget\nItem | Cost\nLaptop | 1200\nRatio | 2.5\nDue | 2026-01-02\n"
                "At | 2026-01-02 13:45:00\n | spaced out",
                None,
            ),
            Segment("Sheet: Staff\nName | Team\nAda | Core", None),
        ],
        num_pages=None,
    )


def test_xlsx_rows_past_the_row_limit_are_left_out(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(parsing, "_XLSX_MAX_ROWS", 3)

    parsed = parse(write, "long.xlsx", xlsx_bytes([[f"Row {number}"] for number in range(1, 8)], title="Log"))

    assert parsed.segments == [Segment("Sheet: Log\nRow 1\nRow 2\nRow 3", None)]


def test_xlsx_with_too_many_cells_is_refused(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(parsing, "_XLSX_MAX_CELLS", 10)

    error = parse_error(write, "wide.xlsx", xlsx_bytes([["a", "b", "c", "d"]] * 3))

    assert error == "The spreadsheet is too large to process."


def test_xlsx_parsing_stops_at_the_deadline_between_rows(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "parse_timeout_seconds", 60)
    monkeypatch.setattr(parsing, "_XLSX_DEADLINE_ROWS", 2)
    clock = iter(range(0, 10_000, 45))  # every reading is 45 s later
    monkeypatch.setattr(parsing, "time", SimpleNamespace(monotonic=lambda: next(clock)))

    error = parse_error(write, "slow.xlsx", xlsx_bytes([[f"Row {number}"] for number in range(10)]))

    assert error == "The file is too large or complex to process (it took over 60 seconds)."


def test_xlsx_builder_output_is_extracted_as_rows(write: WriteFile) -> None:
    parsed = parse(write, "rows.xlsx", xlsx_bytes([["City", "Population"], ["Oslo", 709000]], title="Cities"))

    assert parsed.segments == [Segment("Sheet: Cities\nCity | Population\nOslo | 709000", None)]


@pytest.mark.parametrize("extension", [".docx", ".pptx", ".xlsx"])
def test_zip_bombs_are_refused_before_inflating(
    write: WriteFile, monkeypatch: pytest.MonkeyPatch, extension: str
) -> None:
    monkeypatch.setattr(settings, "max_uncompressed_mb", 1)
    bomb = padded_zip(sample_file(extension, "Harmless looking content."), 2 * 1024 * 1024)

    assert len(bomb) < 100 * 1024
    assert parse_error(write, f"bomb{extension}", bomb) == (
        "The file expands to more than 1 MB when uncompressed, which isn't allowed."
    )


@pytest.mark.parametrize("extension", [".docx", ".pptx", ".xlsx"])
def test_office_files_within_the_uncompressed_limit_are_read(
    write: WriteFile, monkeypatch: pytest.MonkeyPatch, extension: str
) -> None:
    monkeypatch.setattr(settings, "max_uncompressed_mb", 1)
    data = padded_zip(sample_file(extension, "Content next to some padding."), 1024)

    assert "Content next to some padding." in parse(write, f"padded{extension}", data).segments[0].text


@pytest.mark.parametrize("extension", [".docx", ".pptx"])
def test_oversized_xml_parts_are_refused_before_loading(
    write: WriteFile, monkeypatch: pytest.MonkeyPatch, extension: str
) -> None:
    monkeypatch.setattr(settings, "max_xml_part_mb", 1)
    big_part = padded_zip(sample_file(extension, "Harmless looking content."), 2 * 1024 * 1024, "customXml/padding.xml")
    small_part = padded_zip(sample_file(extension, "Content next to a part."), 1024 * 1024, "customXml/padding.xml")

    assert parse_error(write, f"big{extension}", big_part) == TOO_COMPLEX
    assert "Content next to a part." in parse(write, f"small{extension}", small_part).segments[0].text


def test_spreadsheets_stream_their_xml_parts(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "max_xml_part_mb", 1)
    data = padded_zip(sample_file(".xlsx", "Cells are streamed."), 2 * 1024 * 1024, "customXml/padding.xml")

    assert "Cells are streamed." in parse(write, "big.xlsx", data).segments[0].text


@pytest.mark.parametrize("extension", [".docx", ".pptx", ".xlsx"])
def test_corrupt_office_files_are_reported_as_invalid(write: WriteFile, extension: str) -> None:
    other_zip = io.BytesIO()
    with zipfile.ZipFile(other_zip, "w") as archive:
        archive.writestr("hello.txt", "Not an office document")
    expected = f"The file is corrupted or not a valid {extension} file."

    assert parse_error(write, f"plain{extension}", b"definitely not a zip archive") == expected
    assert parse_error(write, f"zip{extension}", other_zip.getvalue()) == expected


def test_empty_office_document_has_no_text(write: WriteFile) -> None:
    assert parse_error(write, "empty.docx", docx_bytes()) == NO_TEXT
    assert parse_error(write, "empty.xlsx", xlsx_bytes([])) == NO_TEXT
    assert parse_error(write, "empty.pptx", pptx_bytes("")) == NO_TEXT


# ---------------- HTML ----------------


@pytest.mark.parametrize("extension", [".html", ".htm"])
def test_html_keeps_title_blocks_and_tables_and_drops_scripts(write: WriteFile, extension: str) -> None:
    page = (
        b"<!doctype html><html><head><title>Field  guide</title><style>.x{color:red}</style>"
        b"<script>var SCRIPT_MARKER = 1;</script></head><body><h1>Heading</h1>"
        b"<p>First <b>bold</b> paragraph</p><p>Line<br>break</p>"
        b"<table><tr><th>Bird</th><th>Habitat</th></tr><tr><td>Heron</td><td>Wetlands</td></tr></table>"
        b"<noscript>NOSCRIPT_MARKER</noscript><svg><text>SVG_MARKER</text></svg>"
        b"<template>TEMPLATE_MARKER</template><div>Last   words</div></body></html>"
    )

    parsed = parse(write, f"guide{extension}", page)

    assert parsed == ParsedDocument(
        segments=[
            Segment(
                "Field guide\n\nHeading\n\nFirst bold paragraph\n\nLine\nbreak\n\nBird | Habitat\n\n"
                "Heron | Wetlands\n\nLast words",
                None,
            )
        ],
        num_pages=None,
    )


def test_html_builder_drops_the_script_marker(write: WriteFile) -> None:
    text = parse(write, "page.html", html_bytes("Paragraph one\nParagraph two", title="Tide tables")).segments[0].text

    assert text == "Tide tables\n\nParagraph one\n\nParagraph two"
    assert "SCRIPT_MARKER" not in text


def test_nested_tables_keep_their_cells_apart(write: WriteFile) -> None:
    page = b"<table><tr><td>a<table><tr><td>x</td><td>y<br>z</td></tr></table></td><td>b</td><td>c</td></tr></table>"

    assert parse(write, "nested.html", page).segments == [Segment("a\n\nx | y\nz\n\n| b | c", None)]


@pytest.mark.parametrize(
    "page",
    ["<p>x</p>" + "<div>" * 100_000, "<p>x</p>" + "<table><tr><td>a</td><td>" * 10_000],
    ids=["divs", "tables"],
)
def test_deeply_nested_html_is_read_in_linear_time(write: WriteFile, page: str) -> None:
    started = time.perf_counter()

    text = parse(write, "deep.html", page.encode()).segments[0].text

    assert time.perf_counter() - started < 5  # it used to take minutes
    assert text.startswith("x")


def test_html_too_large_for_the_text_cap_is_refused_before_parsing(
    write: WriteFile, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "max_extracted_chars", 100)

    assert parse_error(write, "markup.html", b"<b></b>" * 58) == TOO_COMPLEX
    assert parse(write, "fits.html", b"<b></b>" * 55 + b"<p>ok</p>").segments == [Segment("ok", None)]


def test_html_with_only_scripts_has_no_text(write: WriteFile) -> None:
    assert parse_error(write, "empty.html", b"<html><head><script>x()</script></head><body> </body></html>") == NO_TEXT


# ---------------- Plain text, CSV, JSON ----------------


@pytest.mark.parametrize("extension", [".txt", ".md", ".markdown", ".csv"])
@pytest.mark.parametrize(
    ("data", "expected"),
    [
        ("\ufeffCaf\u00e9 au lait".encode(), "Caf\u00e9 au lait"),
        (b"Caf\xe9 \x93quoted\x94", "Caf\u00e9 \u201cquoted\u201d"),
        ("UTF-16 encoded text".encode("utf-16"), "UTF-16 encoded text"),
    ],
    ids=["utf8-bom", "cp1252", "utf16"],
)
def test_text_files_are_decoded(write: WriteFile, extension: str, data: bytes, expected: str) -> None:
    assert parse(write, f"notes{extension}", data) == ParsedDocument(segments=[Segment(expected, None)], num_pages=None)


def test_text_is_normalised(write: WriteFile) -> None:
    raw = (
        "\n\n  Line one   \r\nLine two\x00 has a NUL\r\n\r\n\r\n\r\n"
        "After blanks\x0cNext page\x07bell\u2028sep\t tab \x7f\n\n"
    )

    assert parse(write, "messy.txt", raw.encode()).segments == [
        Segment("Line one\nLine two has a NUL\n\nAfter blanks\nNext pagebell\nsep\t tab", None)
    ]


def test_json_is_pretty_printed(write: WriteFile) -> None:
    value = {"name": "Caf\u00e9", "skills": ["math", "engines"], "active": True}

    assert parse(write, "data.json", json_bytes(value)).segments == [
        Segment(json.dumps(value, indent=2, ensure_ascii=False), None)
    ]


def test_invalid_json_is_kept_as_text(write: WriteFile) -> None:
    assert parse(write, "broken.json", b'{"unterminated": "value').segments == [
        Segment('{"unterminated": "value', None)
    ]


def test_json_too_long_to_pretty_print_keeps_the_compact_text(
    write: WriteFile, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = json_bytes({"items": [{"id": number, "name": f"item {number}"} for number in range(5)]})
    monkeypatch.setattr(settings, "max_extracted_chars", len(data))

    assert parse(write, "compact.json", data).segments == [Segment(data.decode(), None)]


def test_whitespace_only_text_has_no_text(write: WriteFile) -> None:
    assert parse_error(write, "blank.txt", b"  \n\t \r\n \x00 ") == NO_TEXT


@pytest.mark.parametrize("extension", [".exe", ".doc", ".rtf", ""])
def test_unsupported_extensions_are_refused(write: WriteFile, extension: str) -> None:
    path = write(f"file{extension}", b"Some content that would otherwise be readable.")

    with pytest.raises(DocumentParseError) as caught:
        parse_file(path, extension)

    assert str(caught.value) == f"Unsupported file type: {extension or '(none)'}."


def test_extension_is_case_insensitive(write: WriteFile) -> None:
    path = write("upper.TXT", b"Upper case extensions are fine.")

    assert parse_file(path, ".TXT").segments == [Segment("Upper case extensions are fine.", None)]


# ---------------- Extracted-text cap ----------------


def test_text_at_the_cap_is_accepted(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "max_extracted_chars", 100)

    assert parse(write, "exact.txt", txt_bytes("a" * 100)).segments == [Segment("a" * 100, None)]


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("long.txt", txt_bytes("a" * 101)),
        ("pages.pdf", pdf_bytes("p" * 60, "q" * 60)),
        ("sheet.xlsx", xlsx_bytes([["x" * 40], ["y" * 40], ["z" * 40]])),
        ("long.docx", docx_bytes("d" * 60, "e" * 60)),
        ("long.html", html_bytes("h" * 120)),
        ("long.json", json_bytes(["j" * 120])),
    ],
)
def test_too_much_text_is_refused(write: WriteFile, monkeypatch: pytest.MonkeyPatch, name: str, data: bytes) -> None:
    monkeypatch.setattr(settings, "max_extracted_chars", 100)

    assert parse_error(write, name, data) == "The document has too much text (over 100 characters)."


def test_text_cap_message_uses_thousands_separators(write: WriteFile, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "max_extracted_chars", 5000)

    assert parse_error(write, "huge.md", b"m" * 5001) == "The document has too much text (over 5,000 characters)."
