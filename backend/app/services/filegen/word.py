"""DOCX rendering with python-docx (which never writes macros), from the block tree.

Embedded images reach python-docx only as in-memory JPEG data (images.Embedder), never as a path the content names.
"""

import io
import unicodedata
from collections.abc import Iterable
from datetime import datetime, timezone

import docx
from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Emu, Inches, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.text.run import Run

from . import CREATOR
from .images import Embedder
from .markdown import (
    Block,
    Code,
    Figure,
    Heading,
    ListBlock,
    Paragraph,
    Quote,
    Rule,
    Span,
    Table,
    check_image_caps,
    check_table_caps,
    link_label,
    text_reference,
)

CODE_FONT = "Consolas"
CODE_PT = 9.5
LINK_COLOR = RGBColor(0x1D, 0x4E, 0xD8)
CODE_FILL = "F3F4F6"
RULE_COLOR = "D1D5DB"
LIST_INDENT_IN = 0.25
FIGURE_MAX_HEIGHT = 0.6  # of the page's usable height
_USABLE_IN = (6.0, 9.0)  # the default template's text area (Letter, 1.25 in and 1 in margins)
_MAX_LIST_STYLE_LEVEL = 3  # the default template has List Bullet/Number up to 3
_ALIGNMENTS = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
# Child order of w:pPr (ECMA-376), for inserting elements python-docx has no property for.
_PPR_ORDER = (
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl", "w:numPr",
    "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
    "w:sectPr", "w:pPrChange",
)  # fmt: skip


def render_docx(blocks: list[Block], title: str) -> tuple[bytes, list[str]]:
    """(the DOCX, notes for the tool result)."""
    check_table_caps(blocks)
    check_image_caps(blocks)
    document = docx.Document()
    properties = document.core_properties
    properties.title = title
    properties.author = CREATOR
    properties.last_modified_by = CREATOR
    properties.created = properties.modified = datetime.now(timezone.utc)
    properties.revision = 1
    properties.comments = ""
    writer = _Writer(document)
    writer.blocks(blocks)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), writer.images.notes()


class _Writer:
    def __init__(self, document: Document) -> None:
        self.document = document
        self.images = Embedder()

    def blocks(self, blocks: list[Block], *, level: int = 0, quote: bool = False) -> None:
        for block in blocks:
            if isinstance(block, Heading):
                paragraph = self.document.add_paragraph(style=f"Heading {min(block.level, 6)}")
                self.inline(paragraph, block.spans)
            elif isinstance(block, Paragraph):
                style = "Quote" if quote else (f"List Continue{_suffix(level)}" if level else None)
                self.inline(self.document.add_paragraph(style=style), block.spans)
            elif isinstance(block, ListBlock):
                self.list_block(block, level + 1, quote)
            elif isinstance(block, Quote):
                self.blocks(block.blocks, level=level, quote=True)
            elif isinstance(block, Code):
                self.code(block.text, level)
            elif isinstance(block, Rule):
                self.rule()
            elif isinstance(block, Figure):
                self.figure(block, level, quote)
            else:
                self.table(block)

    def list_block(self, block: ListBlock, level: int, quote: bool) -> None:
        style = ("List Number" if block.ordered else "List Bullet") + _suffix(level)
        number_id = self._restarted_numbering(style, block.start) if block.ordered else None
        for item in block.items:
            first = item[0] if item else None
            paragraph = self.document.add_paragraph(style=style)
            if number_id is not None:
                num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_ilvl().val = 0
                num_pr.get_or_add_numId().val = number_id
            if isinstance(first, Paragraph):
                self.inline(paragraph, first.spans)
                rest = item[1:]
            else:
                rest = item
            self.blocks(rest, level=level, quote=quote)

    def _restarted_numbering(self, style: str, start: int) -> int:
        """A new numbering instance on the style's list definition, so every ordered list starts at its own start."""
        numbering = self.document.part.numbering_part.numbering_definitions._numbering
        style_num_pr = self.document.styles[style].element.pPr.numPr
        abstract_id = numbering.num_having_numId(style_num_pr.numId.val).abstractNumId.val
        number = numbering.add_num(abstract_id)
        number.add_lvlOverride(ilvl=0).add_startOverride(start)
        return int(number.numId)

    def inline(self, paragraph: DocxParagraph, spans: list[Span]) -> None:
        if _right_to_left(span.text for span in spans):
            _insert(paragraph._p.get_or_add_pPr(), OxmlElement("w:bidi"))
            rtl = True
        else:
            rtl = False
        position = 0
        while position < len(spans):
            span = spans[position]
            end = position + 1
            while span.link and end < len(spans) and spans[end].link == span.link:
                end += 1
            runs = [self.run(paragraph, item, rtl) for item in spans[position:end]]
            if span.link:
                _hyperlink(paragraph, runs, span.link)
                label = link_label("".join(item.text for item in spans[position:end]), span.link)
                if label:
                    host = self.run(paragraph, Span(f" ({label})", bold=span.bold, italic=span.italic), rtl)
                    host.font.color.rgb = LINK_COLOR
            position = end

    def run(self, paragraph: DocxParagraph, span: Span, rtl: bool) -> Run:
        if span.text == "\n":
            run = paragraph.add_run()
            run.add_break()
            return run
        run = paragraph.add_run(span.text)
        font = run.font
        if span.bold:
            font.bold = font.cs_bold = True
        if span.italic:
            font.italic = font.cs_italic = True
        if span.strike:
            font.strike = True
        if span.code:
            font.name = CODE_FONT
            _size(run, CODE_PT)
        if span.link:
            font.color.rgb = LINK_COLOR
            font.underline = True
        if rtl:
            font.rtl = True
        return run

    def code(self, text: str, level: int) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), CODE_FILL)
        _insert(paragraph._p.get_or_add_pPr(), shading)
        for number, line in enumerate(text.split("\n")):
            if number:
                paragraph.runs[-1].add_break()
            run = paragraph.add_run(line)
            run.font.name = CODE_FONT
            _size(run, CODE_PT)

    def figure(self, figure: Figure, level: int, quote: bool) -> None:
        """A centered picture as wide as the text (less 0.25 in per list level, at most 60% of the page high), kept
        with its Caption paragraph."""
        loaded = self.images.load(figure.ref)
        if loaded is None:
            style = "Quote" if quote else (f"List Continue{_suffix(level)}" if level else None)
            self.inline(self.document.add_paragraph(style=style), [Span(text_reference(figure))])
            return
        data, pixels_wide, pixels_high = loaded
        usable_width, usable_height = self._usable_area()
        indent = Inches(LIST_INDENT_IN * level)
        available = usable_width - indent
        tallest = FIGURE_MAX_HEIGHT * usable_height
        width, height = float(available), available * pixels_high / pixels_wide
        if height > tallest:
            width, height = tallest * pixels_wide / pixels_high, tallest
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        if level:
            paragraph.paragraph_format.left_indent = indent
        shape = paragraph.add_run().add_picture(io.BytesIO(data), width=Emu(int(width)), height=Emu(int(height)))
        shape._inline.docPr.set("descr", figure.ref.alt_text or figure.caption)
        if figure.caption:
            caption = self.document.add_paragraph(style="Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if level:
                caption.paragraph_format.left_indent = indent
            self.inline(caption, [Span(figure.caption)])

    def _usable_area(self) -> tuple[int, int]:
        """The text area of the last section in EMU (width, height)."""
        section = self.document.sections[-1]
        lengths = (
            section.page_width,
            section.page_height,
            section.left_margin,
            section.right_margin,
            section.top_margin,
            section.bottom_margin,
        )
        values = [int(length) for length in lengths if length is not None]
        if len(values) < len(lengths):
            return Inches(_USABLE_IN[0]), Inches(_USABLE_IN[1])
        width, height, left, right, top, bottom = values
        return width - left - right, height - top - bottom

    def rule(self) -> None:
        paragraph = self.document.add_paragraph()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", RULE_COLOR)):
            bottom.set(qn(key), value)
        borders.append(bottom)
        _insert(paragraph._p.get_or_add_pPr(), borders)

    def table(self, table: Table) -> None:
        columns = len(table.header)
        if not columns:
            return
        grid = self.document.add_table(rows=1, cols=columns)
        grid.style = self.document.styles["Table Grid"]
        header_row = grid.rows[0]
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_row._tr.get_or_add_trPr().append(repeat)
        for cells, row in [(table.header, header_row)] + [(cells, grid.add_row()) for cells in table.rows]:
            for column, (spans, cell) in enumerate(zip(cells, row.cells, strict=False)):
                if row is header_row:
                    spans = [Span(span.text, True, span.italic, span.strike, span.code, span.link) for span in spans]
                self._cell(cell, spans, table.aligns[column] if column < len(table.aligns) else "left")

    def _cell(self, cell: _Cell, spans: list[Span], align: str) -> None:
        paragraph = cell.paragraphs[0]
        paragraph.alignment = _ALIGNMENTS.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        self.inline(paragraph, spans)


def _suffix(level: int) -> str:
    level = min(max(level, 1), _MAX_LIST_STYLE_LEVEL)
    return "" if level == 1 else f" {level}"


def _size(run: Run, points: float) -> None:
    """The size for every script: w:sz and w:szCs (python-docx sets only w:sz)."""
    run.font.size = Pt(points)
    r_pr = run._r.get_or_add_rPr()
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(round(points * 2)))
    r_pr.find(qn("w:sz")).addnext(size_cs)


def _hyperlink(paragraph: DocxParagraph, runs: list[Run], url: str) -> None:
    relationship = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    runs[0]._r.addprevious(link)
    for run in runs:
        link.append(run._r)


def _insert(p_pr: BaseOxmlElement, element: BaseOxmlElement) -> None:
    tag = "w:" + element.tag.split("}", 1)[1]
    p_pr.insert_element_before(element, *_PPR_ORDER[_PPR_ORDER.index(tag) + 1 :])


def _right_to_left(texts: Iterable[str]) -> bool:
    """Whether the first strong character is right-to-left (bidi class R or AL)."""
    for text in texts:
        for ch in text:
            kind = unicodedata.bidirectional(ch)
            if kind == "L":
                return False
            if kind in ("R", "AL"):
                return True
    return False
