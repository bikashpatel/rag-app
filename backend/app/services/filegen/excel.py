"""XLSX rendering with openpyxl: typed cells, a styled header row, and no formulas ever."""

import io
import re
import unicodedata
from datetime import date

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import CREATOR, RenderRequest, SheetData
from .markdown import parse
from .mdwriter import write

Value = str | int | float | bool | None

MAX_SHEET_NAME = 31
_FORBIDDEN_IN_NAMES = re.compile(r"[\[\]:*?/\\]")
_LINE_BREAKS = re.compile(r"[\t\n\r]+")
_RESERVED_NAMES = frozenset({"history"})  # Excel reserves it
_DECIMAL = re.compile(r"-?(0|[1-9]\d{0,14})(\.\d+)?")
_ISO_DATE = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill(fill_type="solid", start_color="F3F4F6", end_color="F3F4F6")
_MIN_WIDTH, _MAX_WIDTH = 8, 60
_NOTES_WIDTH = 100


def render_xlsx(request: RenderRequest) -> bytes:
    workbook = Workbook()
    workbook.properties.title = request.title
    workbook.properties.creator = CREATOR
    workbook.properties.lastModifiedBy = CREATOR
    # Removed first: while it exists, openpyxl renames a sheet the model calls "Sheet" to "Sheet1".
    workbook.remove(workbook.worksheets[0])
    names: set[str] = set()
    for number, sheet in enumerate(request.sheets, start=1):
        _fill(workbook.create_sheet(_sheet_name(sheet.name, number, names)), sheet)
    if request.content.strip():
        notes = workbook.create_sheet(_sheet_name("Notes", len(names) + 1, names))
        notes.column_dimensions["A"].width = _NOTES_WIDTH
        for row, line in enumerate(write(parse(request.content)).rstrip("\n").split("\n"), start=1):
            if line:
                cell = notes.cell(row=row, column=1)
                _set_text(cell, line)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    if request.sources:
        rows: list[list[Value]] = [
            [", ".join(f"[{n}]" for n in source.numbers), source.filename, source.page] for source in request.sources
        ]
        sources = workbook.create_sheet(_sheet_name("Sources", len(names) + 1, names))
        _fill(sources, SheetData("Sources", ["Refs", "Document", "Page"], rows))
    if not workbook.worksheets:
        workbook.create_sheet("Sheet1")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sheet_name(name: str, number: int, taken: set[str]) -> str:
    """A valid, unique (ignoring case) sheet name; `taken` holds the lowercased names used so far."""
    name = ILLEGAL_CHARACTERS_RE.sub("", _LINE_BREAKS.sub(" ", name))
    name = "".join(ch for ch in name if unicodedata.category(ch) not in ("Cc", "Cf"))
    name = _FORBIDDEN_IN_NAMES.sub("", name).strip().strip("'").strip()[:MAX_SHEET_NAME].strip()
    if not name or name.lower() in _RESERVED_NAMES:
        name = f"Sheet{number}"
    candidate, copy = name, 1
    while candidate.lower() in taken:
        copy += 1
        suffix = f" ({copy})"
        candidate = name[: MAX_SHEET_NAME - len(suffix)].rstrip() + suffix
    taken.add(candidate.lower())
    return candidate


def _fill(sheet: Worksheet, data: SheetData) -> None:
    widths: dict[int, int] = {}
    texts = right_to_left = 0
    first_row = 1
    if data.columns:
        for column, title in enumerate(data.columns, start=1):
            cell = sheet.cell(row=1, column=column)
            _set_text(cell, str(title))
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            widths[column] = len(str(title))
            texts += 1
            right_to_left += _starts_right_to_left(str(title))
        first_row = 2
    for row_number, row in enumerate(data.rows, start=first_row):
        for column, value in enumerate(row, start=1):
            if value is None or value == "":
                continue
            cell = sheet.cell(row=row_number, column=column)
            shown = _set_value(cell, value)
            widths[column] = max(widths.get(column, 0), len(shown))
            if cell.data_type == "s":
                texts += 1
                right_to_left += _starts_right_to_left(shown)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, _MIN_WIDTH), _MAX_WIDTH)
    if data.columns:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    if texts and right_to_left * 2 > texts:
        sheet.sheet_view.rightToLeft = True


def _set_value(cell: Cell, value: Value) -> str:
    """Store the value with its type; returns the text shown, for the column width."""
    if isinstance(value, bool):
        cell.value = value
        return str(value).upper()
    if isinstance(value, int | float):
        cell.value = value
        return str(value)
    text = ILLEGAL_CHARACTERS_RE.sub("", value)
    if _DECIMAL.fullmatch(text):
        cell.value = float(text) if "." in text else int(text)
        return text
    if (found := _ISO_DATE.fullmatch(text)) is not None:
        try:
            cell.value = date(*(int(part) for part in found.groups()))
        except ValueError:
            pass
        else:
            cell.number_format = "yyyy-mm-dd"
            return text
    _set_text(cell, text)
    return text


def _set_text(cell: Cell, text: str) -> None:
    """Always a text cell: openpyxl would store text starting with "=" as a formula, and Excel evaluates text that
    starts with = + - @ once the cell is edited, unless the cell has the quote prefix."""
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    cell.value = text
    cell.data_type = "s"
    if text.startswith(_FORMULA_START):
        cell.quotePrefix = True


def _starts_right_to_left(text: str) -> bool:
    for ch in text:
        kind = unicodedata.bidirectional(ch)
        if kind in ("R", "AL"):
            return True
        if kind == "L":
            return False
    return False
