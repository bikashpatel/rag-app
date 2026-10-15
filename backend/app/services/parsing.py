"""Uploaded file -> plain-text segments. Synchronous and CPU-bound: call it through asyncio.to_thread."""

import codecs
import json
import logging
import re
import time
import zipfile
from collections.abc import Callable, Iterable, Iterator, Sequence
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Literal

from bs4 import BeautifulSoup, CData, NavigableString, Tag
from docx import Document as open_docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation
from pptx.shapes.group import GroupShape
from pypdf import PageObject, PasswordType, PdfReader
from pypdf.errors import DependencyError
from pypdf.generic import ArrayObject, IndirectObject

from app.config import settings
from app.services.pdf_images import (
    PageMeasure,
    PageMeasurer,
    Rect,
    Rendered,
    RenderJob,
    page_rotation,
    render_pages,
    union_share,
    visible_box,
)

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".md", ".markdown", ".txt", ".csv", ".json")
LEGACY_EXTENSIONS = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}

_NO_PDF_TEXT = (
    "No extractable text: this PDF looks scanned or image-only. Ask the admin to enable a vision model to read it."
)
_NO_PDF_TEXT_VISION = "No extractable text: this PDF has neither a text layer nor page images to read."
_NO_TEXT = "The document contains no text."
_TOO_COMPLEX = "The file is too large or complex to process."
_UNSUPPORTED_PDF_ENCRYPTION = (
    "This PDF uses encryption that isn't supported. Save an unencrypted copy and upload it again."
)
_ZIP_FORMATS = frozenset({".docx", ".pptx", ".xlsx"})
# python-docx and python-pptx load each XML part into a full DOM (about 25 bytes of memory per byte of XML).
_DOM_FORMATS = frozenset({".docx", ".pptx"})
# Plain text needs at least one byte per character and at most four, so a bigger file has too much text.
_MAX_BYTES_PER_CHAR = 4
_XLSX_DEADLINE_ROWS = 1000
# Excel's own row limit, and a bound on the cells a crafted sheet can make openpyxl generate.
_XLSX_MAX_ROWS = 1_048_576
_XLSX_MAX_CELLS = 50_000_000
_HTML_DROPPED = frozenset({"script", "style", "noscript", "template", "svg", "title"})
_HTML_BLOCKS = frozenset(
    "address article aside blockquote caption dd div dl dt fieldset figcaption figure footer form h1 h2 h3 h4 h5 h6 "
    "header hr li main nav ol p pre section table tr ul".split()
)
_HTML_CELLS = frozenset({"td", "th"})
# The string types get_text() keeps: no comments, doctypes or <rt>/<rp> text.
_HTML_STRING_TYPES = (NavigableString, CData)
_NEWLINE, _CELL_SEPARATOR = object(), object()
# C0 controls except tab/newline are dropped; vertical tab, form feed and Unicode line breaks become newlines.
_CONTROL_CHARS = {code: None for code in (*range(0x00, 0x09), *range(0x0E, 0x20), 0x7F)} | {
    0x0B: "\n",
    0x0C: "\n",
    0x2028: "\n",
    0x2029: "\n",
}
_BLANK_LINES = re.compile(r"\n{3,}")

# Vision model (SPEC section 7, step 4). Images under this share of the page are bullets and icons.
_MIN_IMAGE_SHARE = 0.01
# An image drawn on this many units, or on every unit of a multi-unit PDF, is a letterhead, watermark or template.
_REPEATED_MIN_UNITS = 3
_PAGE_IMAGE_SHARE = 0.9  # images covering this much of a page make it a page image (a scan)
_OCR_LAYER_SHARE = 0.5  # a page image whose text is mostly invisible has an OCR layer
_OVERLAY_CHARS = 200  # less visible text on a page image is an overlay (envelope IDs, Bates numbers, fax headers)
_MAX_SOURCE_PIXELS = 100_000_000


class DocumentParseError(Exception):
    """The file can't be turned into text. str(exc) is safe to show to end users."""


@dataclass
class Segment:
    text: str
    page: int | None


VisionMode = Literal["transcribe", "describe"]


@dataclass(frozen=True)
class VisionRequest:
    out_dir: Path  # page JPEGs; the caller owns and removes it
    cached: frozenset[tuple[int, str]]  # (page, mode) already read with the current fingerprint
    blocked: str | None = None  # the error for an uncached scanned page (the owner's allowance is used up)


@dataclass(frozen=True)
class PageImage:
    page: int  # 1-based representative page of the unit
    pages: tuple[int, ...]  # every page of the unit, `page` first
    mode: VisionMode
    image: Path | None  # None when cached
    size: tuple[int, int] | None


@dataclass
class ParsedDocument:
    segments: list[Segment]
    num_pages: int | None
    page_images: list[PageImage] = field(default_factory=list)


class _TextBudget:
    """Running total of extracted characters that fails as soon as the cap is passed, plus the parse deadline."""

    def __init__(self) -> None:
        self.used = 0
        self.deadline = time.monotonic() + settings.parse_timeout_seconds

    def take(self, text: str) -> str:
        self.used += len(text)
        if self.used > settings.max_extracted_chars:
            raise _too_much_text()
        return text

    def check_time(self) -> None:
        """Called between pages, slides and rows: the parse runs in a thread, which can't be cancelled."""
        if time.monotonic() > self.deadline:
            raise DocumentParseError(
                f"The file is too large or complex to process (it took over {settings.parse_timeout_seconds} seconds)."
            )


_Reader = Callable[[Path, _TextBudget], tuple[list[Segment], int | None]]


def parse_file(path: Path, extension: str, vision: VisionRequest | None = None) -> ParsedDocument:
    """Extract normalized text segments (one per page, slide or sheet where the format has them).

    With `vision`, PDF pages without a usable text layer and pages with significant images are selected for the
    vision model (rendered to `vision.out_dir` unless cached); scanned pages selected that way have no segment.
    """
    extension = extension.lower()
    reader = _READERS.get(extension)
    if reader is None and extension != ".pdf":
        raise DocumentParseError(f"Unsupported file type: {extension or '(none)'}.")
    if extension in _ZIP_FORMATS:
        _check_zip(path, extension)
    page_images: list[PageImage] = []
    try:
        if reader is None:
            raw_segments, num_pages, page_images = _read_pdf(path, _TextBudget(), vision)
        else:
            raw_segments, num_pages = reader(path, _TextBudget())
    except (DocumentParseError, OSError):
        raise
    except Exception as exc:
        logger.warning("Could not parse %s: %r", path.name, exc)
        raise DocumentParseError(f"The file is corrupted or not a valid {extension} file.") from exc

    segments = [Segment(text=text, page=s.page) for s in raw_segments if (text := normalize_text(s.text))]
    if not segments and not any(image.mode == "transcribe" for image in page_images):
        if extension != ".pdf":
            raise DocumentParseError(_NO_TEXT)
        raise DocumentParseError(_NO_PDF_TEXT if vision is None else _NO_PDF_TEXT_VISION)
    check_total_text(segments)
    return ParsedDocument(segments=segments, num_pages=num_pages, page_images=page_images)


def check_total_text(segments: Sequence[Segment]) -> None:
    if sum(len(s.text) for s in segments) > settings.max_extracted_chars:
        raise _too_much_text()


def _too_much_text() -> DocumentParseError:
    return DocumentParseError(f"The document has too much text (over {settings.max_extracted_chars:,} characters).")


def _check_zip(path: Path, extension: str) -> None:
    """Refuse corrupt archives and zip bombs before any library inflates them.

    Python's zipfile never inflates an entry past its declared size, so the declared sizes are hard bounds.
    """
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, EOFError, ValueError) as exc:
        raise DocumentParseError(f"The file is corrupted or not a valid {extension} file.") from exc
    if sum(info.file_size for info in entries) > settings.max_uncompressed_mb * 1024 * 1024:
        raise DocumentParseError(
            f"The file expands to more than {settings.max_uncompressed_mb} MB when uncompressed, which isn't allowed."
        )
    if extension in _DOM_FORMATS and any(
        info.file_size > settings.max_xml_part_mb * 1024 * 1024
        for info in entries
        if info.filename.lower().endswith((".xml", ".rels"))
    ):
        raise DocumentParseError(_TOO_COMPLEX)


def normalize_text(text: str) -> str:
    text = text.encode("utf-8", "replace").decode("utf-8")  # lone surrogates can't be stored or sent
    text = text.replace("\r\n", "\n").replace("\r", "\n").translate(_CONTROL_CHARS)
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return _BLANK_LINES.sub("\n\n", text).strip()


@dataclass
class _Unit:
    """Pages that draw the same content: read, measured, rendered and sent to the vision model once."""

    pages: list[int]
    text: str | None  # None: the page's text couldn't be extracted


@dataclass
class _Candidate:
    unit: _Unit
    measure: PageMeasure
    mode: VisionMode
    tier: int  # 1: scan covered by images, 2: other page without text, 3: figures on a page with text
    figure_share: float  # union of the content images that individually cover less than _PAGE_IMAGE_SHARE

    @property
    def page(self) -> int:
        return self.unit.pages[0]


def _read_pdf(
    path: Path, budget: _TextBudget, vision: VisionRequest | None
) -> tuple[list[Segment], int | None, list[PageImage]]:
    try:
        reader = PdfReader(path)  # checks the empty password of AES-256 files already
    except DependencyError as exc:
        raise DocumentParseError(_UNSUPPORTED_PDF_ENCRYPTION) from exc
    if reader.is_encrypted:
        _decrypt_pdf(reader)
    num_pages = len(reader.pages)
    if num_pages > settings.max_pdf_pages:
        raise DocumentParseError(f"The PDF has more than {settings.max_pdf_pages:,} pages, which isn't supported.")
    segments: list[Segment] = []
    # Text by content streams and resources: pages that draw the same ones (a crafted file can repeat one
    # costly stream on thousands of pages) are read once.
    seen: dict[tuple[object, ...], _Unit] = {}
    units: list[_Unit] = []
    for number, page in enumerate(reader.pages, start=1):
        budget.check_time()
        key = _pdf_content_key(page)
        if key is not None and key in seen:
            unit = seen[key]
            unit.pages.append(number)
        else:
            try:
                text = page.extract_text() or ""
            except DependencyError as exc:
                raise DocumentParseError(_UNSUPPORTED_PDF_ENCRYPTION) from exc
            except Exception as exc:
                logger.warning("Skipping unreadable page %d of %s: %r", number, path.name, exc)
                text = None
            unit = _Unit(pages=[number], text=text)
            units.append(unit)
            if key is not None:
                seen[key] = unit
        if unit.text is not None:
            segments.append(Segment(text=budget.take(unit.text), page=number))
    if vision is None:
        return segments, num_pages, []
    page_images = _vision_pages(path, reader, units, budget, vision)
    transcribed = {page for image in page_images if image.mode == "transcribe" for page in image.pages}
    return [s for s in segments if s.page not in transcribed], num_pages, page_images


def _vision_pages(
    path: Path, reader: PdfReader, units: list[_Unit], budget: _TextBudget, vision: VisionRequest
) -> list[PageImage]:
    """Classify the units, select at most VISION_MAX_PAGES of them and render the ones not cached (SPEC section 7,
    step 4). Raises DocumentParseError when a scanned page can't be read."""
    measures = _measure_units(path, reader, units, budget)
    candidates = _classify(units, measures)
    selected = _select(path, candidates)
    cached = [c for c in selected if (c.page, c.mode) in vision.cached]
    uncached = [c for c in selected if (c.page, c.mode) not in vision.cached]
    if vision.blocked is not None and uncached:
        if any(c.mode == "transcribe" for c in uncached):
            raise DocumentParseError(vision.blocked)
        logger.info(
            "%s: %d page(s) not described: the owner's vision allowance is used up", path.name, len(uncached)
        )
        uncached = []
    rendered = _render(path, reader, uncached, budget, vision.out_dir)
    images = [PageImage(c.page, tuple(c.unit.pages), c.mode, None, None) for c in cached]
    for candidate in uncached:
        result = rendered.get(candidate.page)
        if isinstance(result, Rendered):
            images.append(
                PageImage(candidate.page, tuple(candidate.unit.pages), candidate.mode, result.path, result.size)
            )
    return sorted(images, key=lambda image: image.page)


def _measure_units(
    path: Path, reader: PdfReader, units: list[_Unit], budget: _TextBudget
) -> dict[int, tuple[PageMeasure, Rect]]:
    """Representative page -> (measure, visible box), for the units that may draw images and could be measured."""
    measurer = PageMeasurer(reader)
    measures: dict[int, tuple[PageMeasure, Rect]] = {}
    too_large = 0
    for unit in units:
        budget.check_time()
        number = unit.pages[0]
        page = reader.pages[number - 1]
        try:
            if not measurer.is_candidate(page):
                continue
            measure = measurer.measure(page)
            box = visible_box(page)
        except Exception as exc:
            logger.warning("Could not measure the images on page %d of %s: %r", number, path.name, exc)
            continue
        if measure is None:
            too_large += 1
        else:
            measures[number] = (measure, box)
    if too_large:
        logger.info("%s: %d page(s) too large to check for images; read as text", path.name, too_large)
    return measures


def _classify(units: list[_Unit], measures: dict[int, tuple[PageMeasure, Rect]]) -> list[_Candidate]:
    """The units for the vision model, by the rules in SPEC section 7, step 4 (first match wins)."""
    drawn_on: dict[tuple[int, int], int] = {}
    for measure, _ in measures.values():
        for key in {image.key for image in measure.images if image.key is not None}:
            drawn_on[key] = drawn_on.get(key, 0) + 1
    repeated = {
        key
        for key, count in drawn_on.items()
        if count >= _REPEATED_MIN_UNITS or (len(units) >= 2 and count == len(units))
    }
    candidates: list[_Candidate] = []
    for unit in units:
        if unit.pages[0] not in measures:
            continue  # no images, or unmeasured: text
        measure, box = measures[unit.pages[0]]
        content = [
            image
            for image in measure.images
            if image.share >= _MIN_IMAGE_SHARE and (image.key is None or image.key not in repeated)
        ]
        cover = union_share((image.rect for image in content), box)
        if cover < settings.vision_min_image_ratio:
            continue
        text_chars = len("".join(normalize_text(unit.text or "").split()))
        shown = measure.visible_bytes + measure.invisible_bytes
        invisible = measure.invisible_bytes / shown if shown else 0.0
        page_image = cover >= _PAGE_IMAGE_SHARE
        figures = union_share((image.rect for image in content if image.share < _PAGE_IMAGE_SHARE), box)
        if text_chars < settings.vision_min_text_chars:
            candidates.append(_Candidate(unit, measure, "transcribe", 1 if page_image else 2, figures))
        elif page_image and invisible >= _OCR_LAYER_SHARE:
            continue  # an OCR'd scan: its hidden text layer is the text
        elif page_image and text_chars * (1 - invisible) < _OVERLAY_CHARS:
            candidates.append(_Candidate(unit, measure, "transcribe", 1, figures))
        else:
            candidates.append(_Candidate(unit, measure, "describe", 3, figures))
    return candidates


def _select(path: Path, candidates: list[_Candidate]) -> list[_Candidate]:
    """At most VISION_MAX_PAGES units: scans covered by images first (more of them fail the document), then the
    other pages without text in page order, then figure pages, largest figures first. Pages that can't be rendered
    fail the document if they are scanned, and give their slot to the next figure page otherwise."""
    cap = settings.vision_max_pages
    scans = [c for c in candidates if c.tier == 1]
    if len(scans) > cap:
        raise DocumentParseError(
            f"This PDF has {len(scans):,} scanned pages, but at most {cap:,} pages per document can be read with the "
            "vision model (VISION_MAX_PAGES). Split it into smaller files and upload them separately."
        )
    ordered = [
        *scans,
        *(c for c in candidates if c.tier == 2),
        *sorted((c for c in candidates if c.tier == 3), key=lambda c: (-c.figure_share, c.page)),
    ]
    selected: list[_Candidate] = []
    refused: list[tuple[int, str]] = []
    left_out = 0
    for candidate in ordered:
        if len(selected) >= cap:
            left_out += len(candidate.unit.pages)
            continue
        reason = _render_refusal(candidate)
        if reason is None:
            selected.append(candidate)
        elif candidate.mode == "transcribe":
            refused.append((candidate.page, reason))
        else:
            logger.info("%s page %d: image description skipped: %s", path.name, candidate.page, reason)
    if refused:
        raise DocumentParseError(min(refused)[1])
    if left_out:
        logger.info("%s: %d page(s) not read with the vision model (VISION_MAX_PAGES=%d)", path.name, left_out, cap)
    return selected


def _render_refusal(candidate: _Candidate) -> str | None:
    """Why the page must not reach PDFium: too many operators, or a source image too large to decode."""
    if not candidate.measure.renderable:
        return f"Page {candidate.page} is too complex to render for the vision model."
    if max((image.pixels for image in candidate.measure.images), default=0) > _MAX_SOURCE_PIXELS:
        return f"Page {candidate.page} contains an image too large to process."
    return None


def _render(
    path: Path, reader: PdfReader, candidates: list[_Candidate], budget: _TextBudget, out_dir: Path
) -> dict[int, Rendered | str]:
    """Render the candidates' pages. A scanned page that fails fails the document; a figure page is skipped."""
    if not candidates:
        return {}
    jobs = []
    for candidate in candidates:
        page = reader.pages[candidate.page - 1]
        box = visible_box(page)  # what PDFium renders, so the child's size check matches valid pages
        jobs.append(
            RenderJob(page=candidate.page, width=box[2] - box[0], height=box[3] - box[1], rotate=page_rotation(page))
        )
    results = render_pages(path, jobs, len(reader.pages), out_dir, settings.vision_max_image_px, budget)
    failed = []
    for candidate in candidates:
        result = results.get(candidate.page, "missing")
        if isinstance(result, Rendered):
            continue
        logger.warning("%s page %d couldn't be rendered for the vision model: %s", path.name, candidate.page, result)
        if candidate.mode == "transcribe":
            failed.append(candidate.page)
    if failed:
        raise DocumentParseError(f"Page {min(failed)} couldn't be rendered for the vision model.")
    return results


def _pdf_content_key(page: PageObject) -> tuple[object, ...] | None:
    """Identifies what a page draws: its content stream objects and its resources (the same stream can show
    different text with other fonts or forms). None when the streams aren't all objects of their own."""
    contents = page.raw_get("/Contents") if "/Contents" in page else None
    if isinstance(contents, IndirectObject) and isinstance(contents.get_object(), ArrayObject):
        contents = contents.get_object()
    streams = list(contents) if isinstance(contents, ArrayObject) else [contents]
    if not all(isinstance(stream, IndirectObject) for stream in streams):
        return None
    resources = page.raw_get("/Resources") if "/Resources" in page else None
    # A direct resources dictionary is one Python object for every page that inherits or repeats it.
    shared = (resources.idnum, resources.generation) if isinstance(resources, IndirectObject) else id(resources)
    return (shared, *((stream.idnum, stream.generation) for stream in streams))


def _decrypt_pdf(reader: PdfReader) -> None:
    """Open PDFs that are encrypted with an empty user password (permission-restricted files)."""
    try:
        result = reader.decrypt("")
    except DependencyError as exc:
        raise DocumentParseError(_UNSUPPORTED_PDF_ENCRYPTION) from exc
    except Exception:
        result = PasswordType.NOT_DECRYPTED
    if result == PasswordType.NOT_DECRYPTED:
        raise DocumentParseError("Password-protected PDFs aren't supported.")


def _read_docx(path: Path, budget: _TextBudget) -> tuple[list[Segment], int | None]:
    document = open_docx(str(path))
    lines = [budget.take(line) for line in _docx_lines(document)]
    return [Segment(text="\n".join(lines), page=None)], None


def _docx_lines(container: Any) -> Iterator[str]:
    """Paragraphs and table rows in body order (`container` is the document or a table cell)."""
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            yield block.text
        elif isinstance(block, DocxTable):
            yield from _docx_table_lines(block)


def _docx_table_lines(table: DocxTable) -> Iterator[str]:
    for row in table.rows:
        cells: list[str] = []
        previous = None
        for cell in row.cells:
            if cell is previous:  # a horizontally merged cell is repeated once per grid column
                continue
            previous = cell
            cells.append(" ".join(" ".join(_docx_lines(cell)).split()))
        if any(cells):
            yield " | ".join(cells)


def _read_pptx(path: Path, budget: _TextBudget) -> tuple[list[Segment], int | None]:
    presentation = Presentation(str(path))
    segments: list[Segment] = []
    for number, slide in enumerate(presentation.slides, start=1):
        budget.check_time()
        lines = list(_pptx_shape_lines(slide.shapes))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append(f"Notes: {notes.text}")
        segments.append(Segment(text=budget.take("\n".join(lines)), page=number))
    return segments, len(presentation.slides)


def _pptx_shape_lines(shapes: Iterable[Any]) -> Iterator[str]:
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _pptx_shape_lines(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame.text
        elif shape.has_table:
            for row in shape.table.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells if not cell.is_spanned]
                if any(cells):
                    yield " | ".join(cells)


def _read_xlsx(path: Path, budget: _TextBudget) -> tuple[list[Segment], int | None]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        segments: list[Segment] = []
        cells_seen = 0
        for sheet in workbook.worksheets:
            # The stored dimensions are untrusted; without them openpyxl doesn't pad rows to a claimed width.
            sheet.reset_dimensions()
            lines: list[str] = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_number > _XLSX_MAX_ROWS:
                    break
                if row_number % _XLSX_DEADLINE_ROWS == 0:
                    budget.check_time()
                cells_seen += max(1, len(row))
                if cells_seen > _XLSX_MAX_CELLS:
                    raise DocumentParseError("The spreadsheet is too large to process.")
                if line := _xlsx_row_text(row):
                    lines.append(budget.take(line))
            if lines:
                segments.append(Segment(text="\n".join([f"Sheet: {sheet.title}", *lines]), page=None))
        return segments, None
    finally:
        workbook.close()


def _xlsx_row_text(values: tuple[Any, ...]) -> str:
    cells = [_xlsx_cell_text(value) for value in values]
    while cells and not cells[-1]:
        cells.pop()
    return " | ".join(cells)


def _xlsx_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime) and not (value.hour or value.minute or value.second or value.microsecond):
        return value.date().isoformat()
    return " ".join(str(value).split())


def _read_html(path: Path, budget: _TextBudget) -> tuple[list[Segment], int | None]:
    # The tree takes many times the file's size in memory, and markup doesn't count toward the text cap.
    if path.stat().st_size > settings.max_extracted_chars * _MAX_BYTES_PER_CHAR:
        raise DocumentParseError(_TOO_COMPLEX)
    soup = BeautifulSoup(path.read_bytes(), "lxml")
    title = " ".join(soup.title.get_text(" ").split()) if soup.title else ""
    body = "\n".join(" ".join(line.split()) for line in _html_text(soup).splitlines())
    text = f"{title}\n\n{body}" if title else body
    return [Segment(text=budget.take(text), page=None)], None


def _html_text(root: Tag) -> str:
    """Text with blocks on lines of their own and table cells joined by " | ".

    One iterative pass that leaves the tree untouched: inserting into it costs time quadratic in the nesting depth.
    """
    parts: list[str] = []
    stack: list[Any] = list(reversed(root.contents))
    while stack:
        node = stack.pop()
        if node is _NEWLINE:
            parts.append("\n")
        elif node is _CELL_SEPARATOR:
            parts.append(" | ")
        elif isinstance(node, Tag):
            if node.name in _HTML_DROPPED:
                continue
            if node.name == "br":
                parts.append("\n")
                continue
            children: list[Any] = node.contents
            if node.name == "tr":  # " | " before every direct td/th except the first
                children, first_cell = [], True
                for child in node.contents:
                    if isinstance(child, Tag) and child.name in _HTML_CELLS:
                        if not first_cell:
                            children.append(_CELL_SEPARATOR)
                        first_cell = False
                    children.append(child)
            if node.name in _HTML_BLOCKS:
                parts.append("\n")
                stack.append(_NEWLINE)
            stack.extend(reversed(children))
        elif type(node) in _HTML_STRING_TYPES:
            parts.append(node)
    return "".join(parts)


def _read_text(path: Path, budget: _TextBudget) -> tuple[list[Segment], int | None]:
    _check_text_size(path)
    return [Segment(text=budget.take(_decode(path.read_bytes())), page=None)], None


def _read_json(path: Path, budget: _TextBudget) -> tuple[list[Segment], int | None]:
    _check_text_size(path)
    text = _decode(path.read_bytes())
    if len(text) > settings.max_extracted_chars:  # don't build the object only to refuse its text
        raise _too_much_text()
    try:
        pretty = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except (ValueError, RecursionError):
        pass
    else:
        if len(pretty) <= settings.max_extracted_chars:
            text = pretty
    return [Segment(text=budget.take(text), page=None)], None


def _check_text_size(path: Path) -> None:
    """Refuse a text file that is certain to exceed the text cap before reading it into memory."""
    if path.stat().st_size > settings.max_extracted_chars * _MAX_BYTES_PER_CHAR:
        raise _too_much_text()


def _decode(data: bytes) -> str:
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


# PDFs are read by _read_pdf, which also selects pages for the vision model.
_READERS: dict[str, _Reader] = {
    ".docx": _read_docx,
    ".pptx": _read_pptx,
    ".xlsx": _read_xlsx,
    ".html": _read_html,
    ".htm": _read_html,
    ".md": _read_text,
    ".markdown": _read_text,
    ".txt": _read_text,
    ".csv": _read_text,
    ".json": _read_json,
}
